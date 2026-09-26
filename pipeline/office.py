"""
pipeline/office.py — Word · PowerPoint · Excel · رسوم: إنشاءٌ وقراءةٌ وتعديل
=========================================================================
الواجهةُ التي تناديها مهاراتُ المحرّك (word · powerpoint · excel · charts).
لا تعيد بناءَ شيء: البناءُ بأدوات Weaver الموجودة كما هي
(capabilities/skills/{docx,pptx,xlsx,chart}_builder/scripts). والجديدُ هنا
ما لم يكن موجوداً: **قراءةُ ملفٍّ بأرقام أجزائه، وتعديلُه**.

قِيس قبلها على هاتف المستخدم (tools/probe_office.py): المكتباتُ كلُّها
موجودة، والبناءُ ✓ ١٤، وتعديلُ Excel بـopenpyxl وحده لا يوسّع معادلةَ
المجموع للصفّ المُدرَج (=SUM(B2:B3) تبقى) — فهنا نُزاح المراجعُ كما يفعل
Excel.

    python3 office.py info  FILE [--find TEXT] [--rows N]
    python3 office.py build SPEC.json --out FILE.docx|.pptx|.xlsx
    python3 office.py chart SPEC.json --out FILE.png
    python3 office.py chart SPEC.json --into FILE.docx|.pptx [--out NEW]
    python3 office.py edit  FILE OPS.json [--out NEW]

قواعد:
  · الأرقامُ في OPS هي أرقامُ `info` **قبل** التعديل، ولو تتابعت العمليّات.
  · التعديلُ يُحفظ باسمٍ جديد (<الاسم>-edited.<الامتداد>) — الأصلُ لا يُمَسّ
    إلا بـ--out على الاسم نفسِه.
  · كلُّ ملفٍّ يُكتب يُعاد فتحُه قبل «✓» — لا يكفي أنّه كُتب.
  · Excel: يُقارَن الملفُّ قبل الحفظ وبعده؛ إن نقص منه شيء (رسومٌ، صور،
    أشكال، تعليقات…) لا يُكتب ويُقال السبب (إلا بـ--allow-loss).

الخروج: 0 تمّ كلُّه · 2 بعضُه (يُذكر ما لم يتمّ) · 1 فشل · 3 رُفض حمايةً.
"""
from __future__ import annotations

import copy
import json
import os
import re
import sys
import tempfile
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SK = os.path.join(_ROOT, "capabilities", "skills")
_SCRIPTS = [os.path.join(_SK, d, "scripts") for d in
            ("docx_builder", "pptx_builder", "xlsx_builder", "chart_builder")]
_FONTS = os.path.join(_ROOT, "engines", "fonts-core")


def _paths():
    for p in _SCRIPTS + [_FONTS]:
        if p not in sys.path:
            sys.path.append(p)


_AR = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")


def _is_ar(text):
    t = str(text or "")
    letters = sum(1 for c in t if c.isalpha())
    return letters > 0 and len(_AR.findall(t)) / letters >= 0.4


def _lang_of(spec, *texts):
    lg = str((spec or {}).get("lang") or "").lower()
    if lg in ("ar", "en"):
        return lg
    return "ar" if _is_ar(" ".join(str(t or "") for t in texts)) else "en"


class Fail(Exception):
    """خطأٌ يُقال للنموذج كما هو (لا تتبّعَ بايثون)."""


def _load_json(path):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        raise Fail("spec file not found: %s" % path)
    except json.JSONDecodeError as e:
        raise Fail("bad JSON in %s: %s" % (path, e))


def _ext(path):
    return os.path.splitext(str(path))[1].lower()


def _old_format(path):
    e = _ext(path)
    if e in (".doc", ".ppt", ".xls"):
        raise Fail("%s is the old binary Office format — it cannot be read or "
                   "edited here. Ask the user to save it as %sx and send it "
                   "again." % (os.path.basename(path), e))
    if e in (".xlsm", ".docm", ".pptm"):
        raise Fail("%s contains macros; editing it would drop them. Not edited."
                   % os.path.basename(path))


def _need(path):
    if not os.path.isfile(path):
        raise Fail("file not found: %s" % path)
    _old_format(path)


def _no_dupes(path):
    """ملفٌّ فيه اسمٌ مكرّر ليس ملفّاً سليماً — يُحذف ويُقال."""
    try:
        names = zipfile.ZipFile(path).namelist()
    except Exception as e:
        raise Fail("saved file is not a valid Office file: %s" % e)
    if len(names) != len(set(names)):
        dup = sorted({n for n in names if names.count(n) > 1})
        try:
            os.remove(path)
        except OSError:
            pass
        raise Fail("internal error: duplicate parts %s — nothing written"
                   % ", ".join(dup[:3]))


def _default_out(path):
    stem, ext = os.path.splitext(path)
    return stem + "-edited" + ext


# ═════════════════════════════════ info ═════════════════════════════════
def _short(t, n=160):
    t = " ".join(("" if t is None else str(t)).split())
    return t if len(t) <= n else t[:n] + "…"


def info_docx(path, find=None):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = Document(path)
    paras = d.paragraphs
    pidx = {p._p: i for i, p in enumerate(paras)}
    tidx = {t._tbl: i for i, t in enumerate(d.tables)}
    imgs = sum(1 for n in zipfile.ZipFile(path).namelist()
               if n.startswith("word/media/"))
    rtl = any(p._p.pPr is not None and p._p.pPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi")
        is not None for p in paras)
    out = ["WORD %s — %d paragraphs (¶) · %d tables (T) · %d images · %s"
           % (os.path.basename(path), len(paras), len(d.tables), imgs,
              "RTL" if rtl else "LTR")]
    for s_i, sec in enumerate(d.sections):
        for kind, part in (("header", sec.header), ("footer", sec.footer)):
            try:
                txt = " ".join(p.text for p in part.paragraphs if p.text.strip())
            except Exception:
                txt = ""
            if txt:
                out.append("  [%s] %s" % (kind, _short(txt, 100)))
    body = d.element.body
    for el in body.iterchildren():
        if el in pidx:
            p = Paragraph(el, d)
            i = pidx[el]
            if find and find not in p.text:
                continue
            has_img = bool(el.xpath(".//*[local-name()='drawing']"))
            if not p.text.strip() and not has_img:
                if not find:
                    out.append("  ¶%d (empty)" % i)
                continue
            out.append("  ¶%d [%s]%s %s" % (i, p.style.name if p.style is not None
                                              else "?",
                                              " [image]" if has_img else "",
                                              _short(p.text)))
        elif el in tidx:
            t = Table(el, d)
            i = tidx[el]
            rows = t.rows
            ncols = len(t.columns)
            cells = [[c.text for c in r.cells] for r in rows]
            if find and not any(find in c for r in cells for c in r):
                continue
            out.append("  T%d %d×%d (rows×cols)" % (i, len(rows), ncols))
            for r_i, r in enumerate(cells[:12]):
                out.append("    r%d: %s" % (r_i, " | ".join(_short(c, 30) for c in r)))
            if len(cells) > 12:
                out.append("    … %d more rows" % (len(cells) - 12))
    return "\n".join(out)


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:        # GROUP
            yield from _iter_shapes(sh.shapes)


def info_pptx(path, find=None):
    from pptx import Presentation
    prs = Presentation(path)
    out = ["POWERPOINT %s — %d slides · %.1f×%.1f in"
           % (os.path.basename(path), len(prs.slides),
              prs.slide_width / 914400, prs.slide_height / 914400)]
    for n, s in enumerate(prs.slides, 1):
        lines = []
        for k, sh in enumerate(s.shapes):
            kind = ("title" if sh.is_placeholder and sh.placeholder_format.type
                    in (1, 3) else "table" if sh.has_table else
                    "picture" if sh.shape_type == 13 else
                    "chart" if getattr(sh, "has_chart", False) else
                    "group" if sh.shape_type == 6 else "text"
                    if sh.has_text_frame else "shape")
            if sh.has_text_frame:
                txt = "\n".join(p.text for p in sh.text_frame.paragraphs)
                if not txt.strip():
                    continue
                if find and find not in txt:
                    continue
                paras = [p.text for p in sh.text_frame.paragraphs if p.text.strip()]
                lines.append("    #%d %s: %s" % (k, kind, _short(paras[0], 120)))
                for p in paras[1:8]:
                    lines.append("         · %s" % _short(p, 110))
                if len(paras) > 8:
                    lines.append("         … %d more lines" % (len(paras) - 8))
            elif sh.has_table:
                cells = [[c.text for c in r.cells] for r in sh.table.rows]
                if find and not any(find in c for r in cells for c in r):
                    continue
                lines.append("    #%d table %d×%d" % (k, len(cells),
                                                     len(cells[0]) if cells else 0))
                for r_i, r in enumerate(cells[:8]):
                    lines.append("         r%d: %s" % (r_i, " | ".join(
                        _short(c, 25) for c in r)))
            elif not find and kind in ("picture", "chart", "group"):
                lines.append("    #%d %s" % (k, kind))
        notes = ""
        if s.has_notes_slide:
            notes = s.notes_slide.notes_text_frame.text if \
                s.notes_slide.notes_text_frame is not None else ""
        if find and not lines and find not in notes:
            continue
        out.append("  slide %d [%s]" % (n, s.slide_layout.name))
        out.extend(lines)
        if notes.strip():
            out.append("    notes: %s" % _short(notes, 120))
    return "\n".join(out)


# ما يُسقطه openpyxl عند الحفظ — **يُقاس ولا يُفترض**. كان هنا رفضٌ مسبقٌ
# لكلِّ ملفٍّ فيه رسوم، ثمّ قِيس: openpyxl 3.1.5 يقرأ الرسومَ والصورَ ويحفظها
# (chart1.xml وimage1.png باقيان بعد load+save). فالحكمُ الآن بالمقارنة:
# أجزاءُ الملفّ قبل الحفظ وبعده، وعددُ الرسوم والصور والأشكال.
# وما يُسقطه ولا يضرّ: إعداداتُ الطابعة، وسلسلةُ الحساب (تُعاد عند الفتح).
_XLSX_HARMLESS = ("xl/printerSettings/", "xl/calcChain.xml")


def _xlsx_parts(path):
    """{صنفُ الجزء: عدد} + عددُ الأشكال في الرسومات (xdr:sp)."""
    out = {}
    try:
        z = zipfile.ZipFile(path)
    except Exception:
        return out
    for n in z.namelist():
        if n.endswith("/") or n.startswith(_XLSX_HARMLESS):
            continue
        k = re.sub(r"\d+", "#", n)
        out[k] = out.get(k, 0) + 1
        if n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
            try:
                x = z.read(n).decode("utf-8", "ignore")
                out["<shapes>"] = out.get("<shapes>", 0) + len(
                    re.findall(r"<(?:xdr:)?sp[ >]", x))
            except Exception:
                pass
    return out


_PART_LABEL = (("xl/charts/", "charts"), ("xl/media/", "images"),
               ("xl/pivot", "pivot tables"), ("xl/slicer", "slicers"),
               ("xl/threadedComments", "threaded comments"),
               ("xl/comments", "comments"), ("xl/ctrlProps", "form controls"),
               ("xl/vbaProject", "macros"), ("<shapes>", "shapes/text boxes"),
               ("xl/drawings/", "drawings"), ("xl/tables/", "tables"))


def xlsx_loss(before, after):
    """ما نقص بين ملفَّين — قائمةُ أوصافٍ مقروءة، فارغةٌ إن لم ينقص شيء."""
    a, b = _xlsx_parts(before), _xlsx_parts(after)
    lost = []
    for k, cnt in a.items():
        if b.get(k, 0) < cnt:
            label = next((l for pre, l in _PART_LABEL if k.startswith(pre)), k)
            item = "%s (%d→%d)" % (label, cnt, b.get(k, 0))
            if item not in lost:
                lost.append(item)
    return lost


def xlsx_lossy(path):
    """ما سيُسقطه الحفظُ بهذه الأداة — بحفظٍ تجريبيٍّ في ملفٍّ مؤقّت."""
    from openpyxl import load_workbook
    fd, tmp = tempfile.mkstemp(prefix="weaver-rt-", suffix=".xlsx")
    os.close(fd)
    try:
        load_workbook(path).save(tmp)
        return xlsx_loss(path, tmp)
    except Exception as e:
        return ["could not round-trip: %s" % str(e)[:100]]
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


def info_xlsx(path, find=None, rows=40):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    try:
        vals = load_workbook(path, data_only=True)
    except Exception:
        vals = None
    out = ["EXCEL %s — %d sheets" % (os.path.basename(path), len(wb.worksheets))]
    lossy = xlsx_lossy(path)
    if lossy:
        out.append("  ⚠ editing with this tool would DROP: %s — tell the user "
                   "before editing" % ", ".join(lossy))
    for ws in wb.worksheets:
        out.append("  sheet \"%s\" — %s · %s%s" % (
            ws.title, ws.dimensions, "RTL" if ws.sheet_view.rightToLeft else "LTR",
            (" · merged: " + ", ".join(str(m) for m in list(ws.merged_cells.ranges)[:8]))
            if ws.merged_cells.ranges else ""))
        if getattr(ws, "tables", None):
            for tn, t in ws.tables.items():
                out.append("    table %s: %s" % (tn, t.ref))
        vws = vals[ws.title] if vals is not None else None
        shown = 0
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    cached = vws[c.coordinate].value if vws is not None else None
                    s = "%s %s → %s" % (c.coordinate, v,
                                        "?" if cached is None else cached)
                else:
                    s = "%s %s" % (c.coordinate, _short(v, 40))
                cells.append(s)
            if not cells:
                continue
            line = " · ".join(cells)
            if find and find not in line:
                continue
            out.append("    " + line)
            shown += 1
            if shown >= rows:
                out.append("    … (more rows; use --rows N or --find TEXT)")
                break
    return "\n".join(out)


def cmd_info(path, find=None, rows=40):
    _need(path)
    e = _ext(path)
    if e == ".docx":
        return info_docx(path, find)
    if e == ".pptx":
        return info_pptx(path, find)
    if e == ".xlsx":
        return info_xlsx(path, find, rows)
    raise Fail("info supports .docx .pptx .xlsx — got %s" % (e or "no extension"))


# ═════════════════════════════════ charts ═══════════════════════════════
def chart_png(spec, out):
    """spec: {type, data, title, xlabel, ylabel, theme, lang}."""
    _paths()
    from build_chart import build_chart
    lang = _lang_of(spec, spec.get("title"), " ".join(
        str(x) for x in (spec.get("data") or {}).get("labels", []) or []))
    r = build_chart(spec.get("type", "bar"), spec.get("data") or {}, out,
                    title=spec.get("title", ""),
                    theme_id=spec.get("theme", spec.get("theme_id", "academic_navy")),
                    xlabel=spec.get("xlabel", ""), ylabel=spec.get("ylabel", ""),
                    lang=lang)
    if not r.get("ok"):
        raise Fail("chart failed: %s" % r.get("error"))
    with open(out, "rb") as f:
        if f.read(8) != b"\x89PNG\r\n\x1a\n":
            raise Fail("chart output is not a PNG")
    return out


def _tmp_png(tag="chart"):
    fd, p = tempfile.mkstemp(prefix="weaver-%s-" % tag, suffix=".png")
    os.close(fd)
    return p


# ═════════════════════════════════ build ════════════════════════════════
def build_word(spec, out):
    _paths()
    from docx_advanced import build_rich_docx
    sections = [dict(s) for s in (spec.get("sections") or []) if isinstance(s, dict)]
    lang = _lang_of(spec, spec.get("title"), *[s.get("heading", "") + " " +
                                                str(s.get("body", ""))[:200]
                                                for s in sections[:3]])
    tmp = []
    try:
        for s in sections:
            if s.get("chart") and not s.get("image"):
                ch = dict(s["chart"])
                ch.setdefault("lang", lang)
                png = chart_png(ch, _tmp_png())
                tmp.append(png)
                s["image"] = {"path": png, "caption": ch.get("caption", "")}
        build_rich_docx(
            spec.get("title", ""), sections, output_path=out, lang=lang,
            theme_id=spec.get("theme", "academic_navy"), font=spec.get("font"),
            subtitle=spec.get("subtitle", ""), references=spec.get("references"),
            header_text=spec.get("header"),
            page_numbers=spec.get("page_numbers", True),
            toc=bool(spec.get("toc", False)),
            two_columns=bool(spec.get("two_columns", False)),
            cover=spec.get("cover"))
    finally:
        for p in tmp:
            try:
                os.remove(p)
            except OSError:
                pass
    from docx import Document
    d = Document(out)
    imgs = sum(1 for n in zipfile.ZipFile(out).namelist()
               if n.startswith("word/media/"))
    return "Word %s: %d paragraphs · %d tables · %d images · %s" % (
        lang.upper(), len(d.paragraphs), len(d.tables), imgs,
        "RTL" if lang == "ar" else "LTR")


def _slide_kind(s):
    if s.get("table"):
        return "table"
    if s.get("chart"):
        return "chart"
    return "normal"


def _pptx_chart_slide(prs, spec, title, rtl):
    """شريحةُ رسمٍ على كائن العرض نفسِه (embed_chart يعمل على مسارٍ فقط)."""
    _paths()
    from pptx.util import Inches
    from PIL import Image
    import build_pptx as bp
    png = chart_png(spec, _tmp_png())
    try:
        layouts = prs.slide_layouts
        blank = next((l for l in layouts if l.name.strip().lower() == "blank"),
                     None) or min(layouts, key=lambda l: len(l.placeholders))
        slide = prs.slides.add_slide(blank)
        top = Inches(0.4)
        if title:
            tb = slide.shapes.add_textbox(Inches(0.7), top,
                                          prs.slide_width - Inches(1.4), Inches(0.9))
            tb.text_frame.word_wrap = True
            bp._add_text(tb.text_frame, title, size=26, bold=True, color=bp.NAVY,
                         align=None, rtl=rtl)
            top = Inches(1.4)
        iw, ih = Image.open(png).size
        avail_h = prs.slide_height - top - Inches(0.3)
        w = min(prs.slide_width - Inches(1.4), int(avail_h * iw / ih))
        h = int(w * ih / iw)
        slide.shapes.add_picture(png, int((prs.slide_width - w) / 2), top,
                                 width=w, height=h)
        return slide
    finally:
        try:
            os.remove(png)
        except OSError:
            pass


def _reorder(prs, order_ids):
    """رتّب الشرائحَ بقائمة sldId (عناصرُ XML)."""
    lst = prs.slides._sldIdLst
    for el in list(lst):
        lst.remove(el)
    for el in order_ids:
        lst.append(el)


def build_powerpoint(spec, out):
    _paths()
    import build_pptx as bp
    from pptx_table import add_table_slide
    from pptx import Presentation
    slides = [s for s in (spec.get("slides") or []) if isinstance(s, dict)]
    lang = _lang_of(spec, spec.get("title"), *[s.get("title", "") + " " + " ".join(
        str(p) for p in (s.get("points") or [])[:3]) for s in slides[:4]])
    rtl = lang == "ar"
    normal = [s for s in slides if _slide_kind(s) == "normal"]
    bp.build_deck(title=spec.get("title", ""), slides=normal,
                  subtitle=spec.get("subtitle", ""), output_path=out, lang=lang,
                  closing=spec.get("closing"))
    prs = Presentation(out)
    ids = list(prs.slides._sldIdLst)          # غلاف · العاديّة · الختام
    cover, closing, body = ids[0], ids[-1], iter(ids[1:-1])
    order = [cover]
    for s in slides:
        k = _slide_kind(s)
        if k == "normal":
            order.append(next(body))
            continue
        if k == "table":
            t = s["table"]
            add_table_slide(prs, t.get("headers", []), t.get("rows", []),
                            lang=lang, theme_id=spec.get("theme", "academic_navy"),
                            title=s.get("title", ""), totals=t.get("totals"))
        else:
            ch = dict(s["chart"])
            ch.setdefault("lang", lang)
            _pptx_chart_slide(prs, ch, s.get("title", ""), rtl)
        order.append(prs.slides._sldIdLst[-1])
    order.append(closing)
    _reorder(prs, order)
    prs.save(out)
    p2 = Presentation(out)
    return "PowerPoint %s: %d slides · %s" % (lang.upper(), len(p2.slides),
                                              "RTL" if rtl else "LTR")


def _copy_sheet(src, dst):
    for row in src.iter_rows():
        for c in row:
            n = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                n.font = copy.copy(c.font)
                n.fill = copy.copy(c.fill)
                n.border = copy.copy(c.border)
                n.alignment = copy.copy(c.alignment)
                n.number_format = c.number_format
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    dst.sheet_view.rightToLeft = src.sheet_view.rightToLeft


def _xl_chart(ws, ch):
    """رسمٌ أصليٌّ في Excel (يبقى قابلاً للتعديل فيه). ch: {type,title,data,
    categories,anchor} — data مثل "B1:C5" (بالترويسة)، categories مثل "A2:A5"."""
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.utils.cell import range_boundaries
    kind = str(ch.get("type", "bar")).lower()
    obj = {"bar": BarChart, "column": BarChart, "line": LineChart,
           "pie": PieChart}.get(kind)
    if obj is None:
        raise Fail("excel chart type must be bar/line/pie — got %s" % kind)
    c = obj()
    if kind == "bar" and hasattr(c, "type"):
        c.type = "col"
    c.title = ch.get("title") or None
    c1, r1, c2, r2 = range_boundaries(ch["data"])
    c.add_data(Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2),
               titles_from_data=True)
    if ch.get("categories"):
        a1, b1, a2, b2 = range_boundaries(ch["categories"])
        c.set_categories(Reference(ws, min_col=a1, min_row=b1, max_col=a2,
                                   max_row=b2))
    c.width, c.height = 16, 8
    ws.add_chart(c, ch.get("anchor") or "%s2" % _col_letter(
        (ws.max_column or 1) + 2))


def _col_letter(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def build_excel(spec, out):
    _paths()
    from build_xlsx import build_xlsx
    from openpyxl import Workbook, load_workbook
    sheets = spec.get("sheets")
    if not sheets:
        sheets = [{k: spec.get(k) for k in ("name", "headers", "rows", "totals",
                                            "formats", "chart", "widths")}]
    lang = _lang_of(spec, *[" ".join(str(h) for h in (s.get("headers") or []))
                            for s in sheets])
    tmpdir = tempfile.mkdtemp(prefix="weaver-xlsx-")
    try:
        wb = Workbook()
        wb.remove(wb.active)
        for i, s in enumerate(sheets):
            part = os.path.join(tmpdir, "%d.xlsx" % i)
            build_xlsx(s.get("rows") or [], part, headers=s.get("headers"),
                       lang=lang, sheet_name=(s.get("name") or None),
                       with_totals=bool(s.get("totals")))
            src = load_workbook(part).active
            dst = wb.create_sheet(str(s.get("name") or src.title)[:31])
            _copy_sheet(src, dst)
            for col, fmt in (s.get("formats") or {}).items():
                for c in dst[col]:
                    if isinstance(c.value, (int, float)) or (
                            isinstance(c.value, str) and c.value.startswith("=")):
                        c.number_format = fmt
            for col, w in (s.get("widths") or {}).items():
                dst.column_dimensions[col].width = w
            if s.get("chart"):
                _xl_chart(dst, s["chart"])
            dst.freeze_panes = "A2" if s.get("headers") else None
        wb.calculation.fullCalcOnLoad = True
        wb.save(out)
    finally:
        import shutil
        shutil.rmtree(tmpdir, ignore_errors=True)
    wb2 = load_workbook(out)
    return "Excel %s: %d sheets (%s) · %s" % (
        lang.upper(), len(wb2.worksheets),
        ", ".join("%s %s" % (w.title, w.dimensions) for w in wb2.worksheets),
        "RTL" if lang == "ar" else "LTR")


def cmd_build(spec_path, out):
    spec = _load_json(spec_path)
    e = _ext(out)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    if e == ".docx":
        return build_word(spec, out)
    if e == ".pptx":
        return build_powerpoint(spec, out)
    if e == ".xlsx":
        return build_excel(spec, out)
    raise Fail("--out must end with .docx .pptx or .xlsx — got %s" % (e or "none"))


def cmd_chart(spec_path, out=None, into=None):
    spec = _load_json(spec_path)
    if not into:
        if _ext(out or "") != ".png":
            raise Fail("--out must be a .png (or use --into FILE.docx/.pptx)")
        chart_png(spec, out)
        return "chart %s → %s" % (spec.get("type", "bar"), out)
    _need(into)
    dst = out or _default_out(into)
    if _ext(into) == ".docx":
        ops = [{"op": "add_chart", "chart": spec,
                "caption": spec.get("caption", "")}]
        if spec.get("after") is not None:
            ops[0]["after"] = spec["after"]
        res = edit_docx(into, ops, dst)
    elif _ext(into) == ".pptx":
        op = {"op": "add_chart_slide", "chart": spec,
              "title": spec.get("slide_title", spec.get("title", ""))}
        if spec.get("after") is not None:
            op["after"] = spec["after"]
        res = edit_pptx(into, [op], dst)
    else:
        raise Fail("--into must be a .docx or .pptx")
    if res["failed"]:
        raise Fail("; ".join(res["failed"]))
    return "chart added → %s" % dst


# ═════════════════════════════════ edit: Word ═══════════════════════════
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _replace_in_para(p, find, repl):
    """استبدالٌ في فقرة يحفظ التنسيق: داخلَ الـrun إن كان كلُّه فيها؛ فإن
    تفرّق النصُّ على runs دُمجت هذه الفقرةُ وحدَها في أوّلها. يعيد العدد.

    العدُّ من النصّ الأصليّ مرّةً واحدة — قِيس: بديلٌ يحوي المبحوثَ عنه
    («نصّ» ⟵ «نصّ معدَّل») كان يُستبدل مرّتين."""
    orig = p.text
    total = orig.count(find)
    if not total:
        return 0
    runs = p.runs
    if sum(r.text.count(find) for r in runs) == total:
        for r in runs:
            if find in r.text:
                r.text = r.text.replace(find, repl)
        return total
    joined = "".join(r.text for r in runs)
    if joined != orig or not runs:
        return 0                # روابطُ تشعّبيّة داخلها — لا يُدمج ما لا يُرى
    runs[0].text = orig.replace(find, repl)
    for r in runs[1:]:
        r.text = ""
    return total


def _all_doc_paragraphs(d):
    def from_tables(tables):
        for t in tables:
            for row in t.rows:
                for c in row.cells:
                    yield from c.paragraphs
                    yield from from_tables(c.tables)
    yield from d.paragraphs
    yield from from_tables(d.tables)
    for sec in d.sections:
        for part in (sec.header, sec.footer, sec.first_page_header,
                     sec.first_page_footer):
            try:
                if part.is_linked_to_previous and sec is not d.sections[0]:
                    continue
                yield from part.paragraphs
                yield from from_tables(part.tables)
            except Exception:
                continue


def _set_para_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def _new_para_after(anchor_el, template, text, style=None):
    """فقرةٌ جديدةٌ بعد عنصر: تنسيقُ الفقرة والخطِّ من `template`."""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    new = OxmlElement("w:p")
    anchor_el.addnext(new)
    para = Paragraph(new, template._parent)
    if template._p.pPr is not None:
        new.insert(0, copy.deepcopy(template._p.pPr))
    if style:
        try:
            para.style = style
        except Exception:
            pass
    run = para.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, copy.deepcopy(template.runs[0]._r.rPr))
    if _is_ar(text):
        pPr = new.get_or_add_pPr()
        if pPr.find(W_NS + "bidi") is None:
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            pPr.append(OxmlElement("w:bidi"))
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return para


def _doc_lang(d):
    for p in d.paragraphs[:80]:
        if p._p.pPr is not None and p._p.pPr.find(W_NS + "bidi") is not None:
            return "ar"
    return "ar" if _is_ar(" ".join(p.text for p in d.paragraphs[:40])) else "en"


def _append_then_move(d, fn, anchor_el):
    """أضف في آخر المستند بدالّةٍ موجودة، ثمّ انقل ما أُضيف بعد `anchor_el`."""
    body = d.element.body
    # عناصرُ lxml: الهويّةُ ثابتةٌ ما دام الكائنُ حيّاً — فالقائمةُ تُمسكها.
    # قِيس: id() بلا إمساكٍ نقل عناصرَ قديمةً ظنّها جديدة.
    before_list = list(body.iterchildren())
    fn()
    new = [c for c in body.iterchildren()
           if not any(c is b for b in before_list)
           and c.tag != W_NS + "sectPr"]
    if anchor_el is not None:
        cur = anchor_el
        for el in new:
            cur.addnext(el)
            cur = el
    return new


def edit_docx(path, ops, out):
    _paths()
    from docx import Document
    d = Document(path)
    paras = list(d.paragraphs)                  # أرقامُ info قبل التعديل
    tables = list(d.tables)
    lang = _doc_lang(d)
    done, failed = [], []
    gone = set()

    def P(i, key):
        try:
            i = int(i)
        except (TypeError, ValueError):
            raise Fail("%s must be a paragraph number from info" % key)
        if not 0 <= i < len(paras):
            raise Fail("¶%d does not exist (0…%d)" % (i, len(paras) - 1))
        if i in gone:
            raise Fail("¶%d was deleted by an earlier op" % i)
        return paras[i]

    def anchor_of(op):
        if op.get("after") is not None:
            a = str(op["after"])
            if a.upper().startswith("T"):
                return tables[int(a[1:])]._tbl
            return P(op["after"], "after")._p
        if op.get("before") is not None:
            b = P(op["before"], "before")._p
            prev = b.getprevious()
            if prev is None:
                raise Fail("cannot insert before the first element; use after")
            return prev
        return None

    for n, op in enumerate(ops, 1):
        kind = str(op.get("op", ""))
        try:
            if kind == "replace":
                f, w = str(op.get("find", "")), str(op.get("with", ""))
                if not f:
                    raise Fail("replace needs find")
                c = sum(_replace_in_para(p, f, w) for p in _all_doc_paragraphs(d))
                if not c:
                    raise Fail("text not found: %s" % _short(f, 60))
                done.append("replace «%s» ×%d" % (_short(f, 40), c))
            elif kind == "set_text":
                p = P(op.get("paragraph"), "paragraph")
                _set_para_text(p, str(op.get("text", "")))
                done.append("¶%s text set" % op.get("paragraph"))
            elif kind == "insert_paragraph":
                anc = anchor_of(op)
                if anc is None:
                    kids = [c for c in d.element.body.iterchildren()
                            if c.tag != W_NS + "sectPr"]
                    anc = kids[-1] if kids else None
                style = op.get("style")
                tmpl = None
                if op.get("like") is not None:
                    tmpl = P(op["like"], "like")
                elif style:
                    tmpl = next((p for p in paras if p.style is not None
                                 and p.style.name == style), None)
                if tmpl is None:
                    ref = [p for p in paras if p._p is anc]
                    tmpl = ref[0] if ref else (paras[-1] if paras else None)
                if tmpl is None or anc is None:
                    raise Fail("document has no paragraphs to anchor to")
                cur = anc
                texts = op.get("text")
                texts = texts if isinstance(texts, list) else [texts]
                for t in texts:
                    np_ = _new_para_after(cur, tmpl, str(t or ""), style)
                    cur = np_._p
                done.append("inserted %d paragraph(s)" % len(texts))
            elif kind == "delete_paragraph":
                idx = op.get("paragraph")
                items = idx if isinstance(idx, list) else [idx]
                for i in items:
                    p = P(i, "paragraph")
                    p._p.getparent().remove(p._p)
                    gone.add(int(i))
                done.append("deleted ¶%s" % ",".join(str(i) for i in items))
            elif kind == "add_table":
                from docx_advanced import add_table
                anc = anchor_of(op)
                _append_then_move(d, lambda: add_table(
                    d, op.get("headers") or [], op.get("rows") or [], lang,
                    op.get("theme", "academic_navy"), None, op.get("totals")), anc)
                done.append("table added")
            elif kind == "set_cell":
                t = tables[int(op.get("table", 0))]
                cell = t.cell(int(op["row"]), int(op["col"]))
                if cell.paragraphs:
                    _set_para_text(cell.paragraphs[0], str(op.get("text", "")))
                    for extra in cell.paragraphs[1:]:
                        extra._p.getparent().remove(extra._p)
                else:
                    cell.text = str(op.get("text", ""))
                done.append("T%s r%s c%s set" % (op.get("table", 0), op["row"],
                                                 op["col"]))
            elif kind == "add_row":
                t = tables[int(op.get("table", 0))]
                vals = [str(v) for v in (op.get("values") or [])]
                src_row = t.rows[int(op["like_row"])] if op.get(
                    "like_row") is not None else t.rows[-1]
                new_tr = copy.deepcopy(src_row._tr)
                if op.get("after_row") is not None:
                    t.rows[int(op["after_row"])]._tr.addnext(new_tr)
                else:
                    t.rows[-1]._tr.addnext(new_tr)
                from docx.table import _Row
                row = _Row(new_tr, t)
                for ci, cell in enumerate(row.cells):
                    txt = vals[ci] if ci < len(vals) else ""
                    if cell.paragraphs:
                        _set_para_text(cell.paragraphs[0], txt)
                        for extra in cell.paragraphs[1:]:
                            extra._p.getparent().remove(extra._p)
                done.append("T%s row added" % op.get("table", 0))
            elif kind in ("add_image", "add_chart"):
                from docx_advanced import add_image
                anc = anchor_of(op)
                png, tmp = op.get("path"), None
                if kind == "add_chart":
                    ch = dict(op.get("chart") or {})
                    ch.setdefault("lang", lang)
                    png = tmp = chart_png(ch, _tmp_png())
                if not png or not os.path.isfile(png):
                    raise Fail("image not found: %s" % png)
                try:
                    _append_then_move(d, lambda: add_image(
                        d, png, str(op.get("caption", "")),
                        float(op.get("width", 5.5)), lang=lang), anc)
                finally:
                    if tmp:
                        os.remove(tmp)
                done.append("%s added" % ("chart" if tmp else "image"))
            else:
                raise Fail("unknown Word op: %s" % kind)
        except Fail as e:
            failed.append("op %d %s: %s" % (n, kind, e))
        except Exception as e:
            failed.append("op %d %s: %s: %s" % (n, kind, type(e).__name__,
                                                str(e)[:160]))
    if not done:
        return {"done": done, "failed": failed, "out": None}
    d.save(out)
    _no_dupes(out)
    Document(out)
    return {"done": done, "failed": failed, "out": out}


# ═════════════════════════════════ edit: PowerPoint ═════════════════════
def _text_frames(slide):
    for sh in _iter_shapes(slide.shapes):
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    yield c.text_frame


def _set_frame_text(tf, text):
    """نصٌّ جديدٌ في إطار: سطرٌ لكلِّ فقرة، وتنسيقُ كلِّ فقرةٍ باقٍ."""
    lines = str(text).split("\n")
    paras = list(tf.paragraphs)
    base = paras[-1] if paras else None
    for i, line in enumerate(lines):
        if i < len(paras):
            p = paras[i]
        else:
            new = copy.deepcopy(base._p)
            tf._txBody.append(new)
            p = tf.paragraphs[-1]
        if p.runs:
            p.runs[0].text = line
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.add_run().text = line
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)


def _uses_placeholders(prs):
    lay = None
    for l in prs.slide_layouts:
        types = {ph.placeholder_format.type for ph in l.placeholders}
        if 1 in types and (2 in types or 7 in types):       # TITLE + BODY/OBJECT
            lay = l
            break
    if lay is None:
        return None
    for s in prs.slides:
        if any(sh.is_placeholder and sh.placeholder_format.type in (2, 7)
               and sh.has_text_frame and sh.text_frame.text.strip()
               for sh in s.shapes):
            return lay
    return None


def _move_after(prs, new_id, anchor_id):
    lst = prs.slides._sldIdLst
    lst.remove(new_id)
    if anchor_id is None:
        lst.append(new_id)
    else:
        anchor_id.addnext(new_id)


def edit_pptx(path, ops, out):
    _paths()
    from pptx import Presentation
    import build_pptx as bp
    prs = Presentation(path)
    slides = list(prs.slides)                     # أرقامُ info قبل التعديل
    ids = list(prs.slides._sldIdLst)
    deck_ar = _is_ar(" ".join(tf.text for s in slides[:6] for tf in _text_frames(s)))
    done, failed = [], []
    gone = set()

    def S(n, key="slide"):
        try:
            n = int(n)
        except (TypeError, ValueError):
            raise Fail("%s must be a slide number from info" % key)
        if not 1 <= n <= len(slides):
            raise Fail("slide %d does not exist (1…%d)" % (n, len(slides)))
        if n in gone:
            raise Fail("slide %d was deleted by an earlier op" % n)
        return slides[n - 1], ids[n - 1]

    def anchor(op):
        if op.get("after") is None:
            return None
        return S(op["after"], "after")[1]

    for n, op in enumerate(ops, 1):
        kind = str(op.get("op", ""))
        try:
            if kind == "replace":
                f, w = str(op.get("find", "")), str(op.get("with", ""))
                if not f:
                    raise Fail("replace needs find")
                c = 0
                for i, s in enumerate(slides, 1):
                    if i in gone:
                        continue
                    for tf in _text_frames(s):
                        for p in tf.paragraphs:
                            c += _replace_in_para(p, f, w)
                    if op.get("notes") and s.has_notes_slide:
                        for p in s.notes_slide.notes_text_frame.paragraphs:
                            c += _replace_in_para(p, f, w)
                if not c:
                    raise Fail("text not found: %s" % _short(f, 60))
                done.append("replace «%s» ×%d" % (_short(f, 40), c))
            elif kind == "set_text":
                s, _ = S(op.get("slide"))
                k = int(op.get("shape"))
                shapes = list(s.shapes)
                if not 0 <= k < len(shapes) or not shapes[k].has_text_frame:
                    raise Fail("slide %s has no text shape #%d" % (op.get("slide"), k))
                _set_frame_text(shapes[k].text_frame, op.get("text", ""))
                done.append("slide %s #%d text set" % (op.get("slide"), k))
            elif kind == "notes":
                s, _ = S(op.get("slide"))
                s.notes_slide.notes_text_frame.text = str(op.get("text", ""))
                done.append("slide %s notes set" % op.get("slide"))
            elif kind == "delete_slide":
                items = op.get("slide")
                items = items if isinstance(items, list) else [items]
                for i in items:
                    _, sid = S(i)
                    prs.part.drop_rel(sid.rId)
                    prs.slides._sldIdLst.remove(sid)
                    gone.add(int(i))
                # الشريحةُ الجديدةُ تُسمّى slide<العدد+1>.xml، والأسماءُ لا
                # يُعاد ترقيمُها إلا مرّةً عند فتح العرض (python-pptx) — قِيس:
                # حذفٌ ثمّ إضافةٌ ⟵ slide7.xml مرّتين ⟵ ملفٌّ فاسد.
                prs.part.rename_slide_parts(
                    [x.rId for x in prs.slides._sldIdLst])
                done.append("deleted slide %s" % ",".join(str(i) for i in items))
            elif kind == "move_slide":
                _, sid = S(op.get("slide"))
                to = int(op.get("to"))
                lst = prs.slides._sldIdLst
                lst.remove(sid)
                to = max(1, min(to, len(lst) + 1))
                if to > len(lst):
                    lst.append(sid)
                else:
                    lst[to - 1].addprevious(sid)
                done.append("slide %s → position %d" % (op.get("slide"), to))
            elif kind == "add_slide":
                title = str(op.get("title", ""))
                points = [str(p) for p in (op.get("points") or [])]
                rtl = _is_ar(title + " " + " ".join(points)) if (title or points) \
                    else deck_ar
                lay = _uses_placeholders(prs)
                if lay is not None:
                    sl = prs.slides.add_slide(lay)
                    for ph in sl.placeholders:
                        t = ph.placeholder_format.type
                        if t in (1, 3):
                            ph.text_frame.text = title
                        elif t in (2, 7) and points:
                            _set_frame_text(ph.text_frame, "\n".join(points))
                else:
                    bp._content_slide(prs, title, points, rtl, bp.AR_FONT,
                                      bp.EN_FONT)
                _move_after(prs, prs.slides._sldIdLst[-1], anchor(op))
                done.append("slide added (%s)" % ("deck layout" if lay is not None
                                                  else "Weaver design"))
            elif kind == "add_table_slide":
                from pptx_table import add_table_slide
                t = op.get("table") or op
                lang = "ar" if _is_ar(" ".join(str(h) for h in t.get(
                    "headers") or []) + " " + str(op.get("title", ""))) else "en"
                add_table_slide(prs, t.get("headers") or [], t.get("rows") or [],
                                lang=lang, title=op.get("title", ""),
                                totals=t.get("totals"))
                _move_after(prs, prs.slides._sldIdLst[-1], anchor(op))
                done.append("table slide added")
            elif kind == "add_chart_slide":
                ch = dict(op.get("chart") or {})
                rtl = _is_ar(str(op.get("title", "")) + str(ch.get("title", ""))) \
                    or deck_ar
                ch.setdefault("lang", "ar" if rtl else "en")
                _pptx_chart_slide(prs, ch, op.get("title", ""), rtl)
                _move_after(prs, prs.slides._sldIdLst[-1], anchor(op))
                done.append("chart slide added")
            else:
                raise Fail("unknown PowerPoint op: %s" % kind)
        except Fail as e:
            failed.append("op %d %s: %s" % (n, kind, e))
        except Exception as e:
            failed.append("op %d %s: %s: %s" % (n, kind, type(e).__name__,
                                                str(e)[:160]))
    if not done:
        return {"done": done, "failed": failed, "out": None}
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")      # «Duplicate name» — يفحصه _no_dupes
        prs.save(out)
    _no_dupes(out)
    Presentation(out)
    return {"done": done, "failed": failed, "out": out}


# ═════════════════════════════════ edit: Excel ══════════════════════════
_REF = re.compile(
    r"^(?P<sheet>(?:'(?:[^']|'')+'|[^'!]+)!)?"
    r"(?P<a>\$?[A-Za-z]{1,3}\$?\d+|\$?\d+|\$?[A-Za-z]{1,3})"
    r"(?::(?P<b>\$?[A-Za-z]{1,3}\$?\d+|\$?\d+|\$?[A-Za-z]{1,3}))?$")
_PART = re.compile(r"^(\$?)([A-Za-z]{0,3})(\$?)(\d*)$")


def _sheet_of(prefix, own):
    if not prefix:
        return own
    s = prefix[:-1]
    if s.startswith("'"):
        s = s[1:-1].replace("''", "'")
    return s


def _row_of(part):
    m = _PART.match(part)
    return (m, int(m.group(4))) if m and m.group(4) else (m, None)


def _with_row(m, row):
    return "%s%s%s%d" % (m.group(1), m.group(2), m.group(3), row)


def shift_ref(ref, own_sheet, target, at, n, cell_row=None):
    """أزِح مرجعاً لإدراج `n` صفوفٍ عند `at` (n سالبٌ للحذف) في `target`.

    كـExcel: ما عند `at` فأسفل يُزاح. وزيادة: مدىً عموديٌّ ينتهي فوق `at`
    مباشرةً وخليّةُ معادلته تحته (مجموعٌ أسفلَ جدوله) ⟵ يتّسع للصفوف
    المُدرَجة — فإضافةُ صفٍّ قبل سطر «الإجمالي» تدخل في مجموعه."""
    m = _REF.match(ref)
    if not m or _sheet_of(m.group("sheet"), own_sheet) != target:
        return ref
    pre = m.group("sheet") or ""
    ma, ra = _row_of(m.group("a"))
    b = m.group("b")
    mb, rb = _row_of(b) if b else (None, None)
    if ra is None or (b and rb is None):
        return ref                         # عمودٌ كامل (A:A) — لا صفوف
    if n > 0:
        new_a = ra + n if ra >= at else ra
        if b is None:
            return pre + _with_row(ma, new_a)
        new_b = rb + n if rb >= at else rb
        if (rb == at - 1 and ra < rb and cell_row is not None
                and cell_row >= at):
            new_b = rb + n
        return pre + _with_row(ma, new_a) + ":" + _with_row(mb, new_b)
    k = -n
    last = at + k - 1

    def after(r, is_start):
        if r < at:
            return r
        if r > last:
            return r - k
        return at if is_start else at - 1
    if b is None:
        if at <= ra <= last:
            return "#REF!"
        return pre + _with_row(ma, after(ra, True))
    na, nb = after(ra, True), after(rb, False)
    if nb < na:
        return "#REF!"
    return pre + _with_row(ma, na) + ":" + _with_row(mb, nb)


def shift_formula(formula, own_sheet, target, at, n, cell_row=None):
    from openpyxl.formula.tokenizer import Tokenizer, Token
    if not (isinstance(formula, str) and formula.startswith("=")):
        return formula
    try:
        tok = Tokenizer(formula)
    except Exception:
        return formula
    changed = False
    for t in tok.items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            v = shift_ref(t.value, own_sheet, target, at, n, cell_row)
            if v != t.value:
                t.value = v
                changed = True
    return tok.render() if changed else formula


def _shift_range_str(rng, at, n, extend=False):
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    c1, r1, c2, r2 = range_boundaries(rng)
    if n > 0:
        nr1 = r1 + n if r1 >= at else r1
        nr2 = r2 + n if (r2 >= at or (extend and r2 == at - 1 and r1 < r2)) else r2
    else:
        k, last = -n, at - n - 1
        nr1 = r1 if r1 < at else (r1 - k if r1 > last else at)
        nr2 = r2 if r2 < at else (r2 - k if r2 > last else at - 1)
        if nr2 < nr1:
            return None
    return "%s%d:%s%d" % (get_column_letter(c1), nr1, get_column_letter(c2), nr2)


def _chart_refs(ch):
    """كلُّ مرجعٍ في رسمٍ أصليّ: (كائن، الحقل) — للبيانات والفئات والعناوين."""
    for se in getattr(ch, "series", None) or []:
        for part in ("val", "cat", "xVal", "yVal", "bubbleSize"):
            src = getattr(se, part, None)
            for kind in ("numRef", "strRef", "multiLvlStrRef"):
                ref = getattr(src, kind, None) if src is not None else None
                if ref is not None and getattr(ref, "f", None):
                    yield ref
        tx = getattr(se, "tx", None)
        if tx is not None and getattr(tx, "strRef", None) is not None \
                and tx.strRef.f:
            yield tx.strRef


def _rows_change(wb, ws, at, n, values=None, grow_charts=False):
    """أدرِج (n>0) أو احذف (n<0) صفوفاً، وأزِح كلَّ ما يشير إليها.

    الرسومُ الأصليّة تُزاح مراجعُها كالمعادلات (قِيس: بقي رسمٌ على B2:B3 بعد
    إدراج صفّ). و`grow_charts` (add_rows): صفٌّ أُضيف إلى آخر البيانات يدخل
    في الرسم كما يدخل في المجموع."""
    merged = [str(m) for m in ws.merged_cells.ranges]
    for m in merged:
        ws.unmerge_cells(m)
    height = {r: d.height for r, d in ws.row_dimensions.items() if d.height}
    if n > 0:
        ws.insert_rows(at, n)
    else:
        ws.delete_rows(at, -n)
    # المعادلاتُ في الأوراق كلِّها
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    orig = None               # الامتدادُ لمعادلات الورقة نفسِها
                    if sh is ws:
                        orig = c.row
                        if n > 0 and c.row >= at + n:
                            orig = c.row - n
                        elif n < 0 and c.row >= at:
                            orig = c.row - n
                    c.value = shift_formula(c.value, sh.title, ws.title, at, n,
                                            orig)
    for sh in wb.worksheets:
        for ch in getattr(sh, "_charts", None) or []:
            for ref in _chart_refs(ch):
                ref.f = shift_formula("=" + ref.f, sh.title, ws.title, at, n,
                                      at if grow_charts else None)[1:]
    for m in merged:
        nm = _shift_range_str(m, at, n)
        if nm:
            ws.merge_cells(nm)
    for t in list(getattr(ws, "tables", {}).values()):
        nr = _shift_range_str(t.ref, at, n, extend=True)
        if nr:
            t.ref = nr
            if getattr(t, "autoFilter", None) is not None:
                t.autoFilter.ref = nr
    try:
        for name, dn in list(wb.defined_names.items()):
            txt = dn.attr_text or ""
            new = shift_formula("=" + txt, ws.title, ws.title, at, n)[1:]
            if new != txt:
                dn.attr_text = new
    except Exception:
        pass
    # ارتفاعاتُ الصفوف تتبع صفوفَها
    for r in list(ws.row_dimensions.keys()):
        ws.row_dimensions[r].height = None
    for r, h in height.items():
        if n > 0:
            nr = r + n if r >= at else r
        else:
            if at <= r < at - n:
                continue                  # صفٌّ محذوف
            nr = r + n if r >= at - n else r
        ws.row_dimensions[nr].height = h
    if n > 0:
        above = at - 1
        for r in range(at, at + n):
            if above >= 1:
                for col in range(1, (ws.max_column or 1) + 1):
                    src = ws.cell(row=above, column=col)
                    if src.has_style:
                        ws.cell(row=r, column=col)._style = copy.copy(src._style)
                if above in height:
                    ws.row_dimensions[r].height = height[above]
        for i, vals in enumerate(values or []):
            for j, v in enumerate(vals or []):
                ws.cell(row=at + i, column=1 + j, value=v)


def _totals_row(ws):
    """آخرُ صفٍّ فيه بيانات، وهل هو سطرُ مجموعٍ لما فوقه (معادلةٌ على مدىً
    ينتهي فوقه مباشرةً)."""
    last = 0
    for row in ws.iter_rows():
        if any(c.value not in (None, "") for c in row):
            last = row[0].row
    if not last:
        return 0, False
    for c in ws[last]:
        v = c.value
        if isinstance(v, str) and v.startswith("="):
            for mm in re.finditer(r"\$?[A-Za-z]{1,3}\$?(\d+):\$?[A-Za-z]{1,3}\$?(\d+)",
                                  v):
                if int(mm.group(2)) == last - 1:
                    return last, True
    return last, False


def _ws_of(wb, op):
    name = op.get("sheet")
    if name is None:
        return wb.active
    if name not in wb.sheetnames:
        raise Fail("no sheet named %s (have: %s)" % (name, ", ".join(wb.sheetnames)))
    return wb[name]


def edit_xlsx(path, ops, out, allow_loss=False):
    _paths()
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = load_workbook(path)
    done, failed = [], []
    for n, op in enumerate(ops, 1):
        kind = str(op.get("op", ""))
        try:
            if kind == "set":
                ws = _ws_of(wb, op)
                cells = op.get("cells")
                if cells is None:
                    cells = {op["cell"]: op.get("value")}
                for ref, v in cells.items():
                    ws[ref] = v
                done.append("%s: set %s" % (ws.title, ", ".join(list(cells)[:6])))
            elif kind == "set_range":
                ws = _ws_of(wb, op)
                from openpyxl.utils.cell import coordinate_from_string, \
                    column_index_from_string
                col, row = coordinate_from_string(op["start"])
                c0 = column_index_from_string(col)
                for i, vals in enumerate(op.get("values") or []):
                    for j, v in enumerate(vals or []):
                        ws.cell(row=row + i, column=c0 + j, value=v)
                done.append("%s: range from %s" % (ws.title, op["start"]))
            elif kind == "insert_rows":
                ws = _ws_of(wb, op)
                vals = op.get("values") or []
                cnt = int(op.get("count") or len(vals) or 1)
                _rows_change(wb, ws, int(op["at"]), cnt, vals)
                done.append("%s: %d row(s) inserted at %s" % (ws.title, cnt, op["at"]))
            elif kind == "add_rows":
                ws = _ws_of(wb, op)
                vals = op.get("values") or []
                if not vals:
                    raise Fail("add_rows needs values")
                last, is_tot = _totals_row(ws)
                at = last if is_tot else last + 1
                _rows_change(wb, ws, at, len(vals), vals, grow_charts=True)
                done.append("%s: %d row(s) added at %d%s" % (
                    ws.title, len(vals), at, " (before totals)" if is_tot else ""))
            elif kind == "delete_rows":
                ws = _ws_of(wb, op)
                cnt = int(op.get("count") or 1)
                _rows_change(wb, ws, int(op["at"]), -cnt)
                done.append("%s: %d row(s) deleted at %s" % (ws.title, cnt, op["at"]))
            elif kind == "add_sheet":
                from build_xlsx import build_xlsx
                fd, tmp = tempfile.mkstemp(prefix="weaver-sheet-", suffix=".xlsx")
                os.close(fd)
                try:
                    lang = "ar" if _is_ar(" ".join(str(h) for h in op.get(
                        "headers") or [])) else "en"
                    build_xlsx(op.get("rows") or [], tmp, headers=op.get("headers"),
                               lang=lang, with_totals=bool(op.get("totals")))
                    src = load_workbook(tmp).active
                    dst = wb.create_sheet(str(op.get("name") or "Sheet")[:31])
                    _copy_sheet(src, dst)
                finally:
                    os.remove(tmp)
                done.append("sheet %s added" % dst.title)
            elif kind == "rename_sheet":
                ws = _ws_of(wb, op)
                old = ws.title
                ws.title = str(op["to"])[:31]
                done.append("sheet %s → %s" % (old, ws.title))
            elif kind == "delete_sheet":
                ws = _ws_of(wb, op)
                if len(wb.worksheets) == 1:
                    raise Fail("cannot delete the only sheet")
                wb.remove(ws)
                done.append("sheet %s deleted" % op.get("sheet"))
            elif kind == "format":
                ws = _ws_of(wb, op)
                sel = ws[op["range"]]
                if not isinstance(sel, tuple):
                    sel = ((sel,),)
                for row in sel:
                    for c in (row if isinstance(row, tuple) else (row,)):
                        if op.get("bold") is not None or op.get("color"):
                            f = copy.copy(c.font)
                            c.font = Font(name=f.name, size=f.size,
                                          bold=op.get("bold", f.bold),
                                          italic=f.italic,
                                          color=op.get("color") or f.color)
                        if op.get("fill"):
                            c.fill = PatternFill("solid", fgColor=op["fill"])
                        if op.get("number_format"):
                            c.number_format = op["number_format"]
                        if op.get("align"):
                            c.alignment = Alignment(horizontal=op["align"],
                                                    vertical="center")
                done.append("%s: formatted %s" % (ws.title, op["range"]))
            elif kind == "add_chart":
                ws = _ws_of(wb, op)
                _xl_chart(ws, op)
                done.append("%s: %s chart added" % (ws.title, op.get("type", "bar")))
            elif kind == "rtl":
                ws = _ws_of(wb, op)
                ws.sheet_view.rightToLeft = bool(op.get("value", True))
                done.append("%s: %s" % (ws.title, "RTL" if op.get("value", True)
                                        else "LTR"))
            else:
                raise Fail("unknown Excel op: %s" % kind)
        except Fail as e:
            failed.append("op %d %s: %s" % (n, kind, e))
        except Exception as e:
            failed.append("op %d %s: %s: %s" % (n, kind, type(e).__name__,
                                                str(e)[:160]))
    if not done:
        return {"done": done, "failed": failed, "out": None}
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    _no_dupes(out)
    load_workbook(out)
    # ما نقص من الملفّ بالحفظ — مقيسٌ بالمقارنة. ورقةٌ حُذفت بطلبٍ تُنقص
    # أجزاءَها عمداً، فلا تُحسب.
    lost = [] if any(o.get("op") == "delete_sheet" for o in ops) \
        else xlsx_loss(path, out)
    if lost and not allow_loss:
        try:
            os.remove(out)
        except OSError:
            pass
        raise Refused("saving would DROP %s from %s — nothing written. Tell the "
                      "user; do not work around it." % (", ".join(lost),
                                                        os.path.basename(path)))
    return {"done": done, "failed": failed, "out": out}


class Refused(Fail):
    pass


def cmd_edit(path, ops_path, out=None, allow_loss=False):
    _need(path)
    ops = _load_json(ops_path)
    if isinstance(ops, dict):
        ops = ops.get("ops") or [ops]
    if not isinstance(ops, list) or not ops:
        raise Fail("OPS must be a JSON list of operations")
    out = out or _default_out(path)
    e = _ext(path)
    if _ext(out) != e:
        raise Fail("--out must keep the extension %s" % e)
    if e == ".docx":
        return edit_docx(path, ops, out)
    if e == ".pptx":
        return edit_pptx(path, ops, out)
    if e == ".xlsx":
        return edit_xlsx(path, ops, out, allow_loss)
    raise Fail("edit supports .docx .pptx .xlsx — got %s" % (e or "no extension"))


# ═════════════════════════════════ CLI ══════════════════════════════════
USAGE = __doc__.split("قواعد:")[0].strip().splitlines()[-5:]


def _opt(args, name, default=None):
    if name in args:
        i = args.index(name)
        if i + 1 < len(args):
            v = args[i + 1]
            del args[i:i + 2]
            return v
        del args[i]
    return default


def main(argv=None):
    args = list(sys.argv[1:] if argv is None else argv)
    if not args or args[0] in ("-h", "--help", "help"):
        print("\n".join(USAGE))
        return 0
    cmd = args.pop(0)
    try:
        if cmd == "info":
            find = _opt(args, "--find")
            rows = int(_opt(args, "--rows", 40))
            if not args:
                raise Fail("info needs a file")
            print(cmd_info(args[0], find, rows))
            return 0
        if cmd == "build":
            out = _opt(args, "--out")
            if not args or not out:
                raise Fail("usage: build SPEC.json --out FILE.docx|.pptx|.xlsx")
            print("✓ " + cmd_build(args[0], out))
            print("✓ saved: %s" % out)
            return 0
        if cmd == "chart":
            out = _opt(args, "--out")
            into = _opt(args, "--into")
            if not args:
                raise Fail("usage: chart SPEC.json --out FILE.png | --into FILE")
            print("✓ " + cmd_chart(args[0], out, into))
            return 0
        if cmd == "edit":
            out = _opt(args, "--out")
            allow = "--allow-loss" in args
            args = [a for a in args if a != "--allow-loss"]
            if len(args) < 2:
                raise Fail("usage: edit FILE OPS.json [--out NEW]")
            res = cmd_edit(args[0], args[1], out, allow)
            for d in res["done"]:
                print("✓ " + d)
            for f in res["failed"]:
                print("✗ " + f)
            if res["out"]:
                print("✓ saved: %s" % res["out"])
            else:
                print("✗ nothing applied — no file written")
            return 0 if not res["failed"] else (2 if res["done"] else 1)
        raise Fail("unknown command %s — use info, build, chart, edit" % cmd)
    except Refused as e:
        print("✗ REFUSED: %s" % e)
        return 3
    except Fail as e:
        print("✗ %s" % e)
        return 1
    except Exception as e:
        print("✗ %s: %s" % (type(e).__name__, str(e)[:300]))
        return 1


if __name__ == "__main__":
    sys.exit(main())
