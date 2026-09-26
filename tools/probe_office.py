"""
tools/probe_office.py — الخطوة ١ لملفّات أوفيس: قياسٌ لا بناء
===============================================================
فحصٌ للقراءة فقط: لا يكتب في الإعداد، ولا يثبّت شيئاً، ولا ينادي النموذج.
كلُّ ما يُبنى يُبنى في مجلّدٍ مؤقّتٍ ويُحذف في الآخر (إلا بـ --keep).
يجيب بالدليل قبل أن نربط شيئاً بالمحرّك:

  ١ أيّ المكتبات والبرامج على هاتفك — وهل يراها `python3` الذي يناديه المحرّك؟
  ٢ أتعمل أدواتُ Weaver الموجودة كما هي؟ يبني بها Word وPowerPoint وExcel
     ورسماً بيانياً، عربيّاً وإنجليزيّاً، ثمّ **يفتح كلَّ ملفٍّ ويقرؤه** ليتحقّق
     (الاتّجاه RTL، الجداول، المعادلات، الصور) — لا يكفي أنّ الملفَّ وُجد.
  ٣ أيمكن **تعديلُ** ملفٍّ موجود؟ يعدّل نصّاً وخليّةً وشريحة، ويفكّ الملفَّ
     ويعيد حزمه (ooxml)، ثمّ يقرأ النتيجة.
  ٤ طريقُ المحرّك: أتصل الملفّاتُ المبنيّة بطاقةً؟ وأيُحفظ الملفُّ المرفوع في
     مجلّد العمل ليعدّله النموذج — أم يصله نصّاً فقط؟

وما لم يُقَس يُقال «لم يُقَس» — لا يُفترض.

التشغيل:
    cd ~/weaver-write && python3 tools/probe_office.py
    cd ~/weaver-write && python3 tools/probe_office.py --keep ~/storage/downloads/office-probe
"""
import base64
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

_SK = os.path.join(_ROOT, "capabilities", "skills")
_SCRIPTS = {
    "docx": os.path.join(_SK, "docx_builder", "scripts"),
    "pptx": os.path.join(_SK, "pptx_builder", "scripts"),
    "xlsx": os.path.join(_SK, "xlsx_builder", "scripts"),
    "chart": os.path.join(_SK, "chart_builder", "scripts"),
}
_OOXML = os.path.join(_SCRIPTS["docx"], "ooxml")
for _p in _SCRIPTS.values():
    if _p not in sys.path:
        sys.path.append(_p)

results = []          # (القسم، البند، ✓/✗/–، التفصيل)


def say(s=""):
    print(s, flush=True)


def mark(section, name, state, detail=""):
    """state: True ✓ · False ✗ · None «لم يُقَس»."""
    results.append((section, name, state, detail))
    sym = "✓" if state is True else ("✗" if state is False else "–")
    say("  %s %s%s" % (sym, name, ("   — " + detail) if detail else ""))


def _import(name):
    try:
        m = __import__(name)
        return m, str(getattr(m, "__version__", "") or "✓")
    except BaseException as e:
        return None, "%s: %s" % (type(e).__name__, str(e)[:80])


def _timed(fn):
    t0 = time.time()
    try:
        return fn(), None, time.time() - t0
    except BaseException as e:
        return None, "%s: %s" % (type(e).__name__, str(e)[:160]), time.time() - t0


def _zip_text(path, member):
    with zipfile.ZipFile(path) as z:
        return z.read(member).decode("utf-8", "ignore")


def _zip_names(path):
    with zipfile.ZipFile(path) as z:
        return z.namelist()


AR_TITLE = "أثر التعلّم الرقميّ في التحصيل"
EN_TITLE = "Digital Learning and Achievement"

# ── ١ المكتبات والبرامج ─────────────────────────────────────────────────
say("\n══ ١) المكتبات والبرامج ══")
say("  بايثون هذه الأداة: %s (%s)" % (sys.executable, sys.version.split()[0]))
LIBS = (("python-docx (Word)", "docx", True),
        ("python-pptx (PowerPoint)", "pptx", True),
        ("openpyxl (Excel)", "openpyxl", True),
        ("matplotlib (رسوم)", "matplotlib", True),
        ("numpy", "numpy", True),
        ("arabic-reshaper (عربيّ الرسوم)", "arabic_reshaper", True),
        ("python-bidi (عربيّ الرسوم)", "bidi", True),
        ("lxml", "lxml", True),
        ("Pillow (صور)", "PIL", False),
        ("defusedxml (فكّ ooxml)", "defusedxml", True),
        ("formulas (حسابُ المعادلات)", "formulas", False))
have = {}
for label, mod, needed in LIBS:
    m, info = _import(mod)
    have[mod] = bool(m)
    mark("libs", label, True if m else (False if needed else None),
         info if m else ("غيرُ مثبَّت" + ("" if needed else " (اختياريّ)")))

for b, why in (("soffice", "تحويل PDF وإعادة حساب Excel — اختياريّ"),
               ("pdftoppm", "صورٌ مصغّرة للشرائح — اختياريّ")):
    p = shutil.which(b) or (shutil.which("libreoffice") if b == "soffice" else None)
    mark("libs", b + " (برنامج)", True if p else None,
         p or ("غيرُ موجود — " + why))

# المحرّكُ ينادي `python3` من PATH (كما في مهارة cite-pages) — لا بايثونَ هذه الأداة.
_py = shutil.which("python3")
if not _py:
    mark("libs", "python3 الذي يناديه المحرّك", False, "غيرُ موجود في PATH")
else:
    same = os.path.realpath(_py) == os.path.realpath(sys.executable)
    try:
        chk = subprocess.run(
            [_py, "-c", "import importlib.util as u,json;print(json.dumps({m:bool("
             "u.find_spec(m)) for m in ['docx','pptx','openpyxl','matplotlib',"
             "'arabic_reshaper','bidi','defusedxml']}))"],
            capture_output=True, text=True, timeout=60)
        seen = json.loads(chk.stdout.strip().splitlines()[-1])
        miss = [k for k, v in seen.items() if not v]
        mark("libs", "python3 الذي يناديه المحرّك يرى المكتبات", not miss,
             "%s%s" % (_py, "" if same else " (غيرُ بايثون هذه الأداة)")
             + (" · ينقصه: " + ", ".join(miss) if miss else ""))
    except BaseException as e:
        mark("libs", "python3 الذي يناديه المحرّك", None, "لم يُقَس: %s" % e)

try:
    sys.path.insert(0, os.path.join(_ROOT, "engines", "fonts-core"))
    from fonts import resolve_arabic_font   # noqa: E402
    _f = resolve_arabic_font("Kufyan Arabic Black")
    mark("libs", "خطٌّ عربيٌّ مضمَّن للرسوم", bool(_f), str(_f)[:90])
except BaseException as e:
    mark("libs", "خطٌّ عربيٌّ مضمَّن للرسوم", False, str(e)[:90])

# ── ٢ البناءُ بأدوات Weaver الموجودة ─────────────────────────────────────
keep = None
if "--keep" in sys.argv:
    i = sys.argv.index("--keep")
    keep = os.path.expanduser(sys.argv[i + 1]) if i + 1 < len(sys.argv) else None
OUT = tempfile.mkdtemp(prefix="weaver-office-probe-")
built = {}


def _p(name):
    return os.path.join(OUT, name)


say("\n══ ٢) البناءُ بأدوات Weaver الموجودة — ثمّ فتحُ كلِّ ملفٍّ والتحقّق ══")

# Word — بسيط
if have.get("docx"):
    def _docx_simple():
        from build_docx import build_academic_docx
        out = build_academic_docx(
            AR_TITLE, [{"heading": "المقدّمة", "body": "نصٌّ عربيٌّ للتجربة."}],
            references=["المرجع الأوّل"], output_path=_p("word-ar.docx"), lang="ar")
        from docx import Document
        d = Document(out)
        text = "\n".join(p.text for p in d.paragraphs)
        xml = _zip_text(out, "word/document.xml")
        assert AR_TITLE in text, "العنوان غيرُ موجود"
        assert "<w:bidi" in xml, "بلا اتّجاه RTL"
        assert "المراجع" in text, "بلا عنوان المراجع"
        return out
    r, err, dt = _timed(_docx_simple)
    built["word-ar.docx"] = r
    mark("build", "Word عربيّ (build_docx) ⟵ عنوان · RTL · مراجع", r is not None,
         err or "%.1fث" % dt)

    def _docx_rich(lang):
        from docx_advanced import build_rich_docx
        name = "word-rich-%s.docx" % lang
        ar = lang == "ar"
        build_rich_docx(
            AR_TITLE if ar else EN_TITLE,
            [{"heading": "النتائج" if ar else "Results",
              "body": ("## أرقام\n- نقطة أولى\n- نقطة ثانية" if ar
                       else "## Numbers\n- first point\n- second point"),
              "table": {"headers": ["البند", "القيمة"] if ar else ["Item", "Value"],
                        "rows": [["أ", 10], ["ب", 20]] if ar else [["A", 10], ["B", 20]],
                        "totals": ["المجموع", 30] if ar else ["Total", 30]}}],
            output_path=_p(name), lang=lang, toc=True,
            header_text="Weaver Write", page_numbers=True,
            references=["Ref 1"])
        from docx import Document
        d = Document(_p(name))
        assert len(d.tables) >= 1, "بلا جدول"
        cells = [c.text for c in d.tables[0].rows[0].cells]
        assert (("البند" in cells) if ar else ("Item" in cells)), cells
        xml = _zip_text(_p(name), "word/document.xml")
        assert "TOC" in xml, "بلا فهرس"
        if ar:
            assert "<w:bidi" in xml or "w:bidiVisual" in xml, "بلا RTL"
        return _p(name)
    for lg in ("ar", "en"):
        r, err, dt = _timed(lambda lg=lg: _docx_rich(lg))
        built["word-rich-%s.docx" % lg] = r
        mark("build", "Word احترافيّ %s (docx_advanced) ⟵ جدول · فهرس · ترويسة"
             % ("عربيّ" if lg == "ar" else "إنجليزيّ"), r is not None,
             err or "%.1fث" % dt)
else:
    mark("build", "Word", None, "لم يُقَس — python-docx غيرُ مثبَّت")

# PowerPoint
if have.get("pptx"):
    def _deck(lang):
        from build_pptx import build_deck
        ar = lang == "ar"
        name = "deck-%s.pptx" % lang
        slides = [{"title": "المحاور" if ar else "Outline",
                   "points": ["النقطة الأولى", "النقطة الثانية"] if ar
                   else ["First", "Second"]},
                  {"layout": "section", "title": "القسم" if ar else "Section"}]
        build_deck(AR_TITLE if ar else EN_TITLE, slides, subtitle="Weaver",
                   output_path=_p(name), lang=lang)
        from pptx import Presentation
        prs = Presentation(_p(name))
        n = len(prs.slides)
        assert n == len(slides) + 2, "عددُ الشرائح %d" % n
        allt = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                        if sh.has_text_frame)
        assert (AR_TITLE if ar else EN_TITLE) in allt, "بلا العنوان"
        if ar:
            assert 'rtl="1"' in _zip_text(_p(name), "ppt/slides/slide2.xml"), "بلا RTL"
        return _p(name)
    for lg in ("ar", "en"):
        r, err, dt = _timed(lambda lg=lg: _deck(lg))
        built["deck-%s.pptx" % lg] = r
        mark("build", "PowerPoint %s (build_pptx) ⟵ شرائح · عنوان%s"
             % ("عربيّ" if lg == "ar" else "إنجليزيّ", " · RTL" if lg == "ar" else ""),
             r is not None, err or "%.1fث" % dt)

    def _ptable():
        from pptx_table import add_table_slide
        src = built.get("deck-ar.pptx")
        assert src, "لا عرضَ لإضافة الجدول إليه"
        out = _p("deck-ar-table.pptx")
        res = add_table_slide(src, ["البند", "القيمة"], [["أ", 1], ["ب", 2]],
                              lang="ar", title="جدول", output_path=out)
        assert not isinstance(res, dict) or res.get("ok", True), res
        from pptx import Presentation
        prs = Presentation(out)
        assert any(sh.has_table for sh in prs.slides[-1].shapes), "بلا جدول"
        return out
    r, err, dt = _timed(_ptable)
    built["deck-ar-table.pptx"] = r
    mark("build", "PowerPoint ⟵ شريحةُ جدولٍ أصليّ (pptx_table)", r is not None,
         err or "%.1fث" % dt)
else:
    mark("build", "PowerPoint", None, "لم يُقَس — python-pptx غيرُ مثبَّت")

# Excel
if have.get("openpyxl"):
    def _xlsx(lang):
        from build_xlsx import build_xlsx
        ar = lang == "ar"
        name = "sheet-%s.xlsx" % lang
        build_xlsx([["أ" if ar else "A", 10, 5], ["ب" if ar else "B", 20, 7]],
                   _p(name), headers=["البند", "الكمّيّة", "السعر"] if ar
                   else ["Item", "Qty", "Price"], lang=lang, with_totals=True)
        from openpyxl import load_workbook
        ws = load_workbook(_p(name)).active
        assert bool(ws.sheet_view.rightToLeft) == ar, "الاتّجاه خطأ"
        assert str(ws["B4"].value).startswith("=SUM("), "بلا معادلة المجموع: %r" % ws["B4"].value
        assert ws["A1"].font.bold, "الترويسة بلا تنسيق"
        return _p(name)
    for lg in ("ar", "en"):
        r, err, dt = _timed(lambda lg=lg: _xlsx(lg))
        built["sheet-%s.xlsx" % lg] = r
        mark("build", "Excel %s (build_xlsx) ⟵ ترويسة · SUM · %s"
             % ("عربيّ" if lg == "ar" else "إنجليزيّ", "RTL" if lg == "ar" else "LTR"),
             r is not None, err or "%.1fث" % dt)
else:
    mark("build", "Excel", None, "لم يُقَس — openpyxl غيرُ مثبَّت")

# الرسوم
if have.get("matplotlib"):
    def _chart(kind, lang):
        from build_chart import build_chart
        ar = lang == "ar"
        name = "chart-%s-%s.png" % (kind, lang)
        data = ({"labels": ["الأوّل", "الثاني", "الثالث"], "values": [3, 5, 2]} if ar
                else {"labels": ["One", "Two", "Three"], "values": [3, 5, 2]})
        res = build_chart(kind, data, _p(name), title="عنوان" if ar else "Title",
                          lang=lang)
        assert res.get("ok"), res.get("error")
        with open(_p(name), "rb") as f:
            assert f.read(8) == b"\x89PNG\r\n\x1a\n", "ليست صورةَ PNG"
        return _p(name)
    for kind, lg in (("bar", "ar"), ("pie", "ar"), ("line", "en")):
        r, err, dt = _timed(lambda k=kind, lg=lg: _chart(k, lg))
        built["chart-%s-%s.png" % (kind, lg)] = r
        mark("build", "رسم %s %s (build_chart)" % (kind, "عربيّ" if lg == "ar"
                                                   else "إنجليزيّ"),
             r is not None, err or "%.1fث" % dt)
    def _ar_order():
        """أيظهر العربيُّ في الرسم بترتيبه الصحيح؟ يُرسم «ا ب» بطريق build_chart
        نفسِه: الألفُ (الأضيق) يجب أن تكون يميناً. قِيس: matplotlib 3.11
        يشكّل العربيَّ بنفسه، فتشكيلُه مرّةً ثانيةً يقلبه."""
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import numpy as np
        import build_chart as bc
        fam = None
        try:
            from fonts import register_for_matplotlib
            fam = register_for_matplotlib("Kufyan Arabic Black")
        except Exception:
            pass
        fig = plt.figure(figsize=(3, 1), dpi=100)
        kw = {"fontfamily": fam} if fam else {}
        fig.text(0.1, 0.3, bc._reshape_ar(["ا ب"])[0], fontsize=40, **kw)
        fig.canvas.draw()
        a = np.asarray(fig.canvas.buffer_rgba())[:, :, :3].min(axis=2) < 128
        plt.close(fig)
        cols = a.any(axis=0)
        blobs, start = [], None
        for x, v in enumerate(list(cols) + [False]):
            if v and start is None:
                start = x
            elif not v and start is not None:
                blobs.append((start, x - start))
                start = None
        assert len(blobs) == 2, "كتلُ الحبر %d لا ٢" % len(blobs)
        (x1, w1), (x2, w2) = blobs
        assert w2 < w1, "الألفُ يسارَ الباء ⟵ النصُّ مقلوب"
        return "طبقةُ matplotlib: %s" % ("يشكّل بنفسه (libraqm) ⟵ يُمرَّر كما هو"
                                         if bc._native_shaping()
                                         else "arabic-reshaper + bidi")
    r, err, dt = _timed(_ar_order)
    mark("build", "العربيُّ في الرسم بترتيبه الصحيح (لا مقلوب)", r is not None,
         err or r)
    if not (have.get("arabic_reshaper") and have.get("bidi")):
        mark("build", "حروفُ الرسم العربيّ متّصلة", False,
             "arabic-reshaper/python-bidi غيرُ مثبَّتين ⟵ حروفٌ مقطّعة")

    def _embed(kind):
        from embed_chart import embed_chart_in_docx, add_chart_slide
        spec = {"type": "bar", "data": {"labels": ["أ", "ب"], "values": [1, 2]},
                "lang": "ar"}
        if kind == "docx":
            src = built.get("word-ar.docx")
            assert src, "لا مستندَ للإدراج فيه"
            out = _p("word-ar-chart.docx")
            res = embed_chart_in_docx(src, spec, caption="شكل ١", output_path=out)
            assert res.get("ok", True), res
            assert any(n.startswith("word/media/") for n in _zip_names(out)), "بلا صورة"
        else:
            src = built.get("deck-ar.pptx")
            assert src, "لا عرضَ للإدراج فيه"
            out = _p("deck-ar-chart.pptx")
            res = add_chart_slide(src, spec, title="رسم", output_path=out)
            assert res.get("ok", True), res
            assert any(n.startswith("ppt/media/") for n in _zip_names(out)), "بلا صورة"
        return out
    for kind in ("docx", "pptx"):
        r, err, dt = _timed(lambda k=kind: _embed(k))
        built["embed-" + kind] = r
        mark("build", "إدراجُ رسمٍ في %s (embed_chart)" %
             ("Word" if kind == "docx" else "PowerPoint"), r is not None,
             err or "%.1fث" % dt)
else:
    mark("build", "الرسوم", None, "لم يُقَس — matplotlib غيرُ مثبَّت")

# ── ٣ تعديلُ ملفٍّ موجود ─────────────────────────────────────────────────
say("\n══ ٣) تعديلُ ملفٍّ موجود — ثمّ قراءةُ النتيجة ══")


def _edit_docx():
    src = built.get("word-ar.docx")
    assert src, "لم يُبنَ مستندٌ لتعديله"
    from docx import Document
    d = Document(src)
    tgt = next(p for p in d.paragraphs if "نصٌّ عربيٌّ" in p.text)
    tgt.runs[0].text = "نصٌّ مُعدَّل."
    for r_ in tgt.runs[1:]:
        r_.text = ""
    d.add_paragraph("فقرةٌ مضافة.")
    out = _p("word-ar-edited.docx")
    d.save(out)
    d2 = Document(out)
    t = "\n".join(p.text for p in d2.paragraphs)
    assert "نصٌّ مُعدَّل." in t and "نصٌّ عربيٌّ للتجربة" not in t, "لم يُعدَّل"
    assert "فقرةٌ مضافة." in t, "لم تُضَف"
    assert "<w:bidi" in _zip_text(out, "word/document.xml"), "ضاع RTL"
    return out


def _edit_xlsx():
    src = built.get("sheet-ar.xlsx")
    assert src, "لم يُبنَ مصنَّفٌ لتعديله"
    from openpyxl import load_workbook
    wb = load_workbook(src)
    ws = wb.active
    ws["B2"] = 99
    ws.insert_rows(3)
    ws["A3"], ws["B3"], ws["C3"] = "ج", 1, 1
    out = _p("sheet-ar-edited.xlsx")
    wb.save(out)
    ws2 = load_workbook(out).active
    assert ws2["B2"].value == 99, "الخليّة لم تُعدَّل"
    assert ws2["A3"].value == "ج", "الصفّ لم يُضَف"
    f = str(ws2["B5"].value)
    ok = f.startswith("=SUM(") and "B4" in f
    return out, ok, f


def _edit_pptx():
    src = built.get("deck-ar.pptx")
    assert src, "لم يُبنَ عرضٌ لتعديله"
    from pptx import Presentation
    prs = Presentation(src)
    done = False
    for sh in prs.slides[1].shapes:
        if sh.has_text_frame and "النقطة الأولى" in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                for r_ in p.runs:
                    r_.text = r_.text.replace("النقطة الأولى", "نقطةٌ مُعدَّلة")
            done = True
    assert done, "لم يوجد النصّ"
    n0 = len(prs.slides)
    # حذفُ الشريحة الأخيرة (python-pptx بلا دالّةٍ علنيّة لذلك)
    sld = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)
    out = _p("deck-ar-edited.pptx")
    prs.save(out)
    p2 = Presentation(out)
    allt = " ".join(sh.text_frame.text for s in p2.slides for sh in s.shapes
                    if sh.has_text_frame)
    assert "نقطةٌ مُعدَّلة" in allt, "لم يُعدَّل"
    assert len(p2.slides) == n0 - 1, "لم تُحذف الشريحة"
    return out


def _ooxml_roundtrip():
    src = built.get("word-rich-ar.docx")
    assert src, "لم يُبنَ مستندٌ لفكّه"
    d = os.path.join(OUT, "unpacked")
    r1 = subprocess.run([sys.executable, os.path.join(_OOXML, "unpack.py"), src, d],
                        capture_output=True, text=True, timeout=120)
    assert r1.returncode == 0, (r1.stderr or r1.stdout)[-200:]
    doc_xml = os.path.join(d, "word", "document.xml")
    x = open(doc_xml, encoding="utf-8").read()
    import re
    m = re.search(r"<w:t(?:\s[^>]*)?>", x)
    assert m, "XML بلا نصّ"
    # تعديلٌ على مستوى XML: وسمٌ يُضاف إلى أوّل نصّ
    x2 = x[:m.end()] + "[EDITED] " + x[m.end():]
    open(doc_xml, "w", encoding="utf-8").write(x2)
    out = _p("word-rich-ar-ooxml.docx")
    r2 = subprocess.run([sys.executable, os.path.join(_OOXML, "pack.py"), d, out,
                         "--force"], capture_output=True, text=True, timeout=120)
    assert r2.returncode == 0, (r2.stderr or r2.stdout)[-200:]
    from docx import Document
    t = "\n".join(p.text for p in Document(out).paragraphs)
    t += "\n".join(c.text for tb in Document(out).tables for rw in tb.rows
                   for c in rw.cells)
    assert "[EDITED]" in t, "التعديلُ لم يظهر"
    return out


if have.get("docx"):
    r, err, dt = _timed(_edit_docx)
    built["word-ar-edited.docx"] = r
    mark("edit", "Word ⟵ تغييرُ فقرةٍ وإضافةُ أخرى (الاتّجاه باقٍ)", r is not None,
         err or "%.1fث" % dt)
    if have.get("defusedxml"):
        r, err, dt = _timed(_ooxml_roundtrip)
        built["word-rich-ar-ooxml.docx"] = r
        mark("edit", "Word ⟵ فكٌّ وتعديلُ XML وإعادةُ حزم (ooxml)", r is not None,
             err or "%.1fث" % dt)
    else:
        mark("edit", "ooxml", None, "لم يُقَس — defusedxml غيرُ مثبَّت")
else:
    mark("edit", "Word", None, "لم يُقَس — python-docx غيرُ مثبَّت")

if have.get("openpyxl"):
    r, err, dt = _timed(_edit_xlsx)
    if r is not None:
        out, fok, f = r
        built["sheet-ar-edited.xlsx"] = out
        mark("edit", "Excel ⟵ تغييرُ خليّةٍ وإدراجُ صفّ", True, "%.1fث" % dt)
        # openpyxl وحده لا يُزيح المراجع — معلومةٌ لا حكم؛ والحكمُ على أداة
        # المحرّك (pipeline/office.py) في القسم التالي.
        mark("edit", "  ⟵ openpyxl وحده: معادلةُ المجموع بعد الإدراج", None,
             "%s%s" % (f, "" if fok else " — لا يُزيحها (office.py يُزيحها)"))
    else:
        mark("edit", "Excel ⟵ تغييرُ خليّةٍ وإدراجُ صفّ", False, err)
else:
    mark("edit", "Excel", None, "لم يُقَس — openpyxl غيرُ مثبَّت")

if have.get("pptx"):
    r, err, dt = _timed(_edit_pptx)
    built["deck-ar-edited.pptx"] = r
    mark("edit", "PowerPoint ⟵ تغييرُ نصٍّ وحذفُ شريحة", r is not None,
         err or "%.1fث" % dt)
else:
    mark("edit", "PowerPoint", None, "لم يُقَس — python-pptx غيرُ مثبَّت")

# ── ٣ب أداةُ المحرّك: pipeline/office.py بـpython3 الذي يناديه المحرّك ──────
say("\n══ ٣ب) أداةُ المحرّك (pipeline/office.py) — كما تناديها المهارات ══")
_OFF = os.path.join(_ROOT, "pipeline", "office.py")
_PY = shutil.which("python3") or sys.executable


def _office(*args):
    r = subprocess.run([_PY, _OFF] + [str(a) for a in args], capture_output=True,
                       text=True, timeout=300)
    return r.returncode, (r.stdout or "") + (r.stderr or "")


def _spec(name, obj):
    p = os.path.join(OUT, name)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False)
    return p


def _tool_word():
    sp = _spec("t-word.json", {"title": AR_TITLE, "sections": [
        {"heading": "المقدّمة", "body": "نصٌّ أوّل."},
        {"heading": "النتائج", "body": "جدول:",
         "table": {"headers": ["البند", "القيمة"], "rows": [["أ", 1]]}}]})
    out = _p("tool-word.docx")
    c, o = _office("build", sp, "--out", out)
    assert c == 0 and "✓ saved" in o, o[-200:]
    c, o = _office("info", out)
    assert c == 0 and "¶" in o and "T0" in o, o[-200:]
    ops = _spec("t-word-ops.json", [
        {"op": "replace", "find": "نصٌّ أوّل", "with": "نصٌّ أوّل مُعدَّل"},
        {"op": "set_cell", "table": 0, "row": 1, "col": 1, "text": "9"}])
    c, o = _office("edit", out, ops)
    assert c == 0, o[-200:]
    from docx import Document
    d = Document(_p("tool-word-edited.docx"))
    t = "\n".join(p.text for p in d.paragraphs)
    assert t.count("نصٌّ أوّل مُعدَّل") == 1, "الاستبدال: %d" % t.count("نصٌّ أوّل مُعدَّل")
    assert d.tables[0].cell(1, 1).text == "9", "الخليّة لم تُعدَّل"
    return "بناء · قراءة · تعديل"


def _tool_pptx():
    sp = _spec("t-deck.json", {"title": AR_TITLE, "slides": [
        {"title": "أ", "points": ["١"]}, {"title": "ب", "points": ["٢"]},
        {"title": "جدول", "table": {"headers": ["س"], "rows": [["١"]]}}]})
    out = _p("tool-deck.pptx")
    c, o = _office("build", sp, "--out", out)
    assert c == 0, o[-200:]
    # حذفٌ ثمّ إضافة — قِيس هنا عطبٌ: slide7.xml مرّتين ⟵ ملفٌّ فاسد
    ops = _spec("t-deck-ops.json", [
        {"op": "delete_slide", "slide": 2},
        {"op": "add_slide", "after": 3, "title": "جديدة", "points": ["نقطة"]}])
    c, o = _office("edit", out, ops)
    assert c == 0, o[-200:]
    ed = _p("tool-deck-edited.pptx")
    names = zipfile.ZipFile(ed).namelist()
    assert len(names) == len(set(names)), "أسماءٌ مكرّرة في الملفّ"
    from pptx import Presentation
    n = len(Presentation(ed).slides)
    assert n == 5, "عددُ الشرائح %d لا ٥" % n
    return "بناء · حذفٌ ثمّ إضافة · بلا تكرار"


def _tool_xlsx():
    sp = _spec("t-sheet.json", {"sheets": [{"name": "المصاريف",
        "headers": ["البند", "المبلغ"], "rows": [["إيجار", 1500], ["طعام", 800]],
        "totals": True, "chart": {"type": "bar", "data": "B1:B3",
                                  "categories": "A2:A3"}}]})
    out = _p("tool-sheet.xlsx")
    c, o = _office("build", sp, "--out", out)
    assert c == 0, o[-200:]
    ops = _spec("t-sheet-ops.json", [
        {"op": "add_rows", "values": [["مواصلات", 300]]}])
    c, o = _office("edit", out, ops)
    assert c == 0, o[-200:]
    ed = _p("tool-sheet-edited.xlsx")
    from openpyxl import load_workbook
    ws = load_workbook(ed).active
    f = str(ws["B5"].value)
    assert f == "=SUM(B2:B4)", "المجموع: %s" % f
    x = zipfile.ZipFile(ed).read("xl/charts/chart1.xml").decode("utf-8", "ignore")
    assert "$B$2:$B$4" in x, "الرسمُ لم يتّسع"
    return "صفٌّ قبل الإجمالي ⟵ %s والرسمُ يتّسع" % f


def _tool_chart():
    sp = _spec("t-chart.json", {"type": "bar", "title": "المبيعات",
                                "data": {"labels": ["يناير", "فبراير"], "values": [1, 2]}})
    c, o = _office("chart", sp, "--into", _p("tool-word.docx"))
    assert c == 0, o[-200:]
    names = zipfile.ZipFile(_p("tool-word-edited.docx")).namelist()
    assert sum(1 for n in names if n.startswith("word/media/")) >= 1, "بلا صورة"
    return "رسمٌ داخل Word"


if not os.path.isfile(_OFF):
    mark("tool", "pipeline/office.py", False, "غيرُ موجود")
else:
    for label, fn, need in (("Word عبر الأداة", _tool_word, "docx"),
                            ("PowerPoint عبر الأداة", _tool_pptx, "pptx"),
                            ("Excel عبر الأداة", _tool_xlsx, "openpyxl"),
                            ("رسمٌ عبر الأداة", _tool_chart, "matplotlib")):
        if not have.get(need):
            mark("tool", label, None, "لم يُقَس — %s غيرُ مثبَّت" % need)
            continue
        r, err, dt = _timed(fn)
        mark("tool", label, r is not None, err or "%s — %.1fث" % (r, dt))

# ── ٤ طريقُ المحرّك ─────────────────────────────────────────────────────
say("\n══ ٤) طريقُ المحرّك (ما يرسله الخادم ويستقبله) ══")
S = None
try:
    sys.path.insert(0, os.path.join(_ROOT, "web"))
    import server as S   # noqa: E402
except BaseException as e:
    mark("engine", "قراءةُ web/server.py", None, "لم يُقَس: %s" % str(e)[:120])

if S is not None:
    exts = getattr(S, "_WS_EXTS", set())
    miss = [e for e in (".docx", ".pptx", ".xlsx", ".png") if e not in exts]
    mark("engine", "ملفٌّ مبنيٌّ في مجلّد العمل يصل بطاقة (docx/pptx/xlsx/png)",
         not miss, "ينقص: " + ", ".join(miss) if miss else "")

    # الملفُّ المرفوع: يُرسَل كما ترسله الواجهة (base64)، ثمّ نرى هل كُتب ملفٌّ.
    src = built.get("word-ar.docx")
    if src:
        from pipeline import weaver_core as W
        ws = tempfile.mkdtemp(prefix="weaver-office-ws-")
        _rw = W.workspace_dir
        W.workspace_dir = lambda: ws
        try:
            with open(src, "rb") as f:
                b64 = base64.b64encode(f.read()).decode()
            _up = [{"name": "مرفق.docx", "data": b64}]
            txt, names = S._attach_extract(_up)
            if hasattr(S, "_attach_save"):
                S._attach_save(_up, "probe-chat")     # ما يفعله الخادمُ قبل النوبة
            saved = [os.path.join(dp, fn) for dp, _, fns in os.walk(ws) for fn in fns]
            mark("engine", "الملفُّ المرفوع يُقرأ نصّاً", AR_TITLE in txt,
                 "%d حرفاً" % len(txt))
            same = False
            if saved:
                with open(saved[0], "rb") as fh, open(src, "rb") as fs:
                    same = fh.read() == fs.read()
            mark("engine", "الملفُّ المرفوع يُحفظ في مجلّد العمل (ليُعدَّل هو نفسُه)",
                 bool(saved) and same,
                 (", ".join(os.path.relpath(s, ws) for s in saved)
                  + ("" if same else " — المحتوى مختلف!")) if saved
                 else "لا — يصل النموذجَ نصّاً فقط ⟵ تعديلُه يعني إعادةَ بنائه")
        except BaseException as e:
            mark("engine", "الملفُّ المرفوع", None, "لم يُقَس: %s" % str(e)[:120])
        finally:
            W.workspace_dir = _rw
            shutil.rmtree(ws, ignore_errors=True)
    else:
        mark("engine", "الملفُّ المرفوع", None, "لم يُقَس — لم يُبنَ ملفٌّ لرفعه")

# مهاراتُ المحرّك لملفّات أوفيس — عندنا وفي أوبن كلاو (للمقارنة)
try:
    from pipeline import weaver_core as W
    ours = sorted(os.listdir(os.path.join(_ROOT, "capabilities", "prompts", "skills")))
    office = [s for s in ours if any(k in s for k in
                                     ("word", "docx", "powerpoint", "pptx", "excel",
                                      "xlsx", "chart", "office"))]
    mark("engine", "مهاراتُ Weaver للمحرّك لملفّات أوفيس", bool(office) or False,
         ", ".join(office) if office else "لا شيء بعد — الأدواتُ في المسار القديم فقط")
    if office:
        # مُركَّبةٌ في مساحة عمل المحرّك؟ (--skills apply)
        try:
            st = {r["name"]: r["state"] for r in W.skills_state()["skills"]}
            miss = [n for n in office if st.get(n) != "ours"]
            mark("engine", "  ⟵ مُركَّبةٌ في مساحة عمل المحرّك", not miss,
                 ("غيرُ مُركَّبة: %s ⟵ python3 -m pipeline.weaver_core "
                  "--skills apply" % ", ".join(miss)) if miss else "")
        except BaseException as e:
            mark("engine", "  ⟵ مُركَّبةٌ في مساحة عمل المحرّك", None,
                 "لم يُقَس: %s" % str(e)[:100])
    rt_sk = os.path.join(getattr(W, "RUNTIME", ""), "skills")
    if os.path.isdir(rt_sk):
        import re as _re
        _rx = _re.compile(r"(^|[-_])(docx?|pptx?|xlsx?|excel|word|slides?|sheets?"
                          r"|charts?|diagram|pdf|office)([-_]|$)")
        oc = [s for s in sorted(os.listdir(rt_sk)) if _rx.search(s)]
        mark("engine", "مهاراتُ أوبن كلاو المضمَّنة لملفّاتٍ (للمقارنة)", None,
             ", ".join(oc) or "لا شيء")
    else:
        mark("engine", "مهاراتُ أوبن كلاو المضمَّنة", None,
             "لم يُقَس — المحرّكُ غيرُ مركَّب")
except BaseException as e:
    mark("engine", "المهارات", None, "لم يُقَس: %s" % str(e)[:120])

# ── الخلاصة ─────────────────────────────────────────────────────────────
say("\n══ الخلاصة ══")
for sec, title in (("libs", "المكتبات"), ("build", "البناء"), ("edit", "التعديل"),
                   ("tool", "الأداة"), ("engine", "المحرّك")):
    rows = [r for r in results if r[0] == sec]
    okc = sum(1 for r in rows if r[2] is True)
    bad = [r[1].strip() for r in rows if r[2] is False]
    un = sum(1 for r in rows if r[2] is None)
    say("  %-8s ✓ %d · ✗ %d · لم يُقَس/اختياريّ %d%s" % (
        title, okc, len(bad), un, ("\n      ✗ " + "\n      ✗ ".join(bad)) if bad else ""))

if keep:
    try:
        os.makedirs(keep, exist_ok=True)
        n = 0
        for fn in sorted(os.listdir(OUT)):
            fp = os.path.join(OUT, fn)
            if os.path.isfile(fp):
                shutil.copy2(fp, os.path.join(keep, fn))
                n += 1
        say("\n  النماذجُ (%d ملفّاً) نُسخت إلى: %s — افتحها وانظر جودتها." % (n, keep))
    except BaseException as e:
        say("\n  تعذّر النسخ إلى %s: %s" % (keep, e))
shutil.rmtree(OUT, ignore_errors=True)
say("")
