"""page_render.py — أيُّ ملفٍّ ⟵ صفحاتٌ صوراً، للوحة المعاينة الجانبيّة.

    pages_for(src, cache_root) ⟵ {"kind", "n", "id", "w", "h"}
        kind: "pages"  مستندٌ/PDF/جدولٌ/نصّ — صفحاتٌ كما تُطبع
              "slides" عرضٌ تقديميّ — شريحةٌ لكلِّ صورة
              "image"  صورة — تُعرض كما هي
              "html"   صفحةٌ حيّة (رسمٌ تفاعليّ) — في إطارٍ معزول
              ""       لا معاينة
    page_path(cache_root, id, n) ⟵ مسارُ صورة الصفحة n (من ١) أو ""

· Word ⟵ web/docx_render.py (تخطيطٌ حقيقيٌّ من الملفّ).
· Excel وMarkdown والنصّ وCSV ⟵ تُحوَّل مستندَ Word ثمّ تُرسم بالمحرّك نفسِه:
  فلها صفحاتٌ حقيقيّة (تقسيمٌ، وصفُّ العنوان يتكرّر، واتّجاهٌ صحيح).
· PDF ⟵ PyMuPDF أو pdftoppm؛ وبدونهما نصُّه.
· PowerPoint ⟵ web/thumbs._render_slide لكلِّ شريحة.

الصورُ تُحفظ في مجلّدٍ للمعاينات (مفتاحُه بصمةُ الملفّ)، فلا يُعاد الرسمُ
إلا إن تغيّر الملفّ. لا يرفع.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

VERSION = "1"          # يُرفع إن تغيّر الرسم ⟵ تُعاد الصفحات
MAX_PAGES = 60
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".log", ".py", ".js", ".ts",
            ".css", ".xml", ".yaml", ".yml", ".ini", ".cfg", ".sql", ".sh", ".c",
            ".cpp", ".h", ".java", ".go", ".rs", ".tex", ".bib", ".svg"}
_AR = re.compile("[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")


def kind_of(name):
    ext = os.path.splitext(str(name or ""))[1].lower()
    if ext in IMAGE_EXT:
        return "image"
    if ext in (".html", ".htm"):
        return "html"
    if ext in (".pptx", ".pptm"):
        return "slides"
    if ext in (".docx", ".docm", ".pdf", ".xlsx", ".xlsm", ".csv", ".tsv") or \
            ext in TEXT_EXT:
        return "pages"
    return ""


def _key_for_file(path):
    st = os.stat(path)
    h = hashlib.sha1(("%s|%d|%d|%s" % (os.path.abspath(path), st.st_size,
                                        int(st.st_mtime), VERSION)).encode("utf-8"))
    return h.hexdigest()[:24]


def _prune(root, days=7):
    try:
        now = time.time()
        for d in os.listdir(root):
            p = os.path.join(root, d)
            if os.path.isdir(p) and now - os.path.getmtime(p) > days * 86400:
                shutil.rmtree(p, ignore_errors=True)
    except Exception:
        pass


def pages_for(src, cache_root, name=None):
    """صفحاتُ الملفّ (تُرسم مرّةً وتُحفظ)."""
    name = name or os.path.basename(src)
    kind = kind_of(name)
    if kind in ("image", "html", ""):
        return {"kind": kind, "n": 1 if kind else 0, "id": ""}
    try:
        key = _key_for_file(src)
    except Exception:
        return {"kind": "", "n": 0, "id": ""}
    d = os.path.join(cache_root, key)
    meta = os.path.join(d, "meta.json")
    if os.path.isfile(meta):
        try:
            with open(meta, encoding="utf-8") as fh:
                m = json.load(fh)
            os.utime(d, None)
            return m
        except Exception:
            pass
    os.makedirs(cache_root, exist_ok=True)
    _prune(cache_root)
    try:
        imgs = render(src, name)
    except Exception:
        imgs = []
    if not imgs:
        return {"kind": "", "n": 0, "id": ""}
    tmp = d + ".tmp%d" % os.getpid()
    os.makedirs(tmp, exist_ok=True)
    for i, im in enumerate(imgs[:MAX_PAGES]):
        im.convert("RGB").save(os.path.join(tmp, "p%03d.jpg" % (i + 1)), "JPEG",
                               quality=86, optimize=True)
    m = {"kind": kind, "n": min(len(imgs), MAX_PAGES), "id": key,
         "w": imgs[0].width, "h": imgs[0].height}
    with open(os.path.join(tmp, "meta.json"), "w", encoding="utf-8") as fh:
        json.dump(m, fh)
    try:
        if os.path.isdir(d):
            shutil.rmtree(d, ignore_errors=True)
        os.replace(tmp, d)
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
    return m


def page_path(cache_root, key, n):
    if not re.fullmatch(r"[0-9a-f]{8,40}", str(key or "")):
        return ""
    try:
        n = int(n)
    except (TypeError, ValueError):
        return ""
    p = os.path.join(cache_root, key, "p%03d.jpg" % n)
    return p if os.path.isfile(p) else ""


# ───────────────────────────── الرسم ─────────────────────────────
def render(src, name=None):
    ext = os.path.splitext(name or src)[1].lower()
    if ext in (".docx", ".docm"):
        import docx_render
        return docx_render.render_docx(src)
    if ext == ".pdf":
        imgs = _pdf_images(src)
        if imgs:
            return imgs
        return _via_docx(_text_docx(_pdf_text(src), ".txt"))
    if ext in (".pptx", ".pptm"):
        return _pptx_images(src)
    if ext in (".xlsx", ".xlsm"):
        return _via_docx(_xlsx_docx(src))
    if ext in (".csv", ".tsv"):
        return _via_docx(_csv_docx(src, ext))
    if ext in TEXT_EXT:
        with open(src, "rb") as fh:
            txt = fh.read(600000).decode("utf-8", "replace")
        return _via_docx(_text_docx(txt, ext))
    return []


def _via_docx(doc):
    if doc is None:
        return []
    import docx_render
    tmp = tempfile.mkdtemp(prefix="wv-pg-")
    try:
        p = os.path.join(tmp, "x.docx")
        doc.save(p)
        return docx_render.render_docx(p)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _pdf_images(src, width=900):
    from PIL import Image
    out = []
    try:
        import fitz
        with fitz.open(src) as doc:
            for i in range(min(MAX_PAGES, doc.page_count)):
                pg = doc[i]
                z = width / max(1.0, pg.rect.width)
                pix = pg.get_pixmap(matrix=fitz.Matrix(z, z), alpha=False)
                out.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
        if out:
            return out
    except Exception:
        pass
    if shutil.which("pdftoppm"):
        tmp = tempfile.mkdtemp(prefix="wv-pdf-")
        try:
            subprocess.run(["pdftoppm", "-f", "1", "-l", str(MAX_PAGES), "-jpeg",
                            "-scale-to-x", str(width), "-scale-to-y", "-1",
                            src, os.path.join(tmp, "p")],
                           capture_output=True, timeout=240)
            for fn in sorted(os.listdir(tmp),
                             key=lambda n: int(re.sub(r"\D", "", n) or 0)):
                with Image.open(os.path.join(tmp, fn)) as im:
                    out.append(im.convert("RGB"))
        except Exception:
            pass
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return out


def _pdf_text(src):
    try:
        from pypdf import PdfReader
        r = PdfReader(src)
        return "\n\n".join((p.extract_text() or "") for p in r.pages[:MAX_PAGES])
    except Exception:
        return ""


def _pptx_images(src, width=1280):
    from pptx import Presentation
    from PIL import Image
    import thumbs
    prs = Presentation(src)
    out = []
    for sl in list(prs.slides)[:MAX_PAGES]:
        im = None
        try:
            im = thumbs._render_slide(prs, sl, width)
        except Exception:
            im = None
        if im is None:
            sw = int(prs.slide_width or 12192000)
            sh = int(prs.slide_height or 6858000)
            im = Image.new("RGB", (width, int(width * sh / sw)), (255, 255, 255))
        out.append(im)
    return out


# ─────────────── تحويلٌ إلى Word (ليُرسم بالمحرّك نفسِه) ───────────────
def _new_doc(landscape=False):
    import docx
    from docx.shared import Mm, Pt
    from docx.enum.section import WD_ORIENT
    d = docx.Document()
    s = d.sections[0]
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Mm(297), Mm(210)
    else:
        s.page_width, s.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Mm(15 if landscape else 20))
    st = d.styles["Normal"]
    st.font.name = "Tajawal"
    st.font.size = Pt(10.5)
    _set_cs_font(st.element, "Tajawal")
    st.paragraph_format.space_after = Pt(4)
    return d


def _set_cs_font(el, name, size_half=None):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    rpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else None
    if rpr is None:
        return
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)
    if size_half:
        for t in ("w:sz", "w:szCs"):
            e = rpr.find(qn(t))
            if e is None:
                e = OxmlElement(t)
                rpr.append(e)
            e.set(qn("w:val"), str(size_half))


def _bidi(p, on=True):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    if not on:
        return
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(OxmlElement("w:bidi"))


def _shade(cell, hexfill):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexfill)
    tcpr.append(shd)


def _run(p, text, bold=False, italic=False, color=None, size=None, font=None):
    from docx.shared import Pt, RGBColor
    r = p.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    if bold:
        from docx.oxml import OxmlElement
        r._r.get_or_add_rPr().append(OxmlElement("w:bCs"))
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
        _set_cs_font(r._r, font or "Tajawal", int(size * 2))
    elif font:
        _set_cs_font(r._r, font)
    return r


def _inline(p, text, base_size=None, font=None):
    """**عريض** و*مائل* و`شيفرة` داخل سطر Markdown."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|__[^_]+__)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") or \
                part.startswith("__") and part.endswith("__"):
            _run(p, part[2:-2], bold=True, size=base_size, font=font)
        elif part.startswith("`") and part.endswith("`"):
            _run(p, part[1:-1], font="JetBrains Mono", color="8A2D3B", size=base_size)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _run(p, part[1:-1], italic=True, size=base_size, font=font)
        else:
            _run(p, part, size=base_size, font=font)


def _text_docx(txt, ext):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    d = _new_doc()
    lines = (txt or "").replace("\r\n", "\n").split("\n")
    if ext not in (".md", ".markdown"):
        mono = ext not in (".txt", ".log", ".tex", ".bib")
        for ln in lines[:4000]:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            ar = bool(_AR.search(ln))
            _bidi(p, ar and not mono)
            if ar and not mono:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT      # في فقرةٍ RTL: يمين
            _run(p, ln.replace("\t", "    ") or " ",
                 font="JetBrains Mono" if mono else None, size=9 if mono else None)
        return d
    i = 0
    in_code = False
    while i < len(lines) and i < 6000:
        ln = lines[i].rstrip()
        if ln.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            _run(p, ln or " ", font="JetBrains Mono", size=9)
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F2F3F5")
            p._p.get_or_add_pPr().append(shd)
            i += 1
            continue
        # جدول
        if "|" in ln and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows = []
            while i < len(lines) and "|" in lines[i]:
                if not re.match(r"^\s*\|?\s*:?-{2,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _table(d, rows)
            continue
        ar = bool(_AR.search(ln))
        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            lvl = len(m.group(1))
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(10 if lvl <= 2 else 6)
            p.paragraph_format.keep_with_next = True
            _bidi(p, ar)
            _inline(p, "", None)
            _run(p, re.sub(r"\*\*|__", "", m.group(2)), bold=True,
                 color="1F3864", size={1: 18, 2: 15, 3: 13}.get(lvl, 12), font="Cairo")
        elif re.match(r"^\s*([-*_])\1{2,}\s*$", ln):
            p = d.add_paragraph()
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            bdr = OxmlElement("w:pBdr")
            b = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:color", "BBBBBB")):
                b.set(qn(k), v)
            bdr.append(b)
            p._p.get_or_add_pPr().append(bdr)
        elif re.match(r"^\s*[-*+]\s+", ln):
            p = d.add_paragraph(style="List Bullet")
            _bidi(p, ar)
            _inline(p, re.sub(r"^\s*[-*+]\s+", "", ln))
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            p = d.add_paragraph()
            _bidi(p, ar)
            _inline(p, ln.strip())
        elif ln.startswith(">"):
            p = d.add_paragraph()
            _bidi(p, ar)
            p.paragraph_format.left_indent = Pt(14)
            _run(p, ln.lstrip("> "), italic=True, color="555555")
        elif not ln.strip():
            i += 1
            continue
        else:
            p = d.add_paragraph()
            _bidi(p, ar)
            _inline(p, ln)
        i += 1
    return d


def _table(d, rows, rtl=None, widths=None, fills=None, bolds=None, header=True,
           colors=None):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    rows = [r for r in rows if r]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = d.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    if rtl is None:
        rtl = any(_AR.search(str(c)) for r in rows[:20] for c in r)
    if rtl:
        t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    if widths:
        grid = t._tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w)))
    for ri, r in enumerate(rows):
        if ri == 0 and header:
            trpr = t.rows[0]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:tblHeader"))
        for ci in range(ncol):
            val = str(r[ci]) if ci < len(r) and r[ci] is not None else ""
            c = t.cell(ri, ci)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _bidi(p, rtl and bool(_AR.search(val)) or rtl)
            bold = (ri == 0 and header) or bool(bolds and bolds.get((ri, ci)))
            _run(p, val, bold=bold, size=9, color=(colors or {}).get((ri, ci)))
            fill = (fills or {}).get((ri, ci)) or ("E8EEF6" if ri == 0 and header else None)
            if fill:
                _shade(c, fill)
    return t


def _csv_docx(src, ext):
    with open(src, "rb") as fh:
        txt = fh.read(600000).decode("utf-8", "replace")
    rows = list(csv.reader(io.StringIO(txt), delimiter="\t" if ext == ".tsv" else ","))
    rows = [r[:30] for r in rows[:2000]]
    wide = max((len(r) for r in rows), default=0) > 6
    d = _new_doc(landscape=wide)
    _table(d, rows)
    return d


def _fmt_cell(v):
    import datetime
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return ("%.4f" % v).rstrip("0").rstrip(".")
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if not (v.hour or v.minute) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _xlsx_docx(src):
    import openpyxl
    from docx.shared import Pt
    wb = openpyxl.load_workbook(src, data_only=True)
    d = None
    for si, ws in enumerate(wb.worksheets[:8]):
        max_r = min(ws.max_row or 0, 1000)
        max_c = min(ws.max_column or 0, 26)
        rows, fills, bolds, colors = [], {}, {}, {}
        for r in range(1, max_r + 1):
            row = []
            for c in range(1, max_c + 1):
                cell = ws.cell(r, c)
                row.append(_fmt_cell(cell.value))
                try:
                    f = cell.fill
                    if f is not None and f.fill_type == "solid":
                        rgb = f.fgColor.rgb if f.fgColor is not None else None
                        if isinstance(rgb, str) and len(rgb) >= 6 and rgb[-6:] != "000000":
                            fills[(len(rows), c - 1)] = rgb[-6:]
                    if cell.font is not None and cell.font.b:
                        bolds[(len(rows), c - 1)] = True
                    fc = cell.font.color.rgb if cell.font is not None and \
                        cell.font.color is not None else None
                    if isinstance(fc, str) and len(fc) >= 6 and fc[-6:] != "000000":
                        colors[(len(rows), c - 1)] = fc[-6:].upper()
                except Exception:
                    pass
            rows.append(row)
        # الصفوفُ والأعمدةُ الفارغة في الآخر لا تُعرض
        while rows and not any(v.strip() for v in rows[-1]):
            rows.pop()
        if not rows:
            continue
        ncol = max((max((i + 1 for i, v in enumerate(r) if v.strip()), default=0)
                    for r in rows), default=0)
        rows = [r[:ncol] for r in rows]
        widths = []
        for c in range(1, ncol + 1):
            letter = openpyxl.utils.get_column_letter(c)
            w = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
            widths.append((w or 10) * 7 * 15)          # حرفٌ ≈ ٧px ≈ ١٠٥ twip
        total = sum(widths) or 1
        landscape = total > 9000 or ncol > 6
        avail = (297 - 30 if landscape else 210 - 40) * 56.7
        # الجدولُ يملأ عرضَ الصفحة (أسهلُ قراءةً على الهاتف)
        widths = [w * avail / total for w in widths]
        rtl = False
        try:
            rtl = bool(ws.sheet_view.rightToLeft)
        except Exception:
            pass
        rtl = rtl or any(_AR.search(v) for r in rows[:30] for v in r)
        if d is None:
            d = _new_doc(landscape=landscape)
        else:
            p = d.add_paragraph()
            from docx.enum.text import WD_BREAK
            p.add_run().add_break(WD_BREAK.PAGE)
        p = d.add_paragraph()
        _bidi(p, bool(_AR.search(ws.title)))
        p.paragraph_format.space_after = Pt(6)
        _run(p, ws.title, bold=True, color="1E6E3C", size=12, font="Cairo")
        header_bold = any(bolds.get((0, c)) for c in range(ncol)) or not fills
        _table(d, rows, rtl=rtl, widths=widths, fills=fills, bolds=bolds,
               header=header_bold, colors=colors)
    wb.close()
    return d
