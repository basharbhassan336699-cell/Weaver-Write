# -*- coding: utf-8 -*-
"""إصلاحاتُ اختبارات المستخدم الثمانية (هاتف، ملفّاتٌ حقيقيّة).

  · بطاقةُ «احفظ في التقويم» بعد كلِّ ملفّ ولا موعدَ في الطلب.
  · وورد: «شكل ١» في صفحةٍ والرسمُ في التي قبلها.
  · وورد: طُلبت صفحتان فخرجت أربعٌ وقيل «صفحتان تقريباً» ⟵ تقديرُ الصفحات.
  · العروض: البطاقاتُ والدوائرُ والخطواتُ والأرقامُ في النصف الأعلى وتحتها فراغ.
  · العروض: أرقامُ Kufyan العربيّة بأشكالٍ غيرِ مألوفة («٣» كأنّها «Ш»).
  · المهارات: الموضوعُ والتنسيق، الأرقامُ المختلَقة، الرموزُ الداخليّة، اللغة.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "web"))
sys.path.insert(0, os.path.join(_ROOT, "capabilities/skills/docx_builder/scripts"))
sys.path.insert(0, os.path.join(_ROOT, "capabilities/skills/pptx_builder/scripts"))

P = F = 0


def ok(name, cond, extra=""):
    global P, F
    if cond:
        P += 1
        print(f"  ✓ {name}")
    else:
        F += 1
        print(f"  ✗ {name}" + (f"   — {extra}" if extra else ""))


T = tempfile.mkdtemp(prefix="weaver-test-fixes-")
OFFICE = os.path.join(_ROOT, "pipeline", "office.py")


def cli(*args):
    r = subprocess.run([sys.executable, OFFICE] + [str(a) for a in args],
                       capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout + r.stderr


def J(name, obj):
    p = os.path.join(T, name)
    json.dump(obj, open(p, "w", encoding="utf-8"), ensure_ascii=False)
    return p


AR = ("القراءةُ اليوميّةُ تبني لدى الطالب مفرداتٍ أوسع وقدرةً أكبر على التركيز، "
      "وتمنحه ثقةً حين يكتب نصوصَه بنفسه، لأنّه رأى كيف يرتّب الكتّابُ أفكارَهم "
      "ويصلون بين الجمل. والمكتبةُ المدرسيّةُ مكانٌ هادئٌ لهذه العادة، ففيها كتبٌ "
      "لا تقدر عليها أسرةٌ واحدة، وأمينٌ يرشد القارئ إلى ما يناسب عمرَه واهتمامه.")

# ───────────────────────── ١) بطاقةُ التقويم ─────────────────────────
print("\n— بطاقةُ التقويم: فقط حين يُذكر موعد —")
import server as S   # noqa: E402
for txt, want in (("ارسم لي رسماً بيانياً لعدد الحالات", False),
                  ("Make a 6-slide presentation about time management", False),
                  ("ذكّرني بالاجتماع بعد أسبوع", True),
                  ("الموعد غداً الساعة ١٠", True),
                  ("deadline 2026-11-03", True)):
    _o, m = S._calendar_suggest(txt, with_flag=True)
    ok("«%s» ⟵ mentioned=%s" % (txt[:30], want), m is want, m)
ok("الاستدعاءُ القديم بلا علَمٍ كما كان (قائمة)",
   isinstance(S._calendar_suggest("غداً"), list))
html = open(os.path.join(_ROOT, "web", "index.html"), encoding="utf-8").read()
ok("الواجهة: لا بطاقةَ حين mentioned=false", "d.mentioned===false" in html)

# ───────────────────── ٢) وورد: الصورةُ مع تعليقها ─────────────────────
print("\n— وورد: الصورةُ وتعليقُها في صفحةٍ واحدة —")
from PIL import Image   # noqa: E402
img = os.path.join(T, "img.png")
Image.new("RGB", (800, 500), (40, 80, 160)).save(img)
spec = {"title": "تقرير", "cover": False, "sections": [
    {"heading": "المقدّمة", "body": AR, "image": {"path": img, "caption": "شكل ١: صورة"}}]}
rc, out = cli("build", J("img.json", spec), "--out", os.path.join(T, "img.docx"))
ok("build ⟵ ✓", rc == 0, out)
from docx import Document   # noqa: E402
d = Document(os.path.join(T, "img.docx"))
pi = [i for i, p in enumerate(d.paragraphs) if p._p.xpath(".//w:drawing")]
ok("فقرةُ الصورة «ابقَ مع التالي» (keepNext) وبعدها التعليق", bool(pi) and
   "keepNext" in d.paragraphs[pi[0]]._p.xml
   and d.paragraphs[pi[0] + 1].text.startswith("شكل ١"), pi)
ops = J("ops.json", [{"op": "add_image", "after": 1, "path": img, "caption": "صورة مضافة"}])
rc, out = cli("edit", os.path.join(T, "img.docx"), ops)
d = Document(os.path.join(T, "img-edited.docx"))
ok("add_image بالتعديل: keepNext كذلك", all(
    "keepNext" in p._p.xml for p in d.paragraphs if p._p.xpath(".//w:drawing")), out)
ok("  ⟵ والتعديلُ يذكر الطول (≈ N pages)", "length ≈" in out, out)

# ───────────────────── ٣) وورد: تقديرُ الصفحات ─────────────────────
print("\n— وورد: تقديرُ الصفحات —")
import docx_pages as DP   # noqa: E402
long_spec = {"title": "فوائد القراءة", "cover": False, "pages": 2, "sections": [
    {"heading": "القسم %d" % i, "body": "\n\n".join([AR] * 3)} for i in range(6)]}
rc, out = cli("build", J("long.json", long_spec), "--out", os.path.join(T, "long.docx"))
m = re.search(r"≈ (\d+) pages? \(estimate ([\d.]+)", out)
ok("build يذكر ≈ N pages", rc == 0 and m is not None, out)
ok("  ⟵ وطُلبت صفحتان والملفُّ أطول ⟵ ⚠ cut ~N words", "⚠ requested 2 pages" in out
   and "cut ~" in out, out)
short_spec = dict(long_spec, pages=1, sections=long_spec["sections"][:1])
short_spec["sections"] = [{"heading": "مقدّمة", "body": AR}]
rc, out = cli("build", J("short.json", short_spec), "--out", os.path.join(T, "short.docx"))
ok("نصٌّ قصير وطُلبت صفحة ⟵ ✓ matches", "✓ matches the requested 1 page" in out, out)
e_long = DP.estimate_pages(os.path.join(T, "long.docx"))
e_short = DP.estimate_pages(os.path.join(T, "short.docx"))
ok("الأطولُ أكثرُ صفحات", e_long["exact"] > e_short["exact"] + 1, (e_long, e_short))

# معايرة: كلُّ خطٍّ عربيٍّ عندنا ⟵ عددُ الأسطر في فقرةٍ بعرض ٦ بوصات
for fam in ("Amiri", "Kufyan Arabic Regular", "Cairo", "Tajawal", "Noto Naskh Arabic"):
    met = DP._metrics(fam)
    ok("مقاييسُ %s من ملفّه (سطر، أشكالُ الحروف المتّصلة)" % fam, met is not None
       and 1.0 < met["line"] < 2.5 and any(c >= 0xFE70 for c in met["adv"]), met and met["line"])
# مقياسٌ ثابت: Amiri ١٤pt بعرض ٤٣٢pt — LibreOffice: ٥ أسطر (قِيس على تقرير المستخدم)
_t5 = ("يُعَدّ اختيارُ الخطّ أحدَ العوامل المؤثّرة في تجربة القراءة الطويلة، وإن كان القارئُ "
       "قلّما ينتبه إليه حين يكون الخطُّ جيّداً. خطّ أميري خطٌّ عربيٌّ مجانيّ صمّمه الطبيبُ "
       "والخطّاطُ خالد حسني، مستوحىً من طباعة كتاب أميري الذي طُبِع في المطبعة الأميريّة "
       "ببولاق في مطلع القرن العشرين. اكتسب الخطُّ شهرةً واسعةً لصفاته الوظيفيّة، من وضوح "
       "الحروف واتّزان أشكالها إلى انسيابيّة القراءة في النصوص الطويلة.")
from docx.shared import Pt   # noqa: E402
from docx.oxml import OxmlElement   # noqa: E402
from docx.oxml.ns import qn   # noqa: E402
dd = Document()
pp = dd.add_paragraph()
pp._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
rr = pp.add_run(_t5)
rr.font.size = Pt(14)
rf = OxmlElement("w:rFonts")
for k in ("ascii", "hAnsi", "cs"):
    rf.set(qn("w:" + k), "Amiri")
rr._r.get_or_add_rPr().insert(0, rf)
sz = OxmlElement("w:szCs")
sz.set(qn("w:val"), "28")
rr._r.get_or_add_rPr().append(sz)
St = DP._Styles(dd)
n, lh, _nw = DP._para_lines(pp._p, St, 432.0, DP._para_props(pp._p, St))
ok("فقرةُ التقرير (Amiri ١٤) ⟵ ٥ أسطر بارتفاع ٢٨٫٣pt كما في LibreOffice",
   n == 5 and abs(lh - 28.3) < 0.5, (n, lh))

_SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
if _SOFFICE and shutil.which("pdfinfo"):
    subprocess.run([_SOFFICE, "-env:UserInstallation=file://%s/lo" % T, "--headless",
                    "--convert-to", "pdf", "--outdir", T, os.path.join(T, "long.docx")],
                   capture_output=True, timeout=240)
    pdf = os.path.join(T, "long.pdf")
    if os.path.isfile(pdf):
        real = int(re.search(r"Pages:\s+(\d+)", subprocess.run(
            ["pdfinfo", pdf], capture_output=True, text=True).stdout).group(1))
        ok("LibreOffice: الصفحاتُ الحقيقيّة = التقدير (%d)" % real,
           real == e_long["pages"], e_long)
else:
    print("  ⓘ لا LibreOffice — تُخطّى المقارنةُ الحيّة")

# ───────────────────── ٤) العروض: الفراغ والأرقام ─────────────────────
print("\n— العروض: التوسيط العموديّ، وأرقامٌ مألوفةُ الأشكال —")
import pptx_design as PD   # noqa: E402
from pptx import Presentation   # noqa: E402
deck = {"title": "الاضطرابات السلوكيّة", "font": "Kufyan Arabic Black", "slides_total": 6,
        "slides": [
            {"title": "بطاقات", "cards": [{"title": "المشكلة العابرة", "text": "سلوكٌ موقفيّ يزول"},
                                          {"title": "الاضطراب", "text": "سلوكٌ يطول ويتكرّر"},
                                          {"title": "معيار الحكم", "text": "العمر والموقف"}]},
            {"title": "دوائر", "circles": [{"title": "وراثيّة", "text": "عوامل جينيّة"},
                                          {"title": "أسريّة", "text": "قسوة"},
                                          {"title": "مدرسيّة", "text": "فشل"}]},
            {"title": "خطوات", "steps": [{"title": "الملاحظة"}, {"title": "التشخيص"},
                                         {"title": "الخطّة"}]},
            {"title": "أرقام", "stats": [{"value": "٣٠٪", "text": "نسبة"},
                                         {"value": "ADHD", "text": "قصور الانتباه"}]}]}
out_p = os.path.join(T, "deck.pptx")
rc, out = cli("build", J("deck.json", deck), "--out", out_p)
ok("build ⟵ ✓ 6 slides", rc == 0 and "6 slides" in out, out)
prs = Presentation(out_p)
EMU = 914400
mid = (PD.TOP + PD.BOTTOM) / 2
for idx, kind in ((1, "cards"), (2, "circles"), (3, "steps"), (4, "stats")):
    sl = prs.slides[idx]
    ys = [(sh.top / EMU, (sh.top + sh.height) / EMU) for sh in sl.shapes
          if sh.top / EMU >= PD.TOP - 0.05 and (sh.top + sh.height) / EMU <= PD.BOTTOM + 0.05
          and sh.height / EMU < 5.5]
    top, bot = min(y[0] for y in ys), max(y[1] for y in ys)
    ok("%s: الكتلةُ في وسط منطقة المحتوى (%.2f–%.2f)" % (kind, top, bot),
       abs((top + bot) / 2 - mid) < 0.6 and bot - top > 2.0, (top, bot, mid))


def runs(slide):
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                cs = r._r.find(A + "rPr/" + A + "cs")
                yield r.text, cs.get("typeface") if cs is not None else None


rs = list(runs(prs.slides[2]))
digits = [f for t, f in rs if re.fullmatch(r"[٠-٩]+", t.strip() or "x")]
words = [f for t, f in rs if re.search(r"[ء-ي]", t)]
ok("Kufyan: أرقامُ الدوائر بخطٍّ أرقامُه مألوفة (%s)" % PD.DIGIT_FONT,
   digits and all(f == PD.DIGIT_FONT for f in digits), rs)
ok("  ⟵ والنصُّ كلُّه بالخطّ المطلوب", words and all(f == "Kufyan Arabic Black"
                                                    for f in words), rs)
deck2 = dict(deck, font="Amiri")
rc, out = cli("build", J("deck2.json", deck2), "--out", os.path.join(T, "deck2.pptx"))
rs2 = list(runs(Presentation(os.path.join(T, "deck2.pptx")).slides[2]))
ok("Amiri: الأرقامُ بخطّ العرض نفسِه (لا تبديل)", all(
    f == "Amiri" for t, f in rs2 if re.fullmatch(r"[٠-٩]+", t.strip() or "x")), rs2)
# بطاقةٌ بعنوانٍ إنجليزيٍّ يلتفّ: لا يغطّي النصَّ تحته
en = {"title": "Deck", "slides_total": 3, "slides": [
    {"title": "Problem versus disorder", "cards": [
        {"title": "Passing problem", "text": "Situational behaviour that fades"},
        {"title": "Disorder", "text": "Lasting, frequent and severe"},
        {"title": "Who decides", "text": "A specialist, not guesswork"}]}]}
rc, out = cli("build", J("en.json", en), "--out", os.path.join(T, "en.pptx"))
sl = Presentation(os.path.join(T, "en.pptx")).slides[1]
tb = {sh.text_frame.text: sh for sh in sl.shapes if sh.has_text_frame}
t_, b_ = tb.get("Passing problem"), tb.get("Situational behaviour that fades")
ok("عنوانُ البطاقة يتّسع لسطرين فلا يغطّي نصَّها", t_ is not None and b_ is not None
   and t_.height / EMU >= 0.9 and t_.top + t_.height <= b_.top + 1)

# ───────────────────────── ٥) المهارات ─────────────────────────
print("\n— المهارات: قواعدُ قِيست في الاختبارات —")
SK = os.path.join(_ROOT, "capabilities", "prompts", "skills")


def skill(n):
    return open(os.path.join(SK, n, "SKILL.md"), encoding="utf-8").read()


w, pw, ch, xl = (skill("office-word"), skill("office-powerpoint"),
                 skill("office-charts"), skill("office-excel"))
ok("وورد: «بخطّ X» تنسيقٌ لا موضوع", "تنسيقٌ لا موضوع" in w)
ok("وورد: \"pages\" وتقديرُ الأداة", '"pages": 2' in w and "≈ N pages" in w)
ok("وورد: رموزُ info (¶) لا تظهر للمستخدم", "¶12" in w and "لك لا للمستخدم" in w)
ok("وورد: لا أرقامَ مختلَقة، ولا مؤلّفَ لم يُذكر", "لا أرقامَ مختلَقة" in w
   and "cover.author" in w)
ok("عرض: كلُّ عنصرٍ مطلوب، والعددُ باختصار غيره", "لا تحذف المطلوب" in pw)
ok("عرض: presenter باسمٍ ذكره المستخدم فقط", "مساعدي الذكي" in pw)
ok("الكلّ: طلبٌ بالإنجليزيّة ⟵ ردٌّ بالإنجليزيّة", all(
    "بالإنجليزيّة" in s for s in (w, pw, ch, xl)))
ok("رسوم: لا بياناتٍ مختلَقة و\"source\"", "لا تخترع بيانات" in ch and '"source"' in ch)

shutil.rmtree(T, ignore_errors=True)
print("\n" + "=" * 62)
print(f" RESULT: {'PASS' if F == 0 else 'FAIL'}   ({P}/{P + F})")
print("=" * 62)
sys.exit(1 if F else 0)
