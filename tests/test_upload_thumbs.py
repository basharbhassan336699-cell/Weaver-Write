"""الملفّاتُ المرفوعة: صورةٌ ممّا فيها، تُحفظ مع رسالتها وتبقى.

    python3 tests/test_upload_thumbs.py
"""
import base64
import io
import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "web"))

FAILS = []


def check(name, cond):
    print(("  ✓ " if cond else "  ✗ ") + name)
    if not cond:
        FAILS.append(name)


def _img_ok(path, min_w=100):
    from PIL import Image
    try:
        with Image.open(path) as im:
            return im.width >= min_w and im.height >= 60
    except Exception:
        return False


def test_thumbs(tmp):
    import thumbs
    import docx
    from PIL import Image
    Image.new("RGB", (640, 400), (20, 120, 90)).save(os.path.join(tmp, "p.jpg"))
    d = docx.Document()
    d.add_heading("المسؤولية القانونية عن أخطاء الذكاء الاصطناعي", 0)
    d.add_paragraph("نصٌّ عربيٌّ في متن المستند.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "أ"
    d.save(os.path.join(tmp, "d.docx"))
    docx.Document().save(os.path.join(tmp, "empty.docx"))
    with open(os.path.join(tmp, "n.md"), "w", encoding="utf-8") as fh:
        fh.write("# عنوان\n\nسطر\n")
    cases = [("p.jpg", True), ("d.docx", True), ("n.md", True),
             ("empty.docx", False)]
    try:
        import pptx
        pr = pptx.Presentation()
        pr.slides.add_slide(pr.slide_layouts[0]).shapes.title.text = "عرض"
        pr.save(os.path.join(tmp, "s.pptx"))
        cases.append(("s.pptx", True))
    except Exception:
        pass
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["المنتج", "الكمية"])
        wb.active.append(["قلم", 3])
        wb.save(os.path.join(tmp, "x.xlsx"))
        cases.append(("x.xlsx", True))
    except Exception:
        pass
    for name, want in cases:
        out = os.path.join(tmp, name + ".img")
        got = thumbs.make_thumb(os.path.join(tmp, name), out)
        if want:
            check("thumb: %s ⟵ صورةٌ ممّا فيه" % name, got and _img_ok(out))
        else:
            check("thumb: %s فارغ ⟵ لا تُستعمل صورةُ القالب العامّة" % name,
                  not got)
    check("thumb: نوعٌ مجهول ⟵ False بلا استثناء",
          thumbs.make_thumb(os.path.join(tmp, "nothing.bin"),
                            os.path.join(tmp, "z.img")) is False)


def test_server(tmp):
    import server as srv
    old = srv._UPLOADS_DIR
    srv._UPLOADS_DIR = os.path.join(tmp, "uploads")
    try:
        data = "data:text/plain;base64," + base64.b64encode(
            "مرحبا".encode("utf-8")).decode()
        keys = srv._uploads_store([{"name": "a.txt", "data": data, "uid": "u1"},
                                   {"name": "../../evil.txt", "text": "x",
                                    "uid": "u2"}], "c123")
        check("server: الملفُّ يُحفظ للمحادثة", "u1-a.txt" in keys)
        fp = srv._upload_path("c123", "u1", "a.txt")
        check("server: يُعثر عليه بالرقم والاسم",
              fp and open(fp, encoding="utf-8").read() == "مرحبا")
        ev = srv._upload_path("c123", "u2", "../../evil.txt")
        check("server: لا خروجَ من مجلّد المحادثة",
              ev and os.path.dirname(ev) == srv._upload_dir("c123"))
        check("server: اسمٌ مجهول ⟵ لا شيء",
              srv._upload_path("c123", "zz", "nope.txt") == "")
        data, ctype = srv._thumb_for(fp, os.path.join(os.path.dirname(fp), ".thumbs"))
        check("server: صورةٌ مصغّرةٌ تُحفظ وتُعاد",
              bool(data) and ctype.startswith("image/"))
        srv._uploads_remove("c123")
        check("server: حذفُ المحادثة يحذف ملفّاتها",
              not os.path.isdir(os.path.join(tmp, "uploads", "c123")))
    finally:
        srv._UPLOADS_DIR = old


def test_ui():
    with open(os.path.join(ROOT, "web", "index.html"), encoding="utf-8") as f:
        h = f.read()
    check("ui: الرسالةُ تحفظ attachments",
          "attachments: wvAtts" in h)
    check("ui: تُعاد عند فتح المحادثة",
          "wvAddUserFiles(_um, m.attachments, id)" in h)
    check("ui: لا تُرسَل إلى النموذج في السجلّ",
          "k !== 'attachments'" in h)
    check("ui: صورةٌ في حقل الإدخال", "/api/uploads/preview" in h
          and "attach-file-thumb" in h)
    check("ui: نافذةُ المعاينة تُخفى فعلاً",
          ".file-preview-overlay[hidden] { display: none; }" in h)
    check("ui: «📎» القديمُ غيرُ المحفوظ أُزيل",
          "chip.textContent = '📎 '" not in h)


if __name__ == "__main__":
    print("upload thumbnails")
    tmp = tempfile.mkdtemp(prefix="wv-up-")
    try:
        test_thumbs(tmp)
        test_server(tmp)
        test_ui()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("FAIL: " + ", ".join(FAILS) if FAILS else "PASS")
    sys.exit(1 if FAILS else 0)
