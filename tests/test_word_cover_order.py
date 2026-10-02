# -*- coding: utf-8 -*-
"""وورد: ترتيبُ الإدراج، والغلاف — من ملفّ المستخدم.

  · كلُّ فقرةٍ في عمليّةٍ «بعد ¶24» ⟵ خرج البحثُ مقلوباً: المراجعُ (٥، ٤ … ١)
    أوّلاً، و«أولاً: …» في آخر الملفّ، واستشهاداتُ [12]…[1] تنازليّة.
  · طُلب غلافٌ فخرج «المحتويات» مكانه: "cover": {} كان يُعدّ «لا غلاف»، ومثالُ
    المهارة فيه "toc": true. وحقولُ الغلاف الفارغةُ كانت تُحذف.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OFFICE = os.path.join(_ROOT, "pipeline", "office.py")
from docx import Document   # noqa: E402

P = F = 0


def ok(name, cond, extra=""):
    global P, F
    if cond:
        P += 1
        print(f"  ✓ {name}")
    else:
        F += 1
        print(f"  ✗ {name}" + (f"   — {extra}" if extra else ""))


T = tempfile.mkdtemp(prefix="weaver-test-cover-")


def cli(*a):
    r = subprocess.run([sys.executable, OFFICE] + [str(x) for x in a],
                       capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout + r.stderr


def J(name, obj):
    p = os.path.join(T, name)
    json.dump(obj, open(p, "w", encoding="utf-8"), ensure_ascii=False)
    return p


def body(path):
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    out = []
    for el in Document(path).element.body.iterchildren():
        t = el.tag.split("}")[1]
        if t == "p":
            out.append("".join(x.text or "" for x in el.iter(W + "t")))
        elif t == "tbl":
            out.append("[TABLE]")
    return out


print("\n— ترتيبُ الإدراج بعد المرساة نفسِها —")
src = os.path.join(T, "t.docx")
d = Document()
for t in ("أ", "ب", "ج"):
    d.add_paragraph(t)
d.save(src)
rc, o = cli("edit", src, J("o1.json", [
    {"op": "insert_paragraph", "after": 1, "text": "أولاً: العنوان"},
    {"op": "insert_paragraph", "after": 1, "text": "الفقرة الأولى"},
    {"op": "insert_paragraph", "after": 1, "text": "الخاتمة"},
    {"op": "insert_paragraph", "after": 1, "text": "المراجع"}]))
ok("أربعُ عمليّاتٍ «بعد ¶1» ⟵ بترتيبها لا مقلوبة",
   body(os.path.join(T, "t-edited.docx")) == ["أ", "ب", "أولاً: العنوان",
                                               "الفقرة الأولى", "الخاتمة", "المراجع",
                                               "ج"], body(os.path.join(T, "t-edited.docx")))
rc, o = cli("edit", src, J("o2.json", [
    {"op": "insert_paragraph", "after": 0, "text": "س١"},
    {"op": "add_table", "after": 0, "headers": ["ع١", "ع٢"], "rows": [["١", "٢"]]},
    {"op": "insert_paragraph", "after": 0, "text": ["س٢", "س٣"]},
    {"op": "insert_paragraph", "after": 2, "text": "بعد ج"},
    {"op": "insert_paragraph", "before": 1, "text": "قبل ب ١"},
    {"op": "insert_paragraph", "before": 1, "text": "قبل ب ٢"}]), "--out",
    os.path.join(T, "t2.docx"))
ok("فقرةٌ وجدولٌ وفقرتان بعد المرساة نفسِها، و«قبل» كما كانت",
   body(os.path.join(T, "t2.docx")) == ["أ", "س١", "[TABLE]", "س٢", "س٣", "قبل ب ١",
                                         "قبل ب ٢", "ب", "ج", "بعد ج"],
   body(os.path.join(T, "t2.docx")))

print("\n— الغلاف —")
secs = [{"heading": "المقدّمة", "body": "نصٌّ تجريبيّ. " * 30}]
for name, cov, want in (("فارغ {}", {}, True), ("true", True, True),
                        ("false", False, False), ("بلا مفتاح", None, False)):
    spec = {"title": "عنوانُ البحث", "sections": secs}
    if cov is not None:
        spec["cover"] = cov
    out = os.path.join(T, "c-%s.docx" % abs(hash(name)))
    rc, o = cli("build", J("c.json", spec), "--out", out)
    txt = body(out)
    ok("cover %s ⟵ %s" % (name, "غلاف" if want else "بلا غلاف"),
       rc == 0 and txt.count("عنوانُ البحث") == 1
       and (any("────" in t for t in txt) == want), (rc, txt[:8], o[-200:]))
    ok("  ⟵ ولا «المحتويات» إلا بطلب", "المحتويات" not in txt, txt[:8])
spec = {"title": "عنوانُ البحث", "sections": secs, "cover": {
    "institution": "جامعة الإمارات", "course": "مدخل إلى الإعلام الرقمي",
    "author": "", "supervisor": "د. أحمد", "date": "",
    "fields": [{"label": "الرقم الجامعي", "value": ""},
               {"label": "رمز المساق", "value": "0601101"}]}}
out = os.path.join(T, "full.docx")
rc, o = cli("build", J("full.json", spec), "--out", out)
txt = body(out)
ok("حقلٌ فارغ ⟵ خانةٌ منقّطة (إعداد: ……)", any(t.startswith("إعداد:") and "……" in t
                                             for t in txt), txt[:20])
ok("حقولٌ إضافيّة: الرقمُ الجامعيّ ورمزُ المساق", any("الرقم الجامعي:" in t for t in txt)
   and any("رمز المساق: 0601101" in t for t in txt), txt[:20])
_SO = shutil.which("soffice")
if _SO and shutil.which("pdftotext"):
    subprocess.run([_SO, "-env:UserInstallation=file://%s/lo" % T, "--headless",
                    "--convert-to", "pdf", "--outdir", T, out], capture_output=True,
                   timeout=240)
    p1 = subprocess.run(["pdftotext", "-f", "1", "-l", "1", os.path.join(T, "full.pdf"),
                         "-"], capture_output=True, text=True).stdout
    ok("LibreOffice: الغلافُ كلُّه في الصفحة الأولى", "رمز" in p1 and "المساق" in p1, p1)
else:
    print("  ⓘ لا LibreOffice — تُخطّى فحصُ الصفحة")

shutil.rmtree(T, ignore_errors=True)
print("\n" + "=" * 62)
print(f" RESULT: {'PASS' if F == 0 else 'FAIL'}   ({P}/{P + F})")
print("=" * 62)
sys.exit(1 if F else 0)
