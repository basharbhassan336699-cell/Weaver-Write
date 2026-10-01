# -*- coding: utf-8 -*-
"""ملفّاتُ أوفيس في طريق المحرّك: pipeline/office.py، والملفُّ المرفوعُ نفسُه،
ومهاراتُ office-*.

قِيس على هاتف المستخدم (tools/probe_office.py) قبلها:
  ✗ الملفُّ المرفوع يُحفظ في مجلّد العمل — يصل النموذجَ نصّاً فقط
  ✗ معادلةُ المجموع تتّسع للصفّ المُدرَج — =SUM(B2:B3) تبقى
  ✗ مهاراتُ Weaver للمحرّك لملفّات أوفيس — لا شيء
وعيوبٌ قِيست أثناء البناء، وهنا اختبارُ كلٍّ منها:
  · بديلٌ يحوي المبحوثَ عنه يُستبدل مرّتين («المعدَّل المعدَّل»)
  · id() لعناصر lxml غيرُ ثابت ⟵ نُقلت عناصرُ قديمةٌ ظُنّت جديدة
  · حذفُ شريحةٍ ثمّ إضافةُ أخرى ⟵ slide7.xml مرّتين ⟵ ملفٌّ فاسد
  · رسمُ Excel بقي على B2:B3 بعد إدراج صفّ
  · openpyxl 3.1 يحفظ الرسومَ والصور — فالرفضُ المسبقُ خطأ؛ الحكمُ بالمقارنة
"""
import base64
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.request
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "web"))
from pipeline import office as O          # noqa: E402

P = F = 0


def ok(name, cond, extra=""):
    global P, F
    if cond:
        P += 1
        print(f"  ✓ {name}")
    else:
        F += 1
        print(f"  ✗ {name}" + (f"   — {extra}" if extra else ""))


def has(mod):
    try:
        __import__(mod)
        return True
    except ImportError:
        return False


HAVE_MPL = has("matplotlib")
T = tempfile.mkdtemp(prefix="weaver-test-office-")
OFFICE = os.path.join(_ROOT, "pipeline", "office.py")


def J(name, obj):
    p = os.path.join(T, name)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False)
    return p


def cli(*args):
    r = subprocess.run([sys.executable, OFFICE] + [str(a) for a in args],
                       capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout + r.stderr


def zf(path):
    return zipfile.ZipFile(path).namelist()


# ═══════════════════════════ مراجعُ Excel ═══════════════════════════════
print("\n— إزاحةُ المراجع كما يفعل Excel —")
S_ = "Data"
ok("إدراجٌ فوق المرجع ⟵ يُزاح", O.shift_ref("B5", S_, S_, 3, 2) == "B7")
ok("إدراجٌ تحته ⟵ لا يتغيّر", O.shift_ref("B2", S_, S_, 3, 2) == "B2")
ok("مدىً يعبره الإدراج ⟵ يتّسع", O.shift_ref("B2:B5", S_, S_, 3, 1) == "B2:B6")
ok("المطلقُ يُزاح أيضاً ($B$5)", O.shift_ref("$B$5", S_, S_, 3, 1) == "$B$6")
ok("مدىً ينتهي فوق الإدراج ومعادلتُه تحته ⟵ يتّسع (صفٌّ قبل الإجمالي)",
   O.shift_ref("B2:B3", S_, S_, 4, 1, cell_row=4) == "B2:B4")
ok("  ⟵ وبلا خليّةٍ تحته (Excel) ⟵ لا يتّسع",
   O.shift_ref("B2:B3", S_, S_, 4, 1, cell_row=None) == "B2:B3")
ok("  ⟵ ومرجعٌ مفرد لا يتّسع",
   O.shift_ref("B3", S_, S_, 4, 1, cell_row=9) == "B3")
ok("ورقةٌ أخرى ⟵ لا يُمَسّ", O.shift_ref("B5", "Other", S_, 3, 1) == "B5")
ok("مرجعٌ إلى ورقتنا من أخرى ⟵ يُزاح",
   O.shift_ref("Data!B5", "Other", S_, 3, 1) == "Data!B6")
ok("اسمٌ عربيٌّ بين علامتين ⟵ يُزاح",
   O.shift_ref("'المبيعات'!$B$2:$B$3", "x", "المبيعات", 3, 1)
   == "'المبيعات'!$B$2:$B$4")
ok("عمودٌ كامل (A:A) ⟵ لا يتغيّر", O.shift_ref("A:A", S_, S_, 3, 1) == "A:A")
ok("حذفُ صفٍّ عليه المرجع ⟵ #REF!", O.shift_ref("B4", S_, S_, 4, -1) == "#REF!")
ok("حذفٌ داخل مدىً ⟵ يضيق", O.shift_ref("B2:B6", S_, S_, 4, -2) == "B2:B4")
ok("حذفُ المدى كلِّه ⟵ #REF!", O.shift_ref("B4:B5", S_, S_, 4, -2) == "#REF!")
ok("المعادلةُ كاملة (Tokenizer)", O.shift_formula(
    "=SUM(B2:B3)*C4+Data!D9", S_, S_, 3, 1) == "=SUM(B2:B4)*C5+Data!D10")
ok("نصٌّ لا معادلة ⟵ كما هو", O.shift_formula("B5", S_, S_, 1, 1) == "B5")

# ═══════════════════════════ Excel ══════════════════════════════════════
print("\n— Excel: بناءٌ وتعديل —")
xl_spec = {"sheets": [{"name": "المبيعات", "headers": ["البند", "الكمية", "السعر"],
                       "rows": [["أ", 10, 5], ["ب", 20, 7]], "totals": True,
                       "formats": {"C": "#,##0.00"},
                       "chart": {"type": "bar", "data": "B1:B3",
                                 "categories": "A2:A3", "anchor": "F2"}},
                      {"name": "ملخّص", "headers": ["x"], "rows": [["=المبيعات!B4"]]}]}
X = os.path.join(T, "s.xlsx")
c, o = cli("build", J("x.json", xl_spec), "--out", X)
ok("بناء ⟵ ✓ saved", c == 0 and "✓ saved" in o, o[-300:])
from openpyxl import load_workbook   # noqa: E402
wb = load_workbook(X)
ws = wb["المبيعات"]
ok("ورقتان · RTL · مجموعٌ بمعادلة · تنسيقُ أرقام",
   wb.sheetnames == ["المبيعات", "ملخّص"] and ws.sheet_view.rightToLeft
   and ws["B4"].value == "=SUM(B2:B3)" and ws["C2"].number_format == "#,##0.00",
   (wb.sheetnames, ws["B4"].value, ws["C2"].number_format))
ok("رسمٌ أصليٌّ في Excel", any(n.startswith("xl/charts/") for n in zf(X)))
c, o = cli("info", X)
ok("info ⟵ الخلايا بمعادلاتها، ولا تحذيرَ لرسمٍ يحفظه openpyxl",
   c == 0 and "B4 =SUM(B2:B3)" in o and "DROP" not in o, o[:400])

ops = J("xo.json", [{"op": "add_rows", "sheet": "المبيعات", "values": [["ج", 5, 3]]}])
c, o = cli("edit", X, ops)
E = os.path.join(T, "s-edited.xlsx")
ws = load_workbook(E)["المبيعات"]
ok("add_rows ⟵ قبل الإجمالي، والمجموعُ يتّسع", c == 0 and ws["A4"].value == "ج"
   and ws["B5"].value == "=SUM(B2:B4)", (o[-200:], ws["B5"].value))
ok("  ⟵ والصفُّ الجديد بتنسيق ما فوقه",
   ws["A4"].border.left.style == ws["A3"].border.left.style)
ok("  ⟵ ومعادلةُ الورقة الأخرى تُزاح (=المبيعات!B5)",
   load_workbook(E)["ملخّص"]["A2"].value == "=المبيعات!B5",
   load_workbook(E)["ملخّص"]["A2"].value)
_ch = zipfile.ZipFile(E).read("xl/charts/chart1.xml").decode()
ok("  ⟵ والرسمُ يتّسع (قِيس: بقي على B2:B3)", "$B$2:$B$4" in _ch
   and "$A$2:$A$4" in _ch)
ok("  ⟵ والأصلُ لم يُمَسّ", load_workbook(X)["المبيعات"]["B4"].value
   == "=SUM(B2:B3)")

ops = J("xo2.json", [{"op": "insert_rows", "sheet": "المبيعات", "at": 3,
                      "values": [["د", 1, 1]]},
                     {"op": "delete_rows", "sheet": "المبيعات", "at": 2}])
c, o = cli("edit", X, ops, "--out", os.path.join(T, "s2.xlsx"))
ws = load_workbook(os.path.join(T, "s2.xlsx"))["المبيعات"]
ok("إدراجٌ داخل الجدول ثمّ حذف ⟵ المجموعُ صحيح",
   c == 0 and ws["B4"].value == "=SUM(B2:B3)" and ws["A2"].value == "د",
   (o[-200:], ws["B4"].value, ws["A2"].value))

# دمجٌ وجدولٌ واسمٌ معرَّف
from openpyxl import Workbook                 # noqa: E402
from openpyxl.worksheet.table import Table    # noqa: E402
from openpyxl.workbook.defined_name import DefinedName   # noqa: E402
wb = Workbook()
ws = wb.active
ws.title = "D"
ws.append(["name", "value"])
for r in range(2, 6):
    ws.append(["r%d" % r, r])
ws.merge_cells("C4:D5")
ws.add_table(Table(displayName="T1", ref="A1:B5"))
wb.defined_names["Total"] = DefinedName("Total", attr_text="D!$B$2:$B$5")
ws["E1"] = "=SUM(Total)"
M = os.path.join(T, "m.xlsx")
wb.save(M)
c, o = cli("edit", M, J("mo.json", [{"op": "insert_rows", "at": 3, "count": 2}]))
wb2 = load_workbook(os.path.join(T, "m-edited.xlsx"))
ws2 = wb2["D"]
ok("الخلايا المدموجة تُزاح", [str(m) for m in ws2.merged_cells.ranges] == ["C6:D7"],
   [str(m) for m in ws2.merged_cells.ranges])
ok("جدولُ Excel يتّسع", ws2.tables["T1"].ref == "A1:B7", ws2.tables["T1"].ref)
ok("الاسمُ المعرَّف يتّسع", wb2.defined_names["Total"].attr_text == "D!$B$2:$B$7",
   wb2.defined_names["Total"].attr_text)

print("\n— Excel: لا يُكتب ملفٌّ ناقص —")
L = os.path.join(T, "lossy.xlsx")
shutil.copy(X, L)
with zipfile.ZipFile(L, "a") as z:           # جزءٌ لا يعرفه openpyxl فيُسقطه
    z.writestr("xl/threadedComments/threadedComment1.xml", "<x/>")
c, o = cli("info", L)
ok("info يحذّر بما سيضيع (مقيسٌ بحفظٍ تجريبيّ)", "DROP" in o
   and "threaded comments" in o, o[:400])
c, o = cli("edit", L, J("lo.json", [{"op": "set", "cell": "A1", "value": 1}]))
ok("edit ⟵ مرفوض (3) ولا ملفَّ مكتوب", c == 3 and "REFUSED" in o
   and not os.path.exists(os.path.join(T, "lossy-edited.xlsx")), o[-300:])
c, o = cli("edit", L, J("lo.json", [{"op": "set", "cell": "A1", "value": 1}]),
           "--allow-loss")
ok("  ⟵ و--allow-loss يكتبه (قرارُ المستخدم)", c == 0)

print("\n— رفضُ الصيغ القديمة والماكرو —")
for ext in (".xls", ".doc", ".ppt", ".xlsm"):
    p = os.path.join(T, "old" + ext)
    open(p, "wb").write(b"x")
    c, o = cli("info", p)
    ok("%s ⟵ رفضٌ بسببٍ مفهوم" % ext, c == 1 and ("old binary" in o
                                                   or "macros" in o), o)

print("\n— عمليّةٌ فاشلةٌ لا تُسقط غيرَها —")
c, o = cli("edit", X, J("px.json", [{"op": "set", "cell": "A1", "value": "x"},
                                    {"op": "set", "sheet": "لا_توجد", "cell": "A1",
                                     "value": 1}]),
           "--out", os.path.join(T, "p.xlsx"))
ok("جزئيّ ⟵ 2، وما لم يتمّ مذكورٌ بسببه", c == 2 and "no sheet named" in o
   and "✓ saved" in o, o)
c, o = cli("edit", X, J("nx.json", [{"op": "nope"}]), "--out",
           os.path.join(T, "n.xlsx"))
ok("لا شيءَ تمّ ⟵ 1 ولا ملفّ", c == 1 and not os.path.exists(
    os.path.join(T, "n.xlsx")) and "no file written" in o, o)

# ═══════════════════════════ Word ═══════════════════════════════════════
print("\n— Word —")
wspec = {"title": "أثر التعلّم الرقميّ", "header": "Weaver", "sections": [
    {"heading": "المقدّمة", "body": "هذا نصُّ المقدّمة.\n\n- نقطة أولى\n- نقطة ثانية"},
    {"heading": "النتائج", "body": "الجدول:",
     "table": {"headers": ["البند", "القيمة"], "rows": [["أ", 10], ["ب", 20]]}}],
    "references": ["مرجع"]}
if HAVE_MPL:
    wspec["sections"].append({"heading": "رسم", "chart": {
        "type": "bar", "data": {"labels": ["أ", "ب"], "values": [1, 2]}}})
D = os.path.join(T, "r.docx")
c, o = cli("build", J("w.json", wspec), "--out", D)
ok("بناء ⟵ ✓ RTL", c == 0 and "RTL" in o, o[-300:])
from docx import Document   # noqa: E402
d = Document(D)
ok("  ⟵ جدولٌ وعنوانُ المراجع%s" % (" ورسم" if HAVE_MPL else ""),
   len(d.tables) == 1 and any(p.text == "المراجع" for p in d.paragraphs)
   and (not HAVE_MPL or any(n.startswith("word/media/") for n in zf(D))))
c, o = cli("info", D)
idx = {}
for ln in o.splitlines():
    ln = ln.strip()
    if ln.startswith("¶") and "]" in ln:
        idx[ln.split("]", 1)[1].strip()] = int(ln[1:ln.index(" ")])
ok("info ⟵ ¶ وT0 والترويسة", "T0" in o and "[header] Weaver" in o
   and "هذا نصُّ المقدّمة." in idx, o[:500])
i_intro, i_b2 = idx["هذا نصُّ المقدّمة."], idx["نقطة ثانية"]
i_tbl = idx["الجدول:"]
ops = [{"op": "replace", "find": "نصُّ المقدّمة", "with": "نصُّ المقدّمة المعدَّل"},
       {"op": "insert_paragraph", "after": i_intro, "text": ["جديدة ١", "جديدة ٢"]},
       {"op": "insert_paragraph", "after": i_tbl, "text": "عنوانٌ فرعيّ",
        "style": "Heading 1"},
       {"op": "set_cell", "table": 0, "row": 1, "col": 1, "text": "15"},
       {"op": "add_row", "table": 0, "values": ["ج", "7"]},
       {"op": "add_table", "after": i_b2, "headers": ["س"], "rows": [["١"]]},
       {"op": "replace", "find": "Weaver", "with": "Weaver Write"},
       {"op": "delete_paragraph", "paragraph": i_b2},
       {"op": "set_text", "paragraph": i_b2, "text": "x"}]
c, o = cli("edit", D, J("wo.json", ops))
d = Document(os.path.join(T, "r-edited.docx"))
tx = [p.text for p in d.paragraphs]
ok("بديلٌ يحوي المبحوثَ عنه ⟵ مرّةً واحدة (قِيس: «المعدَّل المعدَّل»)",
   "هذا نصُّ المقدّمة المعدَّل." in tx and not any("المعدَّل المعدَّل" in t
                                                  for t in tx), tx[:12])
ok("الفقراتُ في موضعها بأرقام info قبل التعديل (قِيس: id() نقلها)",
   tx.index("جديدة ١") == tx.index("هذا نصُّ المقدّمة المعدَّل.") + 1
   and tx.index("جديدة ٢") == tx.index("جديدة ١") + 1
   and tx.index("نقطة أولى") == tx.index("جديدة ٢") + 1
   and tx.index("عنوانٌ فرعيّ") == tx.index("الجدول:") + 1, tx[:14])
ok("  ⟵ والحذفُ حذفَ المقصودَ وحدَه", "نقطة ثانية" not in tx
   and "نقطة أولى" in tx)
ok("  ⟵ والعمليّةُ على فقرةٍ محذوفةٍ تُرفض بسببها",
   c == 2 and "was deleted by an earlier op" in o, o[-300:])
_hp = [p for p in d.paragraphs if p.text == "عنوانٌ فرعيّ"][0]
_bp = [p for p in d.paragraphs if p.text == "جديدة ١"][0]
ok("الفقرةُ المُدرجةُ بنمطها واتّجاهها (RTL)", _hp.style.name == "Heading 1"
   and _bp._p.pPr is not None and _bp._p.pPr.find(O.W_NS + "bidi") is not None)
body = list(d.element.body.iterchildren())
_t_new = d.tables[0]._tbl       # الجدولُ الجديد أوّلاً في ترتيب المستند
ok("الجدولُ الجديد بعد «نقطة أولى» (مكانِ المحذوفة)", body.index(_t_new)
   == body.index([p for p in d.paragraphs if p.text == "نقطة أولى"][0]._p) + 1
   and d.tables[0].cell(0, 0).text == "س")
old_t = d.tables[1]
ok("set_cell وadd_row", old_t.cell(1, 1).text == "15" and len(old_t.rows) == 4
   and old_t.cell(3, 0).text == "ج", [[c_.text for c_ in r.cells] for r in old_t.rows])
ok("الاستبدالُ يشمل الترويسة", d.sections[0].header.paragraphs[0].text
   == "Weaver Write", d.sections[0].header.paragraphs[0].text)

# نصٌّ مقسومٌ على runs
d0 = Document()
p0 = d0.add_paragraph()
p0.add_run("نصٌّ ").bold = True
p0.add_run("مقسوم هنا")
d0.add_paragraph("كلمةٌ وكلمةٌ")
SP = os.path.join(T, "split.docx")
d0.save(SP)
c, o = cli("edit", SP, J("so.json", [{"op": "replace", "find": "نصٌّ مقسوم",
                                      "with": "جديد"},
                                     {"op": "replace", "find": "كلمةٌ",
                                      "with": "لفظ"}]))
d1 = Document(os.path.join(T, "split-edited.docx"))
ok("نصٌّ مقسومٌ على runs ⟵ يُستبدل", d1.paragraphs[0].text == "جديد هنا",
   d1.paragraphs[0].text)
ok("وتكرارٌ داخل run ⟵ كلُّه ×2", d1.paragraphs[1].text == "لفظ ولفظ"
   and "×2" in o, (d1.paragraphs[1].text, o))

# ═══════════════════════════ PowerPoint ═════════════════════════════════
print("\n— PowerPoint —")
pspec = {"title": "عنوان العرض", "slides": [
    {"title": "المحاور", "points": ["أ", "ب"]},
    {"title": "جدول", "table": {"headers": ["س", "ص"], "rows": [["1", "2"]]}},
    {"layout": "section", "title": "قسم"},
    {"title": "نهاية", "points": ["ج"]}]}
if HAVE_MPL:
    pspec["slides"].insert(2, {"title": "رسم", "chart": {
        "type": "pie", "data": {"labels": ["أ", "ب"], "values": [1, 2]}}})
PP = os.path.join(T, "d.pptx")
c, o = cli("build", J("p.json", pspec), "--out", PP)
from pptx import Presentation   # noqa: E402
prs = Presentation(PP)
titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame
           and sh.text_frame.text.strip()
           # رقمُ الفاصل («١»، «01») ليس عنوانَه
           and not sh.text_frame.text.strip().translate(
               str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")).isdigit()), "")
      for s in prs.slides]
want = ["عنوان العرض", "المحاور", "جدول"] + (["رسم"] if HAVE_MPL else []) + \
    ["قسم", "نهاية", "شكراً لكم"]
ok("بناء ⟵ الشرائحُ بترتيب المواصفة (الجدولُ والرسمُ في مكانهما)",
   c == 0 and titles == want, (titles, o[-200:]))
ok("  ⟵ شريحةُ الجدول جدولٌ أصليّ",
   any(sh.has_table for sh in prs.slides[2].shapes))
n0 = len(prs.slides)
ops = [{"op": "delete_slide", "slide": 3},                   # الجدول
       {"op": "add_slide", "after": 2, "title": "جديدة", "points": ["١", "٢"]},
       {"op": "set_text", "slide": 2, "shape": 2, "text": "◀ س\n◀ ص\n◀ ع"},
       {"op": "move_slide", "slide": n0, "to": 2},
       {"op": "notes", "slide": 1, "text": "ملاحظة"},
       {"op": "replace", "find": "نهاية", "with": "الخاتمة"}]
c, o = cli("edit", PP, J("po.json", ops))
ED = os.path.join(T, "d-edited.pptx")
names = zf(ED)
ok("حذفٌ ثمّ إضافة ⟵ لا اسمَ مكرّراً (قِيس: slide7.xml مرّتين)",
   c == 0 and len(names) == len(set(names)), (o[-300:], len(names) - len(set(names))))
p2 = Presentation(ED)
t2 = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame
           and sh.text_frame.text.strip()
           # رقمُ الفاصل («١»، «01») ليس عنوانَه
           and not sh.text_frame.text.strip().translate(
               str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")).isdigit()), "")
      for s in p2.slides]
ok("  ⟵ الترتيب: نقلٌ وإضافةٌ بعد المقصود، والمحذوفةُ غائبة",
   t2[:4] == ["عنوان العرض", "شكراً لكم", "المحاور", "جديدة"]
   and "جدول" not in t2 and "الخاتمة" in t2 and len(t2) == n0, t2)
_s = p2.slides[2]
ok("  ⟵ set_text بثلاثة أسطر", [p.text for p in list(_s.shapes)[2]
                                 .text_frame.paragraphs] == ["◀ س", "◀ ص", "◀ ع"],
   [p.text for p in list(_s.shapes)[2].text_frame.paragraphs])
ok("  ⟵ ملاحظاتُ المتحدّث", p2.slides[0].notes_slide.notes_text_frame.text
   == "ملاحظة")
# عرضٌ بقوالب (placeholders) ⟵ الشريحةُ الجديدةُ بتخطيطه
pp = Presentation()
s = pp.slides.add_slide(pp.slide_layouts[1])
s.shapes.title.text = "Agenda"
s.placeholders[1].text = "A\nB"
PH = os.path.join(T, "ph.pptx")
pp.save(PH)
c, o = cli("edit", PH, J("ph.json", [{"op": "add_slide", "after": 1,
                                      "title": "Next", "points": ["x", "y"]}]))
p3 = Presentation(os.path.join(T, "ph-edited.pptx"))
ok("عرضٌ بقوالب ⟵ الشريحةُ الجديدةُ بتخطيطه نفسِه",
   "deck layout" in o and p3.slides[1].slide_layout.name == "Title and Content"
   and p3.slides[1].shapes.title.text == "Next", o)

if HAVE_MPL:
    print("\n— الرسوم —")
    CS = J("c.json", {"type": "bar", "title": "المبيعات",
                      "data": {"labels": ["يناير", "فبراير"], "values": [1, 2]},
                      "caption": "شكل", "after": 0})
    c, o = cli("chart", CS, "--out", os.path.join(T, "c.png"))
    ok("صورة PNG", c == 0 and open(os.path.join(T, "c.png"), "rb").read(4)
       == b"\x89PNG", o)
    c, o = cli("chart", CS, "--into", D, "--out", os.path.join(T, "rc.docx"))
    ok("داخل Word بعد ¶0", c == 0 and sum(1 for n in zf(os.path.join(T, "rc.docx"))
                                          if n.startswith("word/media/"))
       > sum(1 for n in zf(D) if n.startswith("word/media/")), o)
    CP = J("cp.json", dict(json.load(open(CS, encoding="utf-8")), after=1))
    c, o = cli("chart", CP, "--into", PP, "--out", os.path.join(T, "pc.pptx"))
    ok("شريحةُ رسمٍ بعد الأولى، بتخطيطٍ فارغ",
       c == 0 and Presentation(os.path.join(T, "pc.pptx")).slides[1]
       .slide_layout.name == "Blank", o)
else:
    print("\n— الرسوم — matplotlib غيرُ مثبَّت هنا: اختباراتُ الرسم تُخطّى")

# ═══════════════════════════ الملفُّ المرفوع ═════════════════════════════
print("\n— الملفُّ المرفوعُ نفسُه في مجلّد المحادثة —")
import server as S                     # noqa: E402
from pipeline import weaver_core as W  # noqa: E402
WS = tempfile.mkdtemp(prefix="weaver-test-ws-")
OUT = tempfile.mkdtemp(prefix="weaver-test-out-")
_real = {"ws": W.workspace_dir, "out": S._output_dir, "ready": S._engine_ready,
         "eng": S._chat_via_engine}
W.workspace_dir = lambda: WS
S._output_dir = lambda: OUT
S._engine_ready = lambda: True
raw = open(D, "rb").read()
b64 = "data:application/octet-stream;base64," + base64.b64encode(raw).decode()
try:
    got = S._attach_save([{"name": "../../تقرير.docx", "data": b64},
                          {"name": "big.txt", "text": "(الملف كبير جداً — x)"},
                          {"name": "n.csv", "text": "a,b\n1,2"}], "chatQ")
    ok("يُحفظ كما هو، باسمٍ بلا مسار", got == ["chats/chatQ/تقرير.docx",
                                              "chats/chatQ/n.csv"]
       and open(os.path.join(WS, "chats/chatQ/تقرير.docx"), "rb").read() == raw, got)
    ok("  ⟵ ورسالةُ «الملفّ كبير» لا تُحفظ ملفّاً",
       not os.path.exists(os.path.join(WS, "chats/chatQ/big.txt")))
    _m = os.path.getmtime(os.path.join(WS, "chats/chatQ/تقرير.docx"))
    time.sleep(0.05)
    S._attach_save([{"name": "تقرير.docx", "data": b64}], "chatQ")
    ok("  ⟵ ومحتوىً مطابق لا يُعاد كتابتُه",
       os.path.getmtime(os.path.join(WS, "chats/chatQ/تقرير.docx")) == _m)
    ok("  ⟵ وبلا محادثة ⟵ لا شيء", S._attach_save([{"name": "a.docx",
                                                     "data": b64}], "") == [])

    seen = {}

    def _fake(m, h=None, t=120, c=None, mem=None, att=None, session=None):
        seen["att"] = att
        if "عدّل" in m:
            p = os.path.join(WS, "chats", session, "تقرير-edited.docx")
            shutil.copy(os.path.join(WS, "chats", session, "تقرير.docx"), p)
        return {"reply": "تمّ", "engine": "weaver-core", "model": "m"}
    S._chat_via_engine = _fake
    srv = S._ReuseTCPServer(("127.0.0.1", 0), S.Handler)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()

    def chat(msg, cid, files):
        rq = urllib.request.Request(
            "http://127.0.0.1:%d/api/chat/stream" % port,
            data=json.dumps({"message": msg, "history": [], "chatId": cid,
                             "files": files}).encode(),
            headers={"Content-Type": "application/json"})
        body = urllib.request.urlopen(rq, timeout=60).read().decode()
        ev = [json.loads(ln[6:]) for ln in body.splitlines()
              if ln.startswith("data: ")]
        out = []
        for e in ev:
            if e.get("t") == "reply" and e.get("output_path"):
                out.append(os.path.basename(e["output_path"]))
            if e.get("t") == "file" and e.get("path"):
                out.append(os.path.basename(e["path"]))
        return out
    try:
        files = chat("لخّص هذا الملفّ", "chatR", [{"name": "تقرير.docx",
                                                   "data": b64}])
        ok("المحادثة ⟵ الملفُّ في مجلّدها", os.path.isfile(
            os.path.join(WS, "chats/chatR/تقرير.docx")))
        ok("  ⟵ والنموذجُ يُخبَر بمساره في أوّل المرفقات",
           (seen.get("att") or "").startswith("[الملفّاتُ المرفوعة نفسُها")
           and "chats/chatR/تقرير.docx" in seen["att"], (seen.get("att") or "")[:200])
        ok("  ⟵ ولا يعود إليك بطاقةً كأنّه جديد", files == [], files)
        files = chat("عدّل هذا الملفّ", "chatR", [{"name": "تقرير.docx",
                                                   "data": b64}])
        ok("عدّله النموذج ⟵ المعدَّلُ وحدَه يُسلَّم",
           files == ["تقرير-edited.docx"], files)
        chat("Summarise this file", "chatE", [{"name": "a.docx", "data": b64}])
        ok("رسالةٌ إنجليزيّة ⟵ السطرُ بالإنجليزيّة",
           (seen.get("att") or "").startswith("[The uploaded files themselves"))
    finally:
        srv.shutdown()
    ok("ومسوّداتُ المهارات لا تُسلَّم (office-spec.json…)",
       {"office-spec.json", "office-ops.json", "chart-spec.json"}
       <= S._WS_INTERNAL)
finally:
    W.workspace_dir = _real["ws"]
    S._output_dir = _real["out"]
    S._engine_ready = _real["ready"]
    S._chat_via_engine = _real["eng"]

# ═══════════════════════════ المهارات ═══════════════════════════════════
print("\n— المهارات —")
import re   # noqa: E402
_src = open(OFFICE, encoding="utf-8").read()
for n in ("office-word", "office-powerpoint", "office-excel", "office-charts"):
    t = open(os.path.join(W.SKILLS_SRC, n, "SKILL.md"), encoding="utf-8").read()
    ok("%s: تنادي {{WEAVER}}/pipeline/office.py" % n,
       "{{WEAVER}}/pipeline/office.py" in t)
    cmds = set(re.findall(r"office\.py (\w+)", t))
    ok("  ⟵ أوامرُها موجودة (%s)" % ", ".join(sorted(cmds)),
       cmds and cmds <= {"info", "build", "chart", "edit", "fonts"}, cmds)
    ops_ = set(re.findall(r'"op":\s*"(\w+)"', t))
    miss = [x for x in ops_ if '"%s"' % x not in _src]
    ok("  ⟵ وكلُّ عمليّةٍ تذكرها تعرفها الأداة", not miss, miss)
    ok("  ⟵ ولا «تمّ» قبل ✓", "✓" in t and "قبل أن تقول «تمّ»" in t)
    ok("  ⟵ ومتى لا تُستدعى", "## متى لا" in t)
    ok("  ⟵ وطلبٌ طبيعيٌّ لاختبار الاستدعاء (--skills invoke)",
       n in W._SKILL_PROBES and n not in W._SKILL_PROBES[n])
ok("والمعاكس: سؤالٌ نصّيٌّ ⟵ لا ملفّ (--not)", "office-word" in W._SKILL_NEGATIVE)

_rw = {k: getattr(W, k) for k in ("ask", "trajectory")}
try:
    W.ask = lambda *a, **k: {"answer": "تمّ"}
    for tgt, cmd in (("office-word", "python3 /x/pipeline/office.py build "
                                     "chats/a/office-spec.json --out chats/a/r.docx"),
                     ("office-excel", "python3 /x/pipeline/office.py edit m.xlsx o.json"),
                     ("office-charts", "python3 /x/pipeline/office.py chart c.json "
                                       "--out c.png")):
        W.trajectory = lambda sk, _c=cmd: [{"name": "exec", "request": _c}]
        r = W.skills_invoke_test(target=tgt)
        ok("استعمالُ الأداة بلا فتح SKILL.md ⟵ يُحسب لـ%s" % tgt, r["ok"]
           and tgt in r["called"], r["called"])
finally:
    for k, v in _rw.items():
        setattr(W, k, v)

shutil.rmtree(T, ignore_errors=True)
shutil.rmtree(WS, ignore_errors=True)
shutil.rmtree(OUT, ignore_errors=True)
print("\n" + "=" * 62)
print(f" RESULT: {'PASS' if F == 0 else 'FAIL'}   ({P}/{P + F})")
print("=" * 62)
sys.exit(1 if F else 0)
