# -*- coding: utf-8 -*-
"""العربيُّ من اليمين والإنجليزيُّ من اليسار — في Word وPowerPoint والرسوم —
وبالخطّ الذي طلبه المستخدم.

شكا المستخدم: «الرسوم البيانية الكتابة العربية تطلع مقطوع ومن اليسار، وحتى
الكتابة العربية في الملف تطلع من إتجاه اليسار». وقِيس (LibreOffice يحاكي Word):
  · Word: <w:bidi/> مع jc=right ⟵ يسار (Word يقرأ right «نهاية» في فقرةٍ من
    اليمين). بلا jc ⟵ يمين. في الملفّ القديم ١١ فقرةً عربيّةً من ٢١ يساراً.
  · PowerPoint: الأعمدةُ معكوسةٌ في الملفّ ومعها tblPr rtl="1" — LibreOffice
    يتجاهل العلَم، وPowerPoint يطبّقه فيقلبها ثانيةً.
  · الرسوم: الخطُّ يُضبط بعد إنشاء المحاور فيأخذ كلُّ رسمٍ خطَّ الرسم السابق
    (والرسمُ الوحيدُ DejaVu)؛ ومسارُ fonts-core ينقصه مستوى؛ والفئةُ الأولى يساراً.
  · الخطوط: «Kufyan Arabic» اسمٌ لا يطابق ملفَّه («Kufyan Arabic Regular»)،
    وملفٌّ لاتينيٌّ فارغ (IBMPlexSerif-Regular.ttf، ٠ بايت).
"""
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
for _p in ("capabilities/skills/docx_builder/scripts",
           "capabilities/skills/pptx_builder/scripts",
           "capabilities/skills/chart_builder/scripts", "engines/fonts-core"):
    sys.path.insert(0, os.path.join(_ROOT, _p))

import docx_rtl as X          # noqa: E402
import font_catalog as FC     # noqa: E402
from docx import Document     # noqa: E402

P = F = 0
W = X.W


def ok(name, cond, extra=""):
    global P, F
    if cond:
        P += 1
        print(f"  ✓ {name}")
    else:
        F += 1
        print(f"  ✗ {name}" + (f"   — {extra}" if extra else ""))


T = tempfile.mkdtemp(prefix="weaver-test-rtl-")
OFFICE = os.path.join(_ROOT, "pipeline", "office.py")


def cli(*args):
    r = subprocess.run([sys.executable, OFFICE] + [str(a) for a in args],
                       capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout + r.stderr


def J(name, obj):
    p = os.path.join(T, name)
    json.dump(obj, open(p, "w", encoding="utf-8"), ensure_ascii=False)
    return p


def pinfo(p):
    pPr = p.find(W + "pPr")
    bidi = pPr is not None and pPr.find(W + "bidi") is not None
    jc = pPr.find(W + "jc") if pPr is not None else None
    return bidi, (jc.get(W + "val") if jc is not None else None)


def left_arabic(path):
    """فقراتٌ عربيّةٌ تظهر يساراً: بلا bidi، أو bidi مع jc=right/end."""
    bad, n = 0, 0
    for p in Document(path).element.body.iter(W + "p"):
        if not X.has_ar(X._p_text(p)):
            continue
        n += 1
        bidi, jc = pinfo(p)
        if not bidi or jc in ("right", "end"):
            bad += 1
    return n, bad


# ═══════════════════════ فهرسُ الخطوط ═══════════════════════
print("\n— فهرسُ الخطوط: الاسمُ الحقيقيّ من الملفّ —")
cat = {e["family"]: e for e in FC.catalog()}
ok("Kufyan باسمه في ملفّه «Kufyan Arabic Regular»", "Kufyan Arabic Regular" in cat
   and "Kufyan Arabic" not in cat)
ok("الملفُّ الفارغ لا يكسر الفهرس (IBMPlexSerif-Regular.ttf ٠ بايت)",
   os.path.getsize(os.path.join(_ROOT, "engines/fonts-core/latin/IBMPlexSerif-Regular.ttf"))
   == 0 and "IBM Plex Serif" in cat and os.path.getsize(
       cat["IBM Plex Serif"]["files"]["regular"]) > 0)
ok("الخطوطُ العربيّةُ الستّ", {"Amiri", "Cairo", "Tajawal", "Tajawal Black",
                               "Noto Naskh Arabic", "Kufyan Arabic Regular"}
   <= set(cat), sorted(cat)[:8])
for q, fam, w in (("أميري", "Amiri", "regular"), ("خط القاهرة عريض", "Cairo", "bold"),
                  ("Tajawal Black", "Tajawal Black", "regular"),
                  ("tajawal bold", "Tajawal", "bold"), ("كوفيان", "Kufyan Arabic Regular",
                                                         "regular"),
                  ("plex serif", "IBM Plex Serif", "regular")):
    e = FC.find_font(q)
    ok("«%s» ⟵ %s (%s)" % (q, fam, w), e["bundled"] and e["family"] == fam
       and e["weight"] == w, (e["family"], e["weight"]))
e = FC.find_font("Arial")
ok("Arial ⟵ باسمه، غيرُ مضمَّن، وأقربُ خطٍّ للرسم", not e["bundled"]
   and e["family"] == "Arial" and e["stand_in"] == "Cairo")
e = FC.find_font("العربي المبسط")
ok("«العربي المبسط» ⟵ Simplified Arabic", e["family"] == "Simplified Arabic")
e = FC.find_font("خطٌّ لا يوجد")
ok("اسمٌ مجهول ⟵ يُكتب كما طُلب ويُقال ما عندنا", not e["bundled"]
   and "Amiri" in e["note"])
ok("ملفُّ الوزن: bold ⟵ Amiri-Bold", FC.file_for(FC.find_font("أميري عريض"))
   .endswith("Amiri-Bold.ttf"))

# ═══════════════════════ اتّجاهُ Word ═══════════════════════
print("\n— Word: قاعدةُ المحاذاة المقيسة —")
d = Document()
from docx.enum.text import WD_ALIGN_PARAGRAPH   # noqa: E402
p1 = d.add_paragraph("فقرةٌ عربيّة")
X._put(p1._p.get_or_add_pPr(), X._el("bidi"), X._PPR)
p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p2 = d.add_paragraph("English reference in an Arabic doc")
X._put(p2._p.get_or_add_pPr(), X._el("bidi"), X._PPR)
p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p3 = d.add_paragraph("عنوان")
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p1.runs[0]
r.font.size = 177800
r.bold = True
X.finalize_direction(d, "ar")
ok("عربيٌّ + jc=right ⟵ bidi بلا jc (يمين)", pinfo(p1._p) == (True, None), pinfo(p1._p))
ok("لاتينيٌّ في مستندٍ عربيّ ⟵ من اليسار", pinfo(p2._p) == (False, None), pinfo(p2._p))
ok("الوسطُ يبقى وسطاً", pinfo(p3._p) == (True, "center"), pinfo(p3._p))
rPr = r._r.rPr
ok("حجمُ العربيّ وعرضُه (szCs · bCs) ⟵ Word يقرؤهما للعربيّ",
   rPr.find(W + "szCs") is not None and rPr.find(W + "bCs") is not None
   and rPr.find(W + "szCs").get(W + "val") == rPr.find(W + "sz").get(W + "val"))
ok("<w:rtl/> على الـrun العربيّ", rPr.find(W + "rtl") is not None)
names = [X._local(c) for c in rPr]
ok("وبترتيب المخطّط (b، bCs … sz، szCs … rtl)",
   names.index("b") < names.index("bCs") < names.index("sz") < names.index("szCs")
   < names.index("rtl"), names)
pp = [X._local(c) for c in p3._p.pPr]
ok("<w:bidi/> قبل <w:jc/> (كان يُلحق في آخر pPr)", pp.index("bidi") < pp.index("jc"), pp)
ok("القسمُ من اليمين (sectPr/bidi)", d.element.body.find(W + "sectPr")
   .find(W + "bidi") is not None)
d2 = Document()
q = d2.add_paragraph("اقتباسٌ عربيٌّ في مستندٍ إنجليزيّ")
en = d2.add_paragraph("English body")
X.finalize_direction(d2, "en")
ok("مستندٌ إنجليزيّ: العربيُّ فيه من اليمين، والإنجليزيُّ كما هو",
   pinfo(q._p)[0] and not pinfo(en._p)[0]
   and d2.element.body.find(W + "sectPr").find(W + "bidi") is None)

print("\n— Word: تضمينُ الخطّ —")
D = os.path.join(T, "emb.docx")
d3 = Document()
d3.add_paragraph("نصٌّ بخطِّ أميري")
X.apply_fonts(d3, cs="Amiri", latin="Amiri")
d3.save(D)
am = FC.find_font("أميري")
got = X.embed_fonts(D, [(am["family"], am["files"])])
z = zipfile.ZipFile(D)
parts = [n for n in z.namelist() if n.startswith("word/fonts/")]
ft = z.read("word/fontTable.xml").decode()
ok("ضُمِّن Amiri (عاديٌّ وعريض)", got == ["Amiri"] and len(parts) == 2, (got, parts))
ok("في جدول الخطوط: embedRegular وembedBold بمفتاح", 'w:name="Amiri"' in ft
   and "embedRegular" in ft and "embedBold" in ft and "w:fontKey=" in ft)
import re   # noqa: E402
key = re.search(r'embedRegular[^>]*w:fontKey="(\{[^"]+\})"', ft) or \
    re.search(r'w:fontKey="(\{[^"]+\})"[^>]*/>', ft)
rid = re.search(r'embedRegular r:id="([^"]+)"', ft).group(1)
rels = z.read("word/_rels/fontTable.xml.rels").decode()
target = re.search(r'Id="%s"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="%s"'
                   % (rid, rid), rels)
tgt = "word/" + (target.group(1) or target.group(2))
fk = re.search(r'embedRegular r:id="%s" w:fontKey="(\{[^"]+\})"' % rid, ft).group(1)
raw = open(am["files"]["regular"], "rb").read()
ok("فكُّ التعمية يعيد الخطَّ نفسَه بايتاً ببايت (ECMA-376 §17.8.1)",
   X.obfuscate(z.read(tgt), fk) == raw)
ok("  ⟵ والمعمّى غيرُ الأصل (أوّلُ ٣٢ بايتاً)", z.read(tgt)[:32] != raw[:32]
   and z.read(tgt)[32:] == raw[32:])
ok("نوعُ المحتوى odttf · وembedTrueTypeFonts في الإعدادات",
   "odttf" in z.read("[Content_Types].xml").decode()
   and "embedTrueTypeFonts" in z.read("word/settings.xml").decode())
ok("تضمينٌ ثانٍ بالاسم نفسِه ⟵ لا يُكرَّر", X.embed_fonts(
    D, [(am["family"], am["files"])]) == [] and len(
    [n for n in zipfile.ZipFile(D).namelist() if n.startswith("word/fonts/")]) == 2)
ok("والملفُّ يُفتح", len(Document(D).paragraphs) == 1)

# ═══════════════════════ عبر الأداة ═══════════════════════
print("\n— الأداة: Word بالخطّ المطلوب ومن اليمين —")
WS = J("w.json", {"title": "أثر التعلّم", "font": "أميري", "toc": True, "header": "ترويسة",
                  "sections": [{"heading": "المقدّمة", "body": "نصّ.\n\n- نقطة\n- أخرى"},
                               {"heading": "النتائج", "body": "جدول:",
                                "table": {"headers": ["البند", "القيمة"],
                                          "rows": [["أ", 1]]}}],
                  "references": ["مرجعٌ عربيّ", "Smith, J. (2020). English."]})
WD = os.path.join(T, "w.docx")
c, o = cli("build", WS, "--out", WD)
ok("بناء ⟵ «font Amiri (embedded …)»", c == 0 and "font Amiri (embedded" in o, o)
n, bad = left_arabic(WD)
ok("لا فقرةَ عربيّةً يساراً (كانت ١١ من ٢١)", n > 5 and bad == 0, (n, bad))
dd = Document(WD)
rf = [r._r.rPr.find(W + "rFonts") for p in dd.paragraphs for r in p.runs
      if r._r.rPr is not None and r.text.strip()]
ok("كلُّ نصٍّ بـAmiri (cs للعربيّ وascii للاتينيّ)",
   rf and all(x is not None and x.get(W + "cs") == "Amiri"
              and x.get(W + "ascii") == "Amiri" for x in rf))
ok("المرجعُ الإنجليزيُّ من اليسار", any(
    p.text.startswith("Smith") and not pinfo(p._p)[0] for p in dd.paragraphs))
c, o = cli("build", J("w2.json", {"title": "بلا خطّ", "sections": [
    {"heading": "أ", "body": "نصّ"}]}), "--out", os.path.join(T, "w2.docx"))
ok("بلا طلب ⟵ Kufyan باسمه الحقيقيّ مضمَّناً", "font Kufyan Arabic Regular (embedded"
   in o and 'w:name="Kufyan Arabic Regular"' in zipfile.ZipFile(
       os.path.join(T, "w2.docx")).read("word/fontTable.xml").decode(), o)
c, o = cli("build", J("w3.json", {"title": "Report", "font": "Arial", "sections": [
    {"heading": "Intro", "body": "Text"}]}), "--out", os.path.join(T, "w3.docx"))
ok("Arial ⟵ باسمه ويُقال إنّه غيرُ مضمَّن", "font Arial (by name only" in o
   and not any(n.startswith("word/fonts/") for n in zipfile.ZipFile(
       os.path.join(T, "w3.docx")).namelist()), o)

print("\n— الأداة: تعديلُ Word —")
# ملفٌّ قديمٌ كما كانت تبنيه الأدوات: bidi + jc=right
old = Document()
for t in ("عنوانٌ قديم", "فقرةٌ قديمة"):
    pp_ = old.add_paragraph(t)
    X._put(pp_._p.get_or_add_pPr(), X._el("bidi"), X._PPR)
    pp_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
OLD = os.path.join(T, "old.docx")
old.save(OLD)
ok("الملفُّ القديم: فقرتاه يساراً", left_arabic(OLD) == (2, 2))
c, o = cli("edit", OLD, J("o1.json", [{"op": "insert_paragraph", "after": 1,
                                       "text": "فقرةٌ جديدة"}]),
           "--out", os.path.join(T, "o1.docx"))
dd = Document(os.path.join(T, "o1.docx"))
ok("الفقرةُ المُدرجةُ من اليمين (كانت الأداةُ تكتب RIGHT)", pinfo(
    dd.paragraphs[2]._p) == (True, None), pinfo(dd.paragraphs[2]._p))
ok("  ⟵ وملفُّ المستخدم لا يُمَسّ إلا بطلب", left_arabic(os.path.join(T, "o1.docx"))
   == (3, 2))
c, o = cli("edit", OLD, J("o2.json", [{"op": "fix_direction"},
                                      {"op": "set_font", "font": "تجوال"}]),
           "--out", os.path.join(T, "o2.docx"))
ok("fix_direction ⟵ الملفُّ كلُّه من اليمين", c == 0 and left_arabic(
    os.path.join(T, "o2.docx")) == (2, 0), o)
ok("set_font تجوال ⟵ مضمَّنٌ في الملفّ", "font → font Tajawal (embedded" in o and
   'w:name="Tajawal"' in zipfile.ZipFile(os.path.join(T, "o2.docx")).read(
       "word/fontTable.xml").decode(), o)

# ═══════════════════════ PowerPoint ═══════════════════════
print("\n— PowerPoint —")
PS = J("p.json", {"title": "عرض", "font": "Cairo", "slides": [
    {"title": "جدول", "table": {"headers": ["الأوّل", "الثاني", "الثالث"],
                                "rows": [["١", "٢", "٣"]]}},
    {"title": "محاور", "points": ["نقطة", "Point in English"]}]})
PP = os.path.join(T, "p.pptx")
c, o = cli("build", PS, "--out", PP)
ok("بناء ⟵ «font Cairo (by name…)» — العرضُ لا يُضمِّن", c == 0
   and "font Cairo (by name" in o, o)
from pptx import Presentation   # noqa: E402
from pptx.oxml.ns import qn     # noqa: E402
prs = Presentation(PP)
tb = next(sh.table for s in prs.slides for sh in s.shapes if sh.has_table)
ok("الجدول: العمودُ الأوّلُ في آخر الملفّ (يُعرض يميناً) وعلَمُ rtl=\"0\"",
   tb.cell(0, 2).text == "الأوّل" and tb._tbl.find(qn("a:tblPr")).get("rtl") == "0",
   (tb.cell(0, 0).text, tb._tbl.find(qn("a:tblPr")).get("rtl")))
runs = [r for s in prs.slides for sh in s.shapes if sh.has_text_frame
        for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]
runs += [r for c_ in tb.iter_cells() for p in c_.text_frame.paragraphs
         for r in p.runs if r.text.strip()]
ok("كلُّ نصٍّ بـCairo (a:cs وa:latin — خلايا الجدول أيضاً)",
   runs and all(r._r.rPr.find(qn("a:cs")) is not None
                and r._r.rPr.find(qn("a:cs")).get("typeface") == "Cairo"
                and r._r.rPr.find(qn("a:latin")).get("typeface") == "Cairo"
                for r in runs), len(runs))
ar_p = [p for s in prs.slides for sh in s.shapes if sh.has_text_frame
        for p in sh.text_frame.paragraphs if X.has_ar(p.text)]
ok("الفقراتُ العربيّة: rtl=1 ومحاذاةٌ يمين (والختامُ في الوسط عمداً)", ar_p and all(
    p._p.pPr.get("rtl") == "1" and p._p.pPr.get("algn") in ("r", "ctr")
    for p in ar_p), [(p.text[:12], dict(p._p.pPr.attrib)) for p in ar_p])
en_p = [p for s in prs.slides for sh in s.shapes if sh.has_text_frame
        for p in sh.text_frame.paragraphs if "English" in p.text]
ok("والإنجليزيّةُ في العرض نفسِه من اليسار (كانت تُفرض يميناً)",
   en_p and en_p[0]._p.pPr.get("rtl") == "0" and en_p[0]._p.pPr.get("algn") == "l",
   en_p and en_p[0]._p.pPr.attrib)
c, o = cli("edit", PP, J("po.json", [{"op": "set_font", "font": "أميري"}]))
p2 = Presentation(os.path.join(T, "p-edited.pptx"))
ok("set_font أميري", c == 0 and all(
    r._r.rPr.find(qn("a:cs")).get("typeface") == "Amiri"
    for s in p2.slides for sh in s.shapes if sh.has_text_frame
    for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()), o)

# ═══════════════════════ Excel ═══════════════════════
print("\n— Excel —")
XS = J("x.json", {"font": "Tajawal", "sheets": [{"name": "ورقة", "headers": ["البند", "القيمة"],
                                                 "rows": [["أ", 1]], "totals": True}]})
XX = os.path.join(T, "x.xlsx")
c, o = cli("build", XS, "--out", XX)
from openpyxl import load_workbook   # noqa: E402
ws = load_workbook(XX).active
ok("الخلايا بـTajawal · والورقةُ من اليمين", c == 0 and ws["A1"].font.name == "Tajawal"
   and ws["A1"].font.bold and ws.sheet_view.rightToLeft, o)
c, o = cli("edit", XX, J("xo.json", [{"op": "format", "range": "B2", "font": "أميري"}]))
ok("format بخطّ ⟵ Amiri", c == 0 and load_workbook(os.path.join(T, "x-edited.xlsx"))
   .active["B2"].font.name == "Amiri", o)

# ═══════════════════════ الرسوم ═══════════════════════
try:
    import matplotlib   # noqa: F401
    HAVE_MPL = True
except ImportError:
    HAVE_MPL = False
if HAVE_MPL:
    print("\n— الرسوم —")
    import build_chart as bc   # noqa: E402
    ok("مسارُ fonts-core صحيح (كان ينقصه مستوى)", os.path.isdir(bc._FONTS_DIR))
    mode = bc._shaping_mode(None)
    conv = (lambda t: t) if mode == "native" else None
    if conv is None:
        import arabic_reshaper
        from bidi.algorithm import get_display
        conv = lambda t: get_display(arabic_reshaper.reshape(t))
    ok("الطريقُ المقيس يرسم العربيَّ متّصلاً وبترتيبه (%s)" % mode,
       bc._draws_right(conv))
    other = (lambda t: t) if mode != "native" else None
    if other is None:
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display
            other = lambda t: get_display(arabic_reshaper.reshape(t))
        except ImportError:
            other = None
    if other is not None:
        ok("  ⟵ والطريقُ الآخرُ يُكشف خطؤه (مقطّعٌ أو مقلوب)", not bc._draws_right(other))
    import matplotlib.pyplot as plt
    import matplotlib.axes as mx
    from matplotlib import font_manager as fm
    seen = []
    _orig_t = mx.Axes.set_title

    def _spy(self, label, *a, **k):
        t = _orig_t(self, label, *a, **k)
        seen.append(os.path.basename(fm.findfont(t.get_fontproperties())))
        return t
    mx.Axes.set_title = _spy
    inv = []
    _oi, _oy = mx.Axes.invert_xaxis, mx.Axes.invert_yaxis
    mx.Axes.invert_xaxis = lambda self: (inv.append("x"), _oi(self))[1]
    mx.Axes.invert_yaxis = lambda self: (inv.append("y"), _oy(self))[1]
    try:
        res = []
        for f in ("أميري", None, "Cairo", "Arial"):
            res.append(bc.build_chart("bar", {"labels": ["أ", "ب"], "values": [1, 2]},
                                      os.path.join(T, "c.png"), title="عنوان",
                                      lang="ar", font=f))
        ok("كلُّ رسمٍ بخطّه لا بخطّ الذي قبله (كان يتأخّر رسماً)",
           seen == ["Amiri-Bold.ttf", "Kufyan-Arabic-Regular.ttf", "Cairo-Variable.ttf",
                    "Cairo-Variable.ttf"], seen)
        ok("Arial ⟵ Cairo ويُقال ذلك", "chart drawn with 'Cairo'" in res[3]["font_note"])
        inv.clear()
        bc.build_chart("bar", {"labels": ["أ", "ب"], "values": [1, 2]},
                       os.path.join(T, "c.png"), lang="ar")
        ok("أعمدةٌ عربيّة ⟵ الفئةُ الأولى يميناً", inv == ["x"], inv)
        inv.clear()
        bc.build_chart("horizontal_bar", {"labels": ["أ", "ب"], "values": [1, 2]},
                       os.path.join(T, "c.png"), lang="ar")
        ok("أعمدةٌ أفقيّة ⟵ من اليمين، والأولى أعلى", sorted(inv) == ["x", "y"], inv)
        inv.clear()
        bc.build_chart("bar", {"labels": ["A", "B"], "values": [1, 2]},
                       os.path.join(T, "c.png"), lang="en")
        bc.build_chart("scatter", {"x": [1, 2], "y": [2, 1]},
                       os.path.join(T, "c.png"), lang="ar")
        ok("الإنجليزيُّ والمحاورُ الرقميّة كما هي", inv == [], inv)
    finally:
        mx.Axes.set_title = _orig_t
        mx.Axes.invert_xaxis, mx.Axes.invert_yaxis = _oi, _oy
    c, o = cli("fonts", "القاهرة")
    ok("office.py fonts القاهرة ⟵ Cairo", c == 0 and "→ Cairo" in o, o)
    c, o = cli("fonts")
    ok("office.py fonts ⟵ الفهرس", c == 0 and "Amiri" in o and "Work Sans" in o)
else:
    print("\n— الرسوم — matplotlib غيرُ مثبَّت هنا: تُخطّى")

shutil.rmtree(T, ignore_errors=True)
print("\n" + "=" * 62)
print(f" RESULT: {'PASS' if F == 0 else 'FAIL'}   ({P}/{P + F})")
print("=" * 62)
sys.exit(1 if F else 0)
