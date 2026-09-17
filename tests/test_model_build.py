# -*- coding: utf-8 -*-
"""النموذج يبني الملفّ بنفسه — وأثرُ ما يُنفَّذ محفوظٌ دائماً.

الفكرة مقيسةٌ على حزمة openclaw: لا مهارةَ باوربوينت ولا إكسل ولا مراجع
عندها — صفرٌ بحدود الكلمة — وتُنتجها كلَّها، لأن النموذج يكتب السكربت ويشغّله
ويقرأ خطأه فيصلحه. هذه الاختبارات تثبت أن الآلية نفسها تعمل هنا، وأن حدودها
قائمة: مطفأةٌ حتى يأذن المستخدم، والسكربتُ يُكتب قبل أن يُشغَّل، والفشلُ لا
يمسّ المخرَج القائم."""
import sys, os, json, tempfile
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)
for k in list(os.environ):
    if k.startswith("WEAVER_"):
        os.environ.pop(k)
from pipeline.model_build import (build_with_model, build_payload,
                                  build_prompt, script_path_for,
                                  _strip_code_fence)
ok = True
_TMP = tempfile.mkdtemp(prefix="weaver-mb-test-")


def chk(label, good, detail=""):
    global ok
    ok &= bool(good)
    print(f"   {'✅' if good else '❌'} {label}" + (f" — {detail}" if detail else ""))


SECS = [{"heading": "المبحث الأول: أركان عقد الزواج", "body": "ميثاقٌ غليظ…",
         "level": 1},
        {"heading": "المطلب الأول: اشتراط الولي", "body": "يدور التساؤل…",
         "level": 2}]
CARD = {"topic": "الزواج في الإسلام", "citation_style": "APA"}
REQ = ("أريد بحثاً بالكامل عن الزواج في الإسلام… ويكون الترتيب بهذا الشكل: "
       "المبحث الأول: … المطلب الأول: … 1.1 … 1.2 …")

GOOD = """
import os, json
from docx import Document
d = json.load(open(os.environ["WEAVER_PAYLOAD"], encoding="utf-8"))
doc = Document()
doc.add_heading(d["title"], 0)
for s in d["sections"]:
    doc.add_heading(s["heading"], level=s["level"])
    if s.get("body"):
        doc.add_paragraph(s["body"])
doc.save(os.environ["WEAVER_OUT"])
"""
BAD = 'raise RuntimeError("AttributeError: Document has no add_heading2")'


class _Model:
    """نموذجٌ يفشل مرّتين ثم يصلح — كما يحدث فعلاً مع نموذجٍ متوسّط."""

    def __init__(self, succeed_at=3):
        self.n = 0
        self.saw_error = []
        self.saw_request = []
        self.succeed_at = succeed_at

    def __call__(self, prompt, **kw):
        self.n += 1
        self.saw_error.append("السبب" in prompt)
        self.saw_request.append("المبحث الأول" in prompt)
        return GOOD if self.n >= self.succeed_at else BAD


print("═" * 70)
print(" ١) طلبُ المستخدم يصل النموذجَ حرفياً")
print("═" * 70)
p = build_prompt(REQ, "docx", "ar")
chk("نصُّ الطلب داخل الموجّه", "المبحث الأول: … المطلب الأول" in p)
chk("ومعلنٌ أنه فوق كلّ ما دونه", "فوق كلّ ما دونه" in p)
chk("والمسارات مشروحة", "WEAVER_PAYLOAD" in p and "WEAVER_OUT" in p)
chk("والممنوعات مذكورة", "subprocess" in p and "os.system" in p)
p2 = build_prompt(REQ, "pptx", "en")
chk("وبالإنجليزية أيضاً", "outranks everything below" in p2 and ".pptx" in p2)
p3 = build_prompt(REQ, "docx", "ar", error="AttributeError: kaboom")
chk("والخطأ يعود إليه ليصلحه", "kaboom" in p3 and "أصلحه" in p3)

print("\n" + "═" * 70)
print(" ٢) المحتوى يُسلَّم JSON لا لبس فيه")
print("═" * 70)
d = json.loads(build_payload(SECS, CARD, "ar"))
chk("العنوان والأقسام", d["title"] == "الزواج في الإسلام"
    and len(d["sections"]) == 2)
chk("والمستويات محفوظة", [s["level"] for s in d["sections"]] == [1, 2])
chk("والعربية سليمة", "المبحث الأول" in d["sections"][0]["heading"])
chk("ولا يسقط نمطُ التوثيق", d["citation_style"] == "APA")

print("\n" + "═" * 70)
print(" ٣) مطفأٌ حتى يأذن المستخدم — والسكربت يُحفظ على كلّ حال")
print("═" * 70)
out = os.path.join(_TMP, "doc.docx")
for f in (out, script_path_for(out)):
    if os.path.exists(f):
        os.remove(f)
m = _Model(succeed_at=1)
r = build_with_model(m, SECS, CARD, "ar", out, REQ)
chk("لم يُشغَّل شيء", r["ran"] is False and r["ok"] is False)
chk("والسبب معلن", "مطفأ" in r["reason"], r["reason"][:46])
chk("والسكربت محفوظٌ ليقرأه المستخدم",
    bool(r["script"]) and os.path.exists(r["script"]),
    os.path.basename(r["script"] or ""))
chk("ولا ملفَّ مخرَجٍ أُنشئ", not os.path.exists(out))
chk("واسمُ السكربت مقروء", os.path.basename(script_path_for(out)) ==
    "build_doc.py")

print("\n" + "═" * 70)
print(" ٤) وبإذنه: يبني، ويصلح خطأه، ويُنتج ملفاً حقيقياً")
print("═" * 70)
os.environ["WEAVER_EXEC"] = "1"
m = _Model(succeed_at=3)
r = build_with_model(m, SECS, CARD, "ar", out, REQ)
chk("نجح بعد إصلاحين", r["ok"] and r["attempts"] == 3, str(r["attempts"]))
chk("والخطأ عاد إليه في المحاولتين التاليتين",
    m.saw_error == [False, True, True], str(m.saw_error))
chk("وطلبُ المستخدم كان معه في كلّ محاولة", all(m.saw_request))
chk("والملفُّ حقيقيّ",
    os.path.exists(out) and os.path.getsize(out) > 20000,
    f"{os.path.getsize(out) if os.path.exists(out) else 0} بايت")
if os.path.exists(out):
    import zipfile
    x = zipfile.ZipFile(out).read("word/document.xml").decode()
    chk("وفيه عناوين المستند بالعربية", "المبحث الأول" in x)

print("\n" + "═" * 70)
print(" ٥) والفشلُ لا يضرّ")
print("═" * 70)


class _Hopeless:
    def __call__(self, prompt, **kw):
        return BAD


out2 = os.path.join(_TMP, "fail.docx")
r = build_with_model(_Hopeless(), SECS, CARD, "ar", out2, REQ)
chk("ثلاثُ محاولاتٍ ثم استسلام", not r["ok"] and r["attempts"] == 3)
chk("ولا ملفَّ نصفَ مبنيٍّ يُترك", not os.path.exists(out2))
chk("والسببُ محفوظٌ للمستخدم", bool(r["reason"]), r["reason"][:44])


class _Empty:
    def __call__(self, prompt, **kw):
        return "   "


r = build_with_model(_Empty(), SECS, CARD, "ar",
                     os.path.join(_TMP, "e.docx"), REQ)
chk("وردٌّ فارغ ⟶ لا انهيار", not r["ok"])


class _Broken:
    def __call__(self, prompt, **kw):
        raise ConnectionError("403")


r = build_with_model(_Broken(), SECS, CARD, "ar",
                     os.path.join(_TMP, "x.docx"), REQ)
chk("ونموذجٌ يرفض ⟶ سببٌ لا انهيار", not r["ok"] and "403" in r["reason"])
chk("وبلا محتوىً ⟶ لا نداء",
    build_with_model(_Model(), [], CARD, "ar",
                     os.path.join(_TMP, "n.docx"), REQ)["reason"] == "لا محتوى")

print("\n" + "═" * 70)
print(" ٦) وسياجُ الماركداون لا يكسر الكود")
print("═" * 70)
chk("```python … ```", _strip_code_fence("```python\nx = 1\n```") == "x = 1")
chk("``` … ```", _strip_code_fence("```\ny = 2\n```") == "y = 2")
chk("وبلا سياج", _strip_code_fence("z = 3") == "z = 3")

print("\n" + "═" * 70)
print(" ٧) والمسار يستدعيه — مشروطاً بإذن المستخدم")
print("═" * 70)
import inspect
from pipeline.orchestrator import WeaverOrchestrator as _W
_src = inspect.getsource(_W._export)
chk("النداء داخل _export", "build_with_model" in _src)
chk("مشروطٌ بـexec_enabled — فبلا إذنٍ لا نموذجَ يُنادى ولا كلفة",
    "exec_enabled()" in _src)
chk("ويسبق المُصدِّر المعتاد، فالفشلُ يسقط إليه",
    _src.index("build_with_model") < _src.index('if fmt == "docx"'))
chk("ومسارُ السكربت يُحفظ على البطاقة", "build_script" in _src)
chk("والإخفاق يُسجَّل للمستخدم", "بُني بالمُصدِّر المعتاد" in _src)

print("\n" + "═" * 70)
print(" النتيجة: " + ("PASS ✅" if ok else "FAIL ❌"))
print("═" * 70)
sys.exit(0 if ok else 1)
