# -*- coding: utf-8 -*-
"""تصميمُ العروض الاحترافيّ (pptx_design.py) عبر office.py.

قِيس على عرضٍ من هاتف المستخدم (طلب ١٦ شريحة بأشكالٍ مربّعةٍ ودائريّة):
  · خرجت ٢١: غلافٌ وختامٌ يُضافان دائماً + ٥ فواصل + شريحتا أشكالٍ في الآخر،
    والمحتوى ١٢ — وقيل للمستخدم «١٦ محتوى».
  · قالبٌ واحد: نقاطٌ بحجم ١٨ في الزاوية وثلثا الشريحة فارغ.
  · الأشكال: سكربتٌ كتبه النموذجُ بنفسه، في شريحتين منفصلتين.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "capabilities/skills/pptx_builder/scripts"))
import pptx_design as PD                  # noqa: E402
from pptx import Presentation             # noqa: E402
from pptx.oxml.ns import qn               # noqa: E402

P = F = 0


def ok(name, cond, extra=""):
    global P, F
    if cond:
        P += 1
        print(f"  ✓ {name}")
    else:
        F += 1
        print(f"  ✗ {name}" + (f"   — {extra}" if extra else ""))


T = tempfile.mkdtemp(prefix="weaver-test-design-")
OFFICE = os.path.join(_ROOT, "pipeline", "office.py")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def cli(*args):
    r = subprocess.run([sys.executable, OFFICE] + [str(a) for a in args],
                       capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout + r.stderr


def J(name, obj):
    p = os.path.join(T, name)
    json.dump(obj, open(p, "w", encoding="utf-8"), ensure_ascii=False)
    return p


def frames(slide):
    return [sh for sh in slide.shapes if sh.name == PD.FRAME_NAME]


def texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame
            and sh.text_frame.text.strip()]


AR = {"title": "الاضطرابات السلوكية", "subtitle": "المفهوم والعلاج",
      "slides_total": 10, "slides": [
          {"layout": "section", "title": "مدخل"},
          {"title": "نقاط", "points": ["أولى", "ثانية", "English point"]},
          {"title": "نقاطٌ وصورة", "points": ["أ", "ب"], "image": {"shape": "circle"}},
          {"title": "مقارنة", "cards": [{"title": "المشكلة", "text": "عابرة"},
                                        {"title": "الاضطراب", "text": "دائم"},
                                        {"title": "المعيار", "text": "الأخصّائيّ"}]},
          {"title": "الأطراف", "circles": [{"title": "الأسرة"}, {"title": "المعلّم"},
                                           {"title": "الطفل"}], "images": True},
          {"title": "مشاهد", "images": {"shape": "square", "items": [
              {"caption": "١"}, {"caption": "٢"}, {"caption": "٣"}, {"caption": "٤"}]}},
          {"title": "الخطوات", "steps": [{"title": "ملاحظة"}, {"title": "تشخيص"},
                                         {"title": "خطّة"}], "notes": "اشرح ببطء"},
          {"title": "التصنيفات", "stats": [{"value": "ADHD", "text": "قصور الانتباه"},
                                           {"value": "ODD", "text": "التحدّي"}]}]}

print("\n— العددُ بالضبط —")
D1 = os.path.join(T, "ar.pptx")
c, o = cli("build", J("ar.json", AR), "--out", D1)
ok("١٠ مطلوبة ⟵ ١٠ (غلاف + ٨ + ختام)", c == 0 and len(Presentation(D1).slides) == 10
   and "10 slides (cover 1 + 8 + closing 1)" in o, o[-300:])
c, o = cli("build", J("bad.json", dict(AR, slides_total=8)), "--out",
           os.path.join(T, "bad.pptx"))
ok("عددٌ لا يطابق ⟵ رفضٌ بالسبب ولا ملفّ", c == 1 and "asked for 8 slides; this "
   "spec makes 10" in o and not os.path.exists(os.path.join(T, "bad.pptx")), o)
c, o = cli("build", J("nc.json", dict(AR, slides_total=8, cover=False,
                                      closing=False)), "--out",
           os.path.join(T, "nc.pptx"))
ok("بلا غلافٍ ولا ختام ⟵ ٨", c == 0 and len(Presentation(os.path.join(
    T, "nc.pptx")).slides) == 8, o[-200:])

print("\n— الأنواع —")
prs = Presentation(D1)
sl = list(prs.slides)
ok("الغلافُ بعنوانه وعنوانه الفرعيّ", "الاضطرابات السلوكية" in texts(sl[0])
   and "المفهوم والعلاج" in texts(sl[0]))
ok("الفاصلُ مرقّم «١»", "١" in texts(sl[1]) and "مدخل" in texts(sl[1]))
ok("نقاطٌ مرقّمة (١ ٢ ٣)", all(x in texts(sl[2]) for x in ("١", "٢", "٣")))
ok("نقاطٌ وصورة ⟵ إطارٌ دائريٌّ فارغ", len(frames(sl[3])) == 1
   and frames(sl[3])[0]._element.spPr.find(A + "prstGeom").get("prst") == "ellipse")
cards = [sh for sh in sl[4].shapes if sh.shape_type == 1 and sh._element.spPr.find(
    A + "prstGeom").get("prst") == "roundRect"]
ok("ثلاثُ بطاقات", len(cards) == 3, len(cards))
ok("الدوائرُ أطرُ صور (images: true)", len(frames(sl[5])) == 3 and all(
    f._element.spPr.find(A + "prstGeom").get("prst") == "ellipse" for f in frames(sl[5])))
ok("أربعةُ أطرٍ مربّعة بتعليق", len(frames(sl[6])) == 4 and all(
    x in texts(sl[6]) for x in ("١", "٤")))
ok("ملاحظاتُ المتحدّث", sl[7].notes_slide.notes_text_frame.text == "اشرح ببطء")
ok("أرقامٌ كبيرة", "ADHD" in texts(sl[8]) and "ODD" in texts(sl[8]))
ok("الختام", "شكراً لكم" in texts(sl[9]))
ok("القالبُ محفوظٌ في الملفّ (لتأخذه شريحةٌ تُضاف)",
   prs.core_properties.category == "weaver-design:academic_navy")

print("\n— من اليمين للعربيّ، ومن اليسار للإنجليزيّ —")
xs = sorted(((c_.left, c_) for c_ in cards), key=lambda v: -v[0])
first = [sh for sh in sl[4].shapes if sh.has_text_frame
         and sh.text_frame.text == "المشكلة"][0]
ok("البطاقةُ الأولى يميناً", first.left > xs[-1][0])
pp = {sh.text_frame.text: sh.text_frame.paragraphs[0]._p.pPr.attrib
      for sh in sl[2].shapes if sh.has_text_frame}
ok("النقطةُ العربيّة rtl=1 يمين، والإنجليزيّةُ rtl=0 يسار",
   pp["أولى"].get("rtl") == "1" and pp["أولى"].get("algn") == "r"
   and pp["English point"].get("rtl") == "0" and pp["English point"].get("algn") == "l",
   pp)
EN = {"title": "Behaviour", "theme": "modern_blue", "slides_total": 3, "slides": [
    {"title": "Compare", "cards": [{"title": "First", "text": "a"},
                                   {"title": "Second", "text": "b"}]}]}
D2 = os.path.join(T, "en.pptx")
c, o = cli("build", J("en.json", EN), "--out", D2)
s2 = list(Presentation(D2).slides)[1]
f1 = [sh for sh in s2.shapes if sh.has_text_frame and sh.text_frame.text == "First"][0]
f2 = [sh for sh in s2.shapes if sh.has_text_frame and sh.text_frame.text == "Second"][0]
ok("الإنجليزيّ: البطاقةُ الأولى يساراً", c == 0 and f1.left < f2.left, o[-200:])
ok("القالبُ modern_blue", "theme modern_blue" in o)
ok("قالبٌ بالوصف («تقني») ⟵ قالبٌ تقنيّ", "tech" in str(PD.load_theme("تقني")
                                                     .get("mood", [])))
ok("قالبٌ بلون", PD.load_theme("#7A2E33")["id"] == "custom")

print("\n— أطرُ الصور —")
from PIL import Image   # noqa: E402
Image.new("RGB", (900, 600), "#3b7dd8").save(os.path.join(T, "wide.png"))
Image.new("RGB", (600, 900), "#d8703b").save(os.path.join(T, "tall.png"))
c, o = cli("edit", D1, J("ops.json", [
    {"op": "set_image", "slide": 6, "frame": 1, "path": os.path.join(T, "wide.png")},
    {"op": "set_image", "slide": 7, "frame": 2, "path": os.path.join(T, "tall.png")},
    {"op": "set_image", "slide": 7, "frame": 9, "path": os.path.join(T, "tall.png")},
    {"op": "add_slide", "after": 5, "title": "توصيات",
     "cards": [{"title": "الحوار", "text": "يوميّاً"}, {"title": "الثبات", "text": "دائماً"}]}]))
E = os.path.join(T, "ar-edited.pptx")
pe = list(Presentation(E).slides)
f6 = frames(pe[6])[0]           # الدائرةُ الأولى (بعد الشريحة المضافة بعد ٥)
bf = f6._element.spPr.find(A + "blipFill")
ok("الدائرةُ مملوءةٌ بالصورة وتبقى دائرة", bf is not None and f6._element.spPr.find(
    A + "prstGeom").get("prst") == "ellipse")
src = bf.find(A + "srcRect")
ok("  ⟵ صورةٌ عريضةٌ في دائرة: يُقصّ الجانبان بالتساوي", src.get("l") == src.get("r")
   and int(src.get("l")) > 0 and src.get("t") == "0", dict(src.attrib))
ok("  ⟵ وعلامةُ «أضف صورة» زالت", not f6.text_frame.text.strip())
ok("إطارٌ غيرُ موجود ⟵ يُقال بعددها", "has 4 image frames" in o and c == 2, o)
added = pe[5]
ok("الشريحةُ المضافة بالتصميم وألوانِ العرض", "الحوار" in texts(added) and any(
    sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == "1B2A4A"
    for sh in added.shapes if sh.shape_type == 1), texts(added))
c, o = cli("info", E)
ok("info يُظهر الأطرَ وحالَها", "image-frame circle (filled)" in o
   and "image-frame square (empty)" in o)

print("\n— الحجم —")
ok("«ADHD» لا تنكسر (عرضُ الحروف الكبيرة)", PD._lines("ADHD", 2.37 * 72 * 0.94,
                                                     PD.fit("ADHD", 2.37, 1.5, 60, 24,
                                                            bold=True), True) == 1)
ok("نصٌّ طويلٌ يصغر ولا ينزل عن الحدّ", PD.fit("كلمة " * 80, 4, 1, 26, 15) == 15)

print("\n— القالبُ القديمُ باقٍ عند الطلب —")
c, o = cli("build", J("cl.json", {"title": "عرض", "design": "classic", "slides": [
    {"title": "أ", "points": ["١"]}]}), "--out", os.path.join(T, "cl.pptx"))
ok("design: classic ⟵ build_pptx كما كان", c == 0 and any(
    "◀" in t for t in texts(list(Presentation(os.path.join(T, "cl.pptx")).slides)[1])))

print("\n— الرسمُ الأصليّ في PowerPoint (أرقامُه تُعدَّل فيه) —")
L3 = ["بسيطة", "متوسّطة", "حادّة"]
NAT = {"title": "رسوم", "slides_total": 4, "cover": False, "closing": False, "slides": [
    {"title": "أعمدة", "chart": {"type": "bar", "data": {"labels": L3,
                                                        "values": [12, 7, 3]}}},
    {"title": "أفقيّة", "chart": {"type": "horizontal_bar", "data": {
        "labels": L3, "values": [12, 7, 3]}}},
    {"title": "توزيع", "chart": {"type": "histogram", "data": {"values": [1, 2, 2, 3]}}},
    {"title": "صورة", "chart": {"type": "bar", "as_image": True,
                                "data": {"labels": L3, "values": [1, 2, 3]}}}]}
try:
    import matplotlib  # noqa: F401  — شريحتا الصورة تحتاجانها
    HAVE_MPL = True
except ImportError:
    HAVE_MPL = False
    NAT = dict(NAT, slides_total=2, slides=NAT["slides"][:2])
DN = os.path.join(T, "nat.pptx")
c, o = cli("build", J("nat.json", NAT), "--out", DN)
try:
    import xlsxwriter  # noqa: F401
    HAVE_XW = True
except ImportError:
    HAVE_XW = False
if HAVE_XW:
    ok("رسمان أصليّان%s" % ("، وصورتان والسببُ مذكور" if HAVE_MPL else ""),
       c == 0 and "charts: 2 native" in o and (not HAVE_MPL or (
           "type not native" in o and "as_image requested" in o)), o[-400:])
    ns = list(Presentation(DN).slides)
    gf = [sh for sh in ns[0].shapes if sh.has_chart]
    ok("شريحةُ الأعمدة: رسمٌ أصليّ لا صورة", len(gf) == 1 and not any(
        sh.shape_type == 13 for sh in ns[0].shapes))
    ch = gf[0].chart
    ok("  ⟵ أرقامُه في الرسم (١٢ ٧ ٣)", list(ch.plots[0].series[0].values)
       == [12.0, 7.0, 3.0])
    emb = [n_ for n_ in zipfile.ZipFile(DN).namelist() if n_.startswith(
        "ppt/embeddings/") and n_.endswith(".xlsx")]
    ok("  ⟵ وجدولُ بياناته داخل الملفّ (تحرير البيانات في PowerPoint)", len(emb) == 2,
       emb)
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(zipfile.ZipFile(DN).read(emb[0])))
    vals = [r[1].value for r in wb.active.iter_rows(min_row=2, max_row=4)]
    ok("  ⟵ والجدولُ فيه الأرقامُ نفسُها", vals == [12, 7, 3], vals)
    ok("العربيّ: الفئةُ الأولى يميناً (محورٌ معكوس)",
       ch.category_axis.reverse_order is True)
    hb = [sh for sh in ns[1].shapes if sh.has_chart][0].chart
    vx = hb._chartSpace.xpath(".//c:valAx")[0].find(qn("c:crosses")).get("val")
    ok("أفقيّ: الأولى أعلى، والأعمدةُ من اليمين، ومحورُ القيم أسفل (c:crosses=max)",
       hb.category_axis.reverse_order and hb.value_axis.reverse_order and vx == "max",
       vx)
    if HAVE_MPL:
        ok("histogram ⟵ صورة (ليس في PowerPoint)", any(
            sh.shape_type == 13 for sh in ns[2].shapes))
        ok("as_image ⟵ صورةٌ بطلب", any(sh.shape_type == 13 for sh in ns[3].shapes))
# بلا xlsxwriter ⟵ صورةٌ كما كان، والسببُ يُقال
blk = tempfile.mkdtemp()
open(os.path.join(blk, "xlsxwriter.py"), "w").write("raise ImportError('blocked')\n")
env = dict(os.environ, PYTHONPATH=blk + os.pathsep + os.environ.get("PYTHONPATH", ""))
r = subprocess.run([sys.executable, OFFICE, "build", J("nat1.json", dict(
    NAT, slides_total=1, slides=NAT["slides"][:1])), "--out",
    os.path.join(T, "nat1.pptx")], capture_output=True, text=True, env=env, timeout=300)
if HAVE_MPL:
  ok("بلا xlsxwriter ⟵ الرسمُ صورةٌ ويُقال لماذا", r.returncode == 0
   and "xlsxwriter not installed" in r.stdout and any(
       sh.shape_type == 13 for sh in list(Presentation(os.path.join(
           T, "nat1.pptx")).slides)[0].shapes), r.stdout[-300:])

print("\n— الرسمُ التفاعليّ (HTML) —")
HT = os.path.join(T, "c.html")
c, o = cli("chart", J("h.json", {"type": "grouped_bar", "title": "حالات",
                                 "data": {"labels": L3, "series": {"٢٠٢٣": [1, 2, 3],
                                                                   "٢٠٢٤": [2, 3, 4]}}}),
           "--out", HT)
pg = open(HT, encoding="utf-8").read() if os.path.isfile(HT) else ""
ok("chart --out .html ⟵ صفحةٌ واحدة", c == 0 and "interactive chart grouped_bar" in o
   and pg.startswith("<!doctype html>"), o)
ok("  ⟵ Chart.js مضمَّنةٌ فيها (بلا إنترنت)", "Chart.js" in pg or "kurkle" in pg)
ok("  ⟵ من اليمين (dir=rtl) والبياناتُ فيها", 'dir="rtl"' in pg and "٢٠٢٣" in pg
   and '"rtl": true' in pg)
c, o = cli("chart", J("h2.json", {"type": "bar", "data": {"labels": ["A"], "values": [1]}}),
           "--out", os.path.join(T, "c.svg"))
ok("امتدادٌ آخر ⟵ رفضٌ يذكر .png و.html", c == 1 and ".html (interactive)" in o, o)

shutil.rmtree(T, ignore_errors=True)
print("\n" + "=" * 62)
print(f" RESULT: {'PASS' if F == 0 else 'FAIL'}   ({P}/{P + F})")
print("=" * 62)
sys.exit(1 if F else 0)
