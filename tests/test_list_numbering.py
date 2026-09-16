# -*- coding: utf-8 -*-
"""الترقيم: كل قائمةٍ تبدأ من ١، ولا يستمرّ العدّاد عبر الأقسام."""
import sys, os, zipfile, re
# THE TESTS ONLY RAN ON THE MACHINE THEY WERE WRITTEN ON. The repository
# root was hardcoded as an absolute path, so on any other checkout the insert
# pointed at a directory that does not exist and every file died on
# "No module named 'pipeline'" before running a single check. The root is
# where this file lives, one directory up — the way tests/smoke_pipeline.py
# already computes it — so the suite runs from any clone on any device.
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)
import tempfile as _tempfile
# a scratch directory of the machine RUNNING the test, never a path
# baked in from the machine that wrote it
_TMP = _tempfile.mkdtemp(prefix="weaver-test-")
sys.path.insert(0, _ROOT + "/capabilities/skills/docx_builder/scripts")
import docx_advanced as DA
from docx import Document
from docx.oxml.ns import qn
ok = True
OUT = os.path.join(_TMP, "num.docx")

BODY = """## المبحث الأول
تمهيدٌ قصير.

1. البند الأول
2. البند الثاني
3. البند الثالث

فقرةٌ فاصلة تُنهي القائمة.

## المبحث الثاني

1. بندٌ في القائمة الثانية
2. بندٌ آخر

## المبحث الثالث

1. قائمةٌ ثالثة
"""

doc = Document()
DA._add_body_markdown(doc, BODY, "ar", "academic_navy", "Arial")
doc.save(OUT)

d = Document(OUT)
runs, cur = [], None
for p in d.paragraphs:
    pPr = p._p.find(qn("w:pPr"))
    nid = None
    if pPr is not None:
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            n = numPr.find(qn("w:numId"))
            if n is not None:
                nid = n.get(qn("w:val"))
    if nid:
        if cur is None or cur[0] != nid:
            cur = [nid, 0]; runs.append(cur)
        cur[1] += 1
    else:
        cur = None

print("قوائم المستند (numId ← عدد البنود):")
for nid, cnt in runs:
    print(f"   numId={nid}  ← {cnt} بنود")
ids = [r[0] for r in runs]
good = (len(runs) == 3 and len(set(ids)) == 3 and [r[1] for r in runs] == [3, 2, 1])
ok &= good
print(f"  ⟵ ثلاث قوائم بثلاثة عدّادات مستقلّة: {'✅' if good else '❌'}")

# numbering.xml يجب أن يحمل التعريفات الجديدة وكلّها على نفس abstractNum
with zipfile.ZipFile(OUT) as z:
    xml = z.read("word/numbering.xml").decode("utf-8")
pairs = re.findall(r'<w:num [^>]*w:numId="(\d+)"[^>]*>\s*<w:abstractNumId w:val="(\d+)"',
                   xml)
mp = dict(pairs)
abs_ids = {mp.get(i) for i in ids if i in mp}
good2 = (len(abs_ids) == 1 and None not in abs_ids)
ok &= good2
print(f"  التعريفات في numbering.xml: {[(i, mp.get(i)) for i in ids]}")
print(f"  ⟵ كلّها على abstractNum واحد (الشكل لم يتغيّر): {'✅' if good2 else '❌'}")

print("\n=== التدهور الآمن ===")
class NoNum:
    class part:
        @property
        def numbering_part(self): raise KeyError("no numbering")
    styles = {}
r = DA.new_list_numbering(NoNum())
ok &= (r is None)
print(f"  مستندٌ بلا جزء ترقيم ⟶ {r!r} (سلوك الأمس) {'✅' if r is None else '❌'}")
p0 = Document().add_paragraph("x")
DA.apply_list_numbering(p0, None)
print(f"  apply(None) ⟶ لا يرمي ولا يغيّر  ✅")

print(f"\n  الملف: {OUT}  ({os.path.getsize(OUT)} بايت)")
print("\n" + ("PASS ✅" if ok else "FAIL ❌"))
sys.exit(0 if ok else 1)
