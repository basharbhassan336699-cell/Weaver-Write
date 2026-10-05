"""لوحةُ المعاينة الجانبيّة: صفحاتُ أيِّ ملفٍّ كما تُطبع.

    python3 tests/test_side_panel.py

قُورن الرسمُ يدوياً بصفحات LibreOffice (مستندٌ عربيٌّ بجدولٍ مظلّل وترويسةٍ
ورقمِ صفحةٍ وتعداد): الأسطرُ وانكسارُها والجدولُ والصورةُ متطابقة تقريباً.
وهنا ما يُفحص آليّاً.
"""
import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "web"))

FAILS = []


def check(name, cond):
    print(("  ✓ " if cond else "  ✗ ") + name)
    if not cond:
        FAILS.append(name)


def _bidi(p):
    from docx.oxml import OxmlElement
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    return p


def _doc(tmp):
    import docx
    from docx.enum.text import WD_BREAK
    d = docx.Document()
    hp = d.sections[0].header.paragraphs[0]
    hp.add_run("HEADER-TEXT")
    p = _bidi(d.add_paragraph())
    p.add_run("واجب رقم ( 2 + 1 )")
    for _ in range(3):
        _bidi(d.add_paragraph()).add_run("نصٌّ عربيٌّ قصيرٌ في فقرةٍ من اليمين.")
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "أ"
    br = d.add_paragraph().add_run()
    br.add_break(WD_BREAK.PAGE)
    d.add_paragraph("English after the page break.")
    path = os.path.join(tmp, "t.docx")
    d.save(path)
    return path


def test_docx_render(tmp):
    import docx_render as R
    path = _doc(tmp)
    pages = R.render_docx(path)
    check("Word: فاصلُ الصفحة ⟵ صفحتان", len(pages) == 2)
    w, h = pages[0].size
    check("Word: أبعادُ الصفحة من الملفّ (Letter/A4)", 1.25 < h / w < 1.45)
    # العربيُّ RTL: الحبرُ قربَ الهامش الأيمن لا الأيسر
    g = pages[0].convert("L")
    box_y = (int(h * 0.12), int(h * 0.2))

    def ink(x0, x1):
        n = 0
        for y in range(box_y[0], box_y[1], 2):
            for x in range(x0, x1, 2):
                if g.getpixel((x, y)) < 128:
                    n += 1
        return n
    right = ink(int(w * 0.6), int(w * 0.88))
    left = ink(int(w * 0.12), int(w * 0.4))
    check("Word: الفقرةُ العربيّةُ من اليمين", right > 50 and right > left * 3)
    # الترويسة مرسومة فوق الهامش العلويّ
    top = sum(1 for y in range(int(h * 0.03), int(h * 0.09), 2)
              for x in range(0, w, 3) if g.getpixel((x, y)) < 128)
    check("Word: الترويسةُ تُرسم", top > 20)
    # ترتيبُ الكلمات بصريّاً: ( 2 + 1 ) في فقرةٍ من اليمين
    import docx
    doc = docx.Document(path)
    ctx = R._Ctx(doc)
    para = R.layout_para(ctx, doc.paragraphs[0]._p, 450)
    vis = [t.text for t in para.lines[0].toks if t.kind == "text"]
    check("Word: الأقواسُ والأرقامُ في ترتيبها (قواعدُ الاتّجاه)",
          vis == [")", "1", "+", "2", "(", "رقم", "واجب"])
    br = [t for t in para.lines[0].toks if t.kind == "text" and t.text in "()"]
    check("Word: القوسُ في السياق العربيّ يُرسم RTL (ينعكس)", all(t.rtl for t in br))


def test_cover():
    import docx_render as R
    taj = os.path.join(ROOT, "engines", "fonts-core", "arabic", "Tajawal-Regular.ttf")
    p, txt = R._cover(taj, "نص")
    check("الخطوط: حرفٌ موجودٌ يبقى بخطّه", p == taj and txt == "نص")
    p, txt = R._cover(taj, "◀")
    check("الخطوط: حرفٌ غيرُ موجود ⟵ خطٌّ يحمله أو «•» (لا مربّعٌ فارغ)",
          (p != taj and txt == "◀") or txt == "•")


def test_page_render(tmp):
    import page_render as P
    cache = os.path.join(tmp, "cache")
    path = _doc(tmp)
    m = P.pages_for(path, cache)
    check("صفحات: Word", m.get("kind") == "pages" and m.get("n") == 2)
    m2 = P.pages_for(path, cache)
    check("صفحات: تُحفظ ولا تُعاد", m2.get("id") == m.get("id"))
    check("صفحات: صورةُ الصفحة موجودة", bool(P.page_path(cache, m["id"], 1)))
    check("صفحات: لا خروجَ من مجلّد المعاينات",
          P.page_path(cache, "../../etc", 1) == "" and
          P.page_path(cache, m["id"], "x") == "")
    with open(os.path.join(tmp, "a.md"), "w", encoding="utf-8") as fh:
        fh.write("# عنوان\n\n- نقطة\n\n| أ | ب |\n|---|---|\n| 1 | 2 |\n")
    check("صفحات: Markdown", P.pages_for(os.path.join(tmp, "a.md"), cache).get("n") == 1)
    with open(os.path.join(tmp, "a.csv"), "w", encoding="utf-8") as fh:
        fh.write("name,qty\nقلم,3\n")
    check("صفحات: CSV", P.pages_for(os.path.join(tmp, "a.csv"), cache).get("n") == 1)
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        for i in range(120):
            wb.active.append(["منتج %d" % i, i])
        wb.save(os.path.join(tmp, "a.xlsx"))
        mx = P.pages_for(os.path.join(tmp, "a.xlsx"), cache)
        check("صفحات: Excel ⟵ جدولٌ مقسَّمٌ على صفحات", mx.get("n", 0) >= 2)
    except ImportError:
        pass
    try:
        import pptx
        pr = pptx.Presentation()
        for i in range(3):
            pr.slides.add_slide(pr.slide_layouts[0]).shapes.title.text = "شريحة %d" % i
        pr.save(os.path.join(tmp, "a.pptx"))
        ms = P.pages_for(os.path.join(tmp, "a.pptx"), cache)
        check("صفحات: PowerPoint ⟵ شرائح", ms.get("kind") == "slides" and ms.get("n") == 3)
    except ImportError:
        pass
    check("صفحات: صورة ⟵ تُعرض كما هي", P.kind_of("x.jpg") == "image")
    check("صفحات: HTML ⟵ صفحةٌ حيّة", P.kind_of("x.html") == "html")


def test_server():
    import server as srv
    s = open(os.path.join(ROOT, "web", "server.py"), encoding="utf-8").read()
    check("server: /api/pages و/api/pages/img",
          'path == "/api/pages"' in s and 'path == "/api/pages/img"' in s)
    check("server: HTML مرفوعٌ يُعرض نصّاً لا يعمل", "as_text=True" in s)
    check("server: مجلّدُ المعاينات خارجَ المستودع",
          ".weaver-write" in srv._PAGE_CACHE)


def test_ui():
    h = open(os.path.join(ROOT, "web", "index.html"), encoding="utf-8").read()
    check("ui: لوحةٌ جانبيّة (تحميل، تكبير، إغلاق)",
          'id="wvPanel"' in h and "wvp-dl" in h and "wvp-wide" in h and "wvp-close" in h)
    check("ui: رقمُ الصفحة على كلِّ صفحة", "wvp-num" in h and "' / ' + m.n" in h)
    check("ui: الملفُّ الناتجُ يُعاين بالضغط عليه",
          "card.addEventListener('click'" in h and "wvPanelable(ext)" in h)
    check("ui: والمرفوعُ قبل الإرسال وبعده", "metaBody: o" in h and "src=upload" in h)
    check("ui: على الهاتف تملأ الشاشة",
          ".wv-panel { position: fixed; inset: 0;" in h)


if __name__ == "__main__":
    print("side panel / page rendering")
    tmp = tempfile.mkdtemp(prefix="wv-sp-")
    try:
        test_docx_render(tmp)
        test_cover()
        test_page_render(tmp)
        test_server()
        test_ui()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("FAIL: " + ", ".join(FAILS) if FAILS else "PASS")
    sys.exit(1 if FAILS else 0)
