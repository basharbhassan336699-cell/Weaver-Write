# -*- coding: utf-8 -*-
"""اليدُ التي تنفّذ — وحدودُها.

الأداةُ تُمكّن النموذج من كتابة سكربتٍ يبني الملفّ النهائيّ بدل أن يملأ
فراغاتٍ في بايثون مكتوبةٍ سلفاً. وحدودُها مقصودة: مطفأةٌ افتراضياً، وبمهلة،
وفي دليلٍ مؤقّت، وتُرفض الأوامر المدمّرة، والفشلُ لا يضرّ."""
import sys, os, json, tempfile
_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)
for k in list(os.environ):
    if k.startswith("WEAVER_"):
        os.environ.pop(k)
from capabilities.tools.tool_exec_python import (
    run_python, screen_script, exec_enabled)
ok = True
_TMP = tempfile.mkdtemp(prefix="weaver-exec-test-")


def chk(label, good, detail=""):
    global ok
    ok &= bool(good)
    print(f"   {'✅' if good else '❌'} {label}" + (f" — {detail}" if detail else ""))


print("═" * 70)
print(" ١) مطفأةٌ حتى يأذن المستخدم")
print("═" * 70)
chk("بلا WEAVER_EXEC ⟶ لا تعمل", not exec_enabled())
r = run_python("print(1)", os.path.join(_TMP, "a.docx"))
chk("والنداء يُرفض بسببٍ معلن", not r["ok"] and "مطفأ" in r["reason"], r["reason"])
for v in ("1", "true", "yes", "on"):
    os.environ["WEAVER_EXEC"] = v
    chk(f"وتعمل عند WEAVER_EXEC={v}", exec_enabled())
os.environ["WEAVER_EXEC"] = "1"

print("\n" + "═" * 70)
print(" ٢) الأوامرُ المدمّرة تُرفض قبل التشغيل")
print("═" * 70)
for code, why in (
        ("import shutil; shutil.rmtree('/')", "حذفُ جذر"),
        ("import os; os.system('rm -rf /')", "os.system"),
        ("import subprocess; subprocess.run(['ls'])", "subprocess"),
        ("import socket; socket.socket()", "شبكة"),
        ("", "فارغ")):
    chk(f"يُرفض: {why}", bool(screen_script(code)), screen_script(code)[:40])
chk("ويُقبل سكربتٌ سليم", not screen_script("from docx import Document"))

print("\n" + "═" * 70)
print(" ٣) وتبني ملفاً حقيقياً")
print("═" * 70)
CODE = """
import os, json
from docx import Document
data = json.load(open(os.environ["WEAVER_PAYLOAD"], encoding="utf-8"))
d = Document()
d.add_heading(data["title"], 0)
for sec in data["sections"]:
    d.add_heading(sec["h"], level=sec["lv"])
    if sec.get("body"):
        d.add_paragraph(sec["body"])
    if sec.get("table"):
        t = d.add_table(rows=0, cols=len(sec["table"][0]))
        t.style = "Table Grid"
        for row in sec["table"]:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
d.save(os.environ["WEAVER_OUT"])
"""
payload = json.dumps({"title": "الزواج في الإسلام", "sections": [
    {"h": "المبحث الأول: أركان عقد الزواج", "lv": 1, "body": "ميثاقٌ غليظ…"},
    {"h": "المطلب الأول: اشتراط الولي", "lv": 2, "body": "يدور التساؤل…",
     "table": [["المذهب", "موقف الولي"], ["الحنفية", "ليس شرطاً"],
               ["المالكية", "شرط"]]}]}, ensure_ascii=False)
out = os.path.join(_TMP, "built.docx")
r = run_python(CODE, out, payload)
chk("نجح التشغيل", r["ok"], r["reason"] or "")
if r["ok"]:
    import zipfile
    x = zipfile.ZipFile(out).read("word/document.xml").decode()
    chk(f"وفيه جدولٌ حقيقيّ ({os.path.getsize(out)} بايت)",
        x.count("<w:tbl>") == 1)
    chk("وعناوينُ مرتّبة", x.count("Heading") >= 2)
    chk("والعربيةُ سليمة", "المبحث الأول" in x)

print("\n" + "═" * 70)
print(" ٤) والفشلُ لا يضرّ — يعود بالسبب لا بالانهيار")
print("═" * 70)
r2 = run_python("raise ValueError('kaboom')", os.path.join(_TMP, "b.docx"))
chk("سكربتٌ ينهار ⟶ ok=False وسببٌ مقروء",
    not r2["ok"] and "kaboom" in (r2["stderr"] or ""), r2["reason"])
r3 = run_python("print('done')", os.path.join(_TMP, "c.docx"))
chk("سكربتٌ لا ينتج ملفاً ⟶ يُكشف",
    not r3["ok"] and "صالح" in r3["reason"], r3["reason"])
os.environ["WEAVER_EXEC_TIMEOUT"] = "2"
r4 = run_python("import time\nwhile True: time.sleep(0.1)",
                os.path.join(_TMP, "d.docx"))
chk("وحلقةٌ لا تنتهي ⟶ تُقتل بالمهلة",
    not r4["ok"] and "المهلة" in r4["reason"], r4["reason"])
chk("ولا يبقى دليلٌ مؤقّتٌ معلَّق",
    not any(n.startswith("weaver-exec-") and "test" not in n
            for n in os.listdir(tempfile.gettempdir())))

print("\n" + "═" * 70)
print(" النتيجة: " + ("PASS ✅" if ok else "FAIL ❌"))
print("═" * 70)
sys.exit(0 if ok else 1)
