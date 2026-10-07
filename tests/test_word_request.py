"""البحثُ الأكاديميُّ في Word: قسماً قسماً، وتنسيقُ المستخدم، وفحصُ ما يُنسى.

    python3 tests/test_word_request.py

قِيس على هاتف المستخدم: بحثٌ من ١٥ صفحة بتنسيقٍ محدّد ⟵ النموذجُ ترك الأداةَ
(لا أحجامَ ولا تباعدَ فيها) وكتب سكربتاً واحداً طويلاً فانقطع عند حدّ الردّ،
وما خرج نسي فيه المراجعَ والتقسيمات. وهنا يُفحص العلاج، ومعه أنّ مواصفةً بلا
خياراتٍ جديدة تُبنى كما كانت.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "web"))
OFFICE = os.path.join(ROOT, "pipeline", "office.py")
FAILS = []


def check(name, cond):
    print(("  ✓ " if cond else "  ✗ ") + name)
    if not cond:
        FAILS.append(name)


def J(d, name, obj):
    p = os.path.join(d, name)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False)
    return p


def build(spec_path, out):
    r = subprocess.run([sys.executable, OFFICE, "build", spec_path, "--out", out],
                       capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout + r.stderr


P = "نصٌّ في فقرةٍ يشرح سموّ الدستور — وأثرَه في النظام (جفال، 2010). "


def sec(h, lv, body=None):
    return {"heading": h, "level": lv, "body": body or P * 3}


def test_sections_format_check(tmp):
    import docx
    for c in (1, 2):
        parts = [sec("المبحث %d" % c, 1, "تمهيد.")]
        for m in (1, 2, 3):
            parts.append(sec("المطلب %d.%d" % (c, m), 2))
            if not (c == 2 and m == 2):
                parts.append(sec("أولاً %d.%d" % (c, m), 3, "فقرة (الخطيب، 2021)."))
        J(tmp, "sections/s%02d.json" % c, parts)
    J(tmp, "sections/s09.json", [sec("الخاتمة", 1, "خاتمة (علوان، 2019).")])
    J(tmp, "refs.json", ["جفال، زياد. (2010). التنظيم الدستوري.",
                         "الخطيب، نعمان. (2021). الوسيط."])
    spec = {"title": "بحث", "cover": {"author": ""}, "font": "Times New Roman",
            "sections": [sec("المقدمة", 1), "sections/s01.json", "sections/s02.json",
                         "sections/s09.json"],
            "references": "refs.json", "page": {"size": "A4"},
            "format": {"body_size": 14, "heading_sizes": [16, 16, 14, 14],
                       "line_spacing": 1.25, "heading_underline": False,
                       "dashes": "-"},
            "outline": {"chapters": 2, "per_chapter": 3, "sub_min": 1}}
    out = os.path.join(tmp, "a.docx")
    rc, o = build(J(tmp, "spec.json", spec), out)
    check("قسماً قسماً: الأقسامُ والمراجعُ من ملفّاتٍ منفصلة ⟵ مستندٌ واحد",
          rc == 0 and os.path.isfile(out))
    d = docx.Document(out)
    texts = [p.text for p in d.paragraphs]
    check("قسماً قسماً: كلُّ المباحث بترتيبها",
          texts.index("المبحث 1") < texts.index("المبحث 2") < texts.index("الخاتمة"))
    check("قسماً قسماً: المراجعُ من refs.json", any("الوسيط" in t for t in texts))
    sec_ = d.sections[0]
    check("تنسيق: صفحة A4", abs(sec_.page_width.cm - 21.0) < 0.05)
    h1 = next(p for p in d.paragraphs if p.text == "المبحث 1")
    b = next(p for p in d.paragraphs if p.text.startswith("نصٌّ"))
    x1, xb = h1.runs[0]._r.xml, b.runs[0]._r.xml
    check("تنسيق: العنوانُ ١٦ للعربيّ أيضاً (szCs لا sz وحدَه)",
          'w:szCs w:val="32"' in x1 and 'w:sz w:val="32"' in x1)
    check("تنسيق: النصُّ ١٤", 'w:szCs w:val="28"' in xb)
    check("تنسيق: تباعدُ ١٫٢٥", b.paragraph_format.line_spacing == 1.25)
    check("تنسيق: «—» ⟵ «-» في النصّ", "—" not in b.text and " - " in b.text)
    cover = d.paragraphs[0]
    check("تنسيق: الغلافُ لا يُمَسّ", 'w:sz w:val="28"' not in
          (cover.runs[0]._r.xml if cover.runs else ""))
    check("فحص: سطرُ التنسيق", "format: page A4" in o and "line spacing 1.25" in o)
    check("فحص: استشهادٌ بلا مرجع يُذكر بعينه", "(علوان, 2019)" in o)
    check("فحص: مطلبٌ بلا تقسيم يُذكر بعينه", "«المطلب 2.2» has 0 level-3" in o)
    check("فحص: ✗ ⟵ لا يقول «تمّ»", "✗ check: 2 problems" in o)


def test_missing_refs_and_dash(tmp):
    spec = {"title": "Report", "sections": [
        {"heading": "Intro", "body": "The chair — designed 1955–1956 (Smith, 2020)."}]}
    rc, o = build(J(tmp, "n.json", spec), os.path.join(tmp, "n.docx"))
    check("فحص: استشهاداتٌ بلا قائمة مراجع ⟵ ✗", "NO references list" in o)
    check("فحص: الشرطاتُ الطويلةُ تُعدّ", "2 long dashes" in o)


def test_unchanged_without_options(tmp):
    """مواصفةٌ بلا format/page/outline: المستندُ كما كان، والإضافةُ أسطرُ فحصٍ فقط."""
    import docx
    sys.path.insert(0, os.path.join(ROOT, "pipeline"))
    import office as O
    spec = {"title": "تقرير", "sections": [{"heading": "مقدمة", "body": "نص — عربي."}]}
    out = os.path.join(tmp, "u.docx")
    O.build_word(dict(spec), out)
    d = docx.Document(out)
    p = next(p for p in d.paragraphs if p.text.startswith("نص"))
    check("بلا خيارات: الشرطةُ باقيةٌ كما كتبها", "—" in p.text)
    check("بلا خيارات: لا تباعدَ يُفرض", p.paragraph_format.line_spacing is None)
    check("بلا خيارات: _apply_word_format لا يلمس الملف",
          O._apply_word_format(out, spec, spec["sections"], "ar") == "")


def test_pinned_requests():
    import server as S
    from pipeline import weaver_core as W
    saved = (W.available, W.node_bin, W.ask)
    cap = {}
    try:
        W.available = lambda: True
        W.node_bin = lambda: "node"
        W.ask = lambda text, **k: (cap.__setitem__("t", text) or {"answer": "ok"})
        brief = "أريد بحثاً من ١٥ صفحة بخط Times New Roman وحجم ١٦ للعناوين. " * 5
        hist = [{"role": "user", "content": brief},
                {"role": "assistant", "content": "حسناً"}] + \
            [{"role": "user", "content": "أكمل"},
             {"role": "assistant", "content": "…"}] * 4
        S._chat_via_engine("أكمل", hist, timeout=5)
        t = cap.get("t", "")
        check("محادثةٌ طويلة: الطلبُ المفصّلُ القديمُ يبقى",
              "[طلباتك السابقة المفصّلة" in t and brief.strip() in t)
        cap.clear()
        S._chat_via_engine("أكمل", hist[-4:], timeout=5)
        check("محادثةٌ قصيرة: لا تثبيتَ ولا تكرار", "[طلباتك" not in cap.get("t", ""))
    finally:
        W.available, W.node_bin, W.ask = saved


def test_skill():
    s = open(os.path.join(ROOT, "capabilities", "prompts", "skills", "office-word",
                          "SKILL.md"), encoding="utf-8").read()
    check("المهارة: قسماً قسماً", "sections/s01.json" in s and "refs.json" in s)
    check("المهارة: لا سكربتَ بدلَ الأداة", "لا تكتب سكربتَ python-docx بدلَ الأداة" in s)
    check("المهارة: format و outline", '"format"' in s and '"outline"' in s)
    check("المهارة: لا تسلّم ملفّاً عليه ✗", "لا تسلّم" in s)


if __name__ == "__main__":
    print("Word research: sections · format · check")
    tmp = tempfile.mkdtemp(prefix="wv-wr-")
    try:
        test_sections_format_check(tmp)
        test_missing_refs_and_dash(tmp)
        test_unchanged_without_options(tmp)
        test_pinned_requests()
        test_skill()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("FAIL: " + ", ".join(FAILS) if FAILS else "PASS")
    sys.exit(1 if FAILS else 0)
