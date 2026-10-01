# -*- coding: utf-8 -*-
"""التصميمُ الحرّ للعروض (pptx_free.py) عبر office.py — تجريبيّ.

قِيس: القوالبُ تُخرج عروضاً متشابهة (عرضان بقالبٍ واحدٍ في اختبارات المستخدم).
هنا يصمّم النموذجُ كلَّ شريحةٍ بعناصرها، والأداةُ تبنيها أصليّةً وترسم صورتها
وتفحص بالقياس. وسطرُ PowerPoint ١٫٢ × الحجم — قِيس بـLibreOffice على خمسة خطوط.
"""
import json
import os
import subprocess
import sys
import tempfile

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(_ROOT, "capabilities/skills/pptx_builder/scripts"))
import pptx_free as PF                    # noqa: E402
from pptx import Presentation             # noqa: E402

P = F = 0


def ok(name, cond, extra=""):
    global P, F
    if cond:
        P += 1
        print(f"  ✓ {name}")
    else:
        F += 1
        print(f"  ✗ {name}" + (f"   — {extra}" if extra else ""))


T = tempfile.mkdtemp(prefix="weaver-test-free-")
OFFICE = os.path.join(_ROOT, "pipeline", "office.py")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def cli(*args, env=None):
    e = dict(os.environ)
    e.update(env or {})
    r = subprocess.run([sys.executable, OFFICE] + [str(a) for a in args],
                       capture_output=True, text=True, timeout=300, env=e)
    return r.returncode, r.stdout + r.stderr


def J(name, obj):
    p = os.path.join(T, name)
    json.dump(obj, open(p, "w", encoding="utf-8"), ensure_ascii=False)
    return p


GOOD = {"design": "free", "title": "الاضطرابات", "font": "Kufyan Arabic Regular",
        "slides_total": 3, "slides": [
            {"bg": {"gradient": ["#0B1F3A", "#1E4D7B"], "angle": 45}, "elements": [
                {"type": "shape", "shape": "circle", "x": -1.5, "y": 3.2, "w": 6, "h": 6,
                 "fill": "#C8A04A", "opacity": 0.18},
                {"type": "text", "x": 4.2, "y": 2.0, "w": 7.5, "h": 1.9,
                 "text": "الاضطرابات السلوكية والانفعالية", "size": 50, "bold": True,
                 "color": "#FFFFFF", "valign": "bottom"}]},
            {"bg": "#F7F4EE", "elements": [
                {"type": "image", "shape": "rounded", "x": 0.6, "y": 0.6, "w": 5.2,
                 "h": 6.3},
                {"type": "text", "x": 6.4, "y": 0.7, "w": 6.3, "h": 1.0,
                 "text": "ما المقصود بالاضطراب؟", "size": 36, "bold": True,
                 "color": "#0B1F3A"},
                {"type": "shape", "shape": "pill", "x": 6.4, "y": 6.0, "w": 3, "h": 0.6,
                 "fill": "#0B1F3A", "text": "Workshop 2026", "size": 16,
                 "align": "center", "valign": "middle"}]},
            {"bg": "#FFFFFF", "elements": [
                {"type": "chart", "x": 0.6, "y": 1.6, "w": 7.6, "h": 5.3,
                 "chart": {"type": "bar", "data": {"labels": ["أ", "ب"], "values": [3, 5]}},
                 "colors": ["#1E4D7B", "#C8A04A"]},
                {"type": "table", "x": 8.6, "y": 1.6, "w": 4.1, "h": 2.4,
                 "headers": ["العامل", "الأثر"], "rows": [["الحوار", "وقاية"]],
                 "size": 18}]}]}

print("\n— البناء: أشكالٌ أصليّة —")
out = os.path.join(T, "free.pptx")
rc, o = cli("build", J("good.json", GOOD), "--out", out)
ok("build ⟵ ✓ (free design) 3 slides", rc == 0 and "free design" in o
   and "3 slides" in o, o)
prs = Presentation(out)
ok("٣ شرائح بالضبط", len(prs.slides) == 3)
bg = prs.slides[0].background.fill
ok("الغلاف: خلفيّةٌ متدرّجة", bg.type is not None and
   prs.slides[0].background._cSld.bg.find(".//" + A + "gradFill") is not None)
t1 = [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text]
p0 = t1[0].text_frame.paragraphs[0]._p.find(A + "pPr")
ok("النصُّ العربيّ rtl=1", p0 is not None and p0.get("rtl") == "1")
pill = [sh for sh in prs.slides[1].shapes if sh.has_text_frame
        and sh.text_frame.text == "Workshop 2026"]
ok("نصٌّ داخل شكل، والإنجليزيُّ rtl=0", pill and pill[0].text_frame.paragraphs[0]
   ._p.find(A + "pPr").get("rtl") == "0")
ok("إطارُ صورةٍ فارغٌ يملؤه المستخدم", any(sh.name == "Weaver Image Frame"
                                       for sh in prs.slides[1].shapes))
s3 = prs.slides[2]
ch = [sh for sh in s3.shapes if sh.has_chart]
ok("رسمٌ أصليٌّ قابلٌ للتعديل", bool(ch) or "as image" in o, o)
if ch:
    pts = ch[0].chart.plots[0].series[0].points
    ok("  ⟵ بالألوان المختارة", str(pts[0].format.fill.fore_color.rgb) == "1E4D7B")
ok("جدولٌ أصليّ", any(sh.has_table for sh in s3.shapes))

print("\n— المعاينةُ والفحص —")
try:
    import PIL  # noqa: F401
    has_pil = True
except ImportError:
    has_pil = False
if has_pil:
    pdir = os.path.join(T, "free-preview")
    ok("صورةٌ جامعةٌ وصورةٌ لكلِّ شريحة", os.path.isfile(os.path.join(
        pdir, "slide-overview.png")) and all(os.path.isfile(os.path.join(
            pdir, "slide-s%02d.png" % i)) for i in (1, 2, 3)), os.listdir(T))
    ok("المواصفةُ السليمة ⟵ ✓ design check", "✓ design check" in o, o)
    ok("  ⟵ ويقول للنموذج: افتحها وانظر", "open them with read" in o)
    BAD = {"design": "free", "slides": [{"bg": "#FFFFFF", "elements": [
        {"type": "text", "x": 1, "y": 1, "w": 4, "h": 0.8, "size": 40, "bold": True,
         "text": "عنوانٌ طويلٌ جداً لا يتّسع في هذا المربّع الصغير أبداً"},
        {"type": "text", "x": 2, "y": 1.2, "w": 4, "h": 1, "text": "نصٌّ يقع فوق العنوان",
         "size": 24},
        {"type": "text", "x": 7, "y": 3, "w": 4, "h": 1, "text": "نصٌّ باهت", "size": 20,
         "color": "#EEEEEE"},
        {"type": "text", "x": 11, "y": 5, "w": 3, "h": 1, "text": "خارج الشريحة",
         "size": 20},
        {"type": "text", "x": 1, "y": 5, "w": 1.2, "h": 1,
         "text": "Internationalization", "size": 28},
        {"type": "text", "x": 1, "y": 6.5, "w": 5, "h": 0.5, "text": "حاشيةٌ صغيرة",
         "size": 9}]}]}
    rc, o2 = cli("build", J("bad.json", BAD), "--out", os.path.join(T, "bad.pptx"))
    for what, key in (("نصٌّ يتجاوز إطاره (والارتفاعُ اللازم)", "text needs"),
                      ("تداخلُ نصّين", "overlaps"), ("تباينٌ ضعيف", "low contrast"),
                      ("خارجَ الشريحة", "outside the slide"),
                      ("كلمةٌ أعرضُ من مربّعها", "wider than the box"),
                      ("خطٌّ صغير", "too small")):
        ok("الفحص: " + what, key in o2, o2)
    ok("  ⟵ والملفُّ يُحفظ ليُصلَح (لا يُمنع)", rc == 0
       and os.path.isfile(os.path.join(T, "bad.pptx")))
    # سطرُ PowerPoint: ١٫٢ × الحجم (لا مقاييسُ الخطّ كما في Word)
    lines, need, _w, _b = PF.layout({"text": "سطر", "size": 24}, 6, "Kufyan Arabic Regular",
                                    None)
    ok("ارتفاعُ السطر ١٫٢ × الحجم (٢٤ ⟵ ٢٨٫٨pt)",
       abs(need - 24 * 1.2 * PF.PX / 72) < 0.5, need)
else:
    print("  ⓘ Pillow غيرُ مثبَّت — تُخطّى المعاينة")

print("\n— العددُ والمفتاح —")
bad_n = dict(GOOD, slides_total=5)
rc, o3 = cli("build", J("n.json", bad_n), "--out", os.path.join(T, "n.pptx"))
ok("slides_total لا يطابق ⟵ رفضٌ بلا ملفّ", rc != 0 and "asked for 5" in o3
   and not os.path.isfile(os.path.join(T, "n.pptx")), o3)
rc, o4 = cli("build", J("g2.json", GOOD), "--out", os.path.join(T, "off.pptx"),
             env={"WEAVER_PPTX_FREE": "off"})
ok("WEAVER_PPTX_FREE=off ⟵ يعود إلى القوالب (رفضٌ يقول ذلك)", rc != 0
   and "template" in o4 and not os.path.isfile(os.path.join(T, "off.pptx")), o4)

print("\n" + "=" * 62)
print(f" RESULT: {'PASS' if F == 0 else 'FAIL'}   ({P}/{P + F})")
print("=" * 62)
sys.exit(1 if F else 0)
