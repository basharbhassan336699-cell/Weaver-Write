# -*- coding: utf-8 -*-
"""
pptx_free.py — التصميمُ الحرّ للعروض (الافتراضيّ؛ والقوالبُ احتياطٌ)
==============================================
القوالبُ (pptx_design.py) تُخرج عروضاً مرتّبةً متشابهة: ١٣ نوعاً ثابتاً، ويتغيّر
اللونُ أكثرَ ممّا يتغيّر التصميم — قِيس في اختبارات المستخدم: عرضان بقالبٍ واحد.
هنا يصمّم النموذجُ كلَّ شريحةٍ بنفسه — عناصرُ بمواضعها (بوصة، ١٣٫٣٣×٧٫٥) —
والأداةُ:
  ١) تبنيها أشكالاً أصليّةً قابلةً للتعديل في PowerPoint (نصّ، أشكال، صور، رسمٌ
     أصليّ، جدول) — بأدوات pptx_design نفسِها (الاتّجاه، الخطوط، الأرقام).
  ٢) ترسم كلَّ شريحةٍ صورةً (Pillow + raqm: العربيُّ متّصلٌ من اليمين) لينظر
     إليها النموذجُ ويصحّح — كما يفعل المصمّم.
  ٣) تفحص بالقياس: نصٌّ يتجاوز إطاره، وكلمةٌ لا تتّسع، وعنصرٌ خارجَ الشريحة،
     وتداخلُ نصوص، وخطٌّ صغير، وتباينٌ ضعيف — بأرقامٍ يُصلَح بها.

المواصفة:
  {"design": "free", "title": "…", "font": "…", "slides_total": N,
   "slides": [{"bg": "#0F172A" | {"gradient": ["#0F172A", "#1E3A8A"], "angle": 90},
               "elements": [
                 {"type": "text", "x":…, "y":…, "w":…, "h":…, "text": "…" | […],
                  "size": 40, "bold": true, "color": "#…", "align": "right",
                  "valign": "middle", "font": "…", "line_spacing": 1.1, "spacing": 6},
                 {"type": "shape", "shape": "rounded|rect|oval|circle|pill|triangle|
                  diamond|hexagon|chevron|arrow|left_arrow|pentagon|star|
                  parallelogram|donut|line", "fill": "#…" | {"gradient": […]},
                  "opacity": 0.6, "line": "#…", "line_w": 1.5, "radius": 0.2,
                  "rotation": 0, "text": …},
                 {"type": "image", "path": "…" (أو بلا مسار ⟵ إطارٌ فارغ),
                  "shape": "rect|rounded|circle"},
                 {"type": "chart", "chart": {"type": "bar", "data": {…}}, "colors": […]},
                 {"type": "table", "headers": […], "rows": [[…]], "size": 16,
                  "header_fill": "#…", "header_color": "#…", "fill": "#…",
                  "alt_fill": "#…", "color": "#…", "border": "#…"}],
               "notes": "…"}]}
"""
from __future__ import annotations

import os
import re
import sys

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)
import pptx_design as PD   # noqa: E402

W, H = PD.W_IN, PD.H_IN
PX = 120                                   # بكسل لكلِّ بوصة في صور المعاينة
MARGIN_X, MARGIN_Y = 0.05, 0.02            # هوامشُ مربّع النصّ (كما في Deck.text)
_FONTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.dirname(_HERE)))), "engines", "fonts-core")

SHAPES = {
    "rect": "RECTANGLE", "rectangle": "RECTANGLE", "square": "RECTANGLE",
    "rounded": "ROUNDED_RECTANGLE", "pill": "ROUNDED_RECTANGLE",
    "oval": "OVAL", "circle": "OVAL", "ellipse": "OVAL",
    "triangle": "ISOSCELES_TRIANGLE", "right_triangle": "RIGHT_TRIANGLE",
    "diamond": "DIAMOND", "hexagon": "HEXAGON", "octagon": "OCTAGON",
    "chevron": "CHEVRON", "arrow": "RIGHT_ARROW", "right_arrow": "RIGHT_ARROW",
    "left_arrow": "LEFT_ARROW", "pentagon": "PENTAGON", "star": "STAR_5_POINT",
    "parallelogram": "PARALLELOGRAM", "donut": "DONUT", "wave": "WAVE",
    "block_arc": "BLOCK_ARC", "pie": "PIE", "chord": "CHORD", "tear": "TEAR",
}


# ─────────────────────────── تطبيع ───────────────────────────
def _f(v, d=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return d


def _box(e):
    return _f(e.get("x")), _f(e.get("y")), max(0.0, _f(e.get("w"))), \
        max(0.0, _f(e.get("h")))


def _hex(c, d=None):
    if c is None or c == "":
        return d
    s = str(c).strip().lstrip("#")
    if len(s) == 3:
        s = "".join(ch * 2 for ch in s)
    return s[:6].upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", s[:6] or "") else d


def _rgb(c):
    c = _hex(c, "000000")
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


def _rel_lum(rgb):
    def ch(v):
        v = v / 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(v) for v in rgb[:3])
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b):
    la, lb = _rel_lum(a), _rel_lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def _bg_colors(bg):
    """[لون] أو [لون١، لون٢] (تدرّج)، وزاويته."""
    if isinstance(bg, dict):
        g = [c for c in (_hex(x) for x in (bg.get("gradient") or [])) if c]
        if len(g) >= 2:
            return g[:2], _f(bg.get("angle"), 90)
        c = _hex(bg.get("color"))
        return [c or "FFFFFF"], 0
    return [_hex(bg, "FFFFFF")], 0


def paragraphs(e):
    """نصُّ العنصر فقراتٍ: [{text,size,bold,color,align,font}] — «\\n» يفصل."""
    raw = e.get("text")
    if raw is None:
        return []
    items = raw if isinstance(raw, list) else [raw]
    out = []
    for it in items:
        o = dict(it) if isinstance(it, dict) else {"text": it}
        for part in str(o.get("text") if o.get("text") is not None else "").split("\n"):
            p = dict(o)
            p["text"] = part
            for k in ("size", "bold", "color", "align", "font"):
                if p.get(k) is None and e.get(k) is not None:
                    p[k] = e.get(k)
            p["size"] = _f(p.get("size"), 24) or 24
            p["bold"] = bool(p.get("bold"))
            out.append(p)
    return out


def _auto_color(backdrop_rgb):
    return "FFFFFF" if _rel_lum(backdrop_rgb) < 0.33 else "1F2937"


# ─────────────────────────── البناء (pptx أصليّ) ───────────────────────────
def _set_bg(s, bg):
    cols, ang = _bg_colors(bg)
    f = s.background.fill
    if len(cols) == 2:
        f.gradient()
        f.gradient_angle = ang
        st = f.gradient_stops
        st[0].color.rgb = RGBColor.from_string(cols[0])
        st[0].position = 0
        st[1].color.rgb = RGBColor.from_string(cols[1])
        st[1].position = 1.0
    else:
        f.solid()
        f.fore_color.rgb = RGBColor.from_string(cols[0])


def _fill_gradient(sh, spec):
    cols = [c for c in (_hex(x) for x in spec.get("gradient") or []) if c]
    if len(cols) < 2:
        return False
    sh.fill.gradient()
    sh.fill.gradient_angle = _f(spec.get("angle"), 90)
    st = sh.fill.gradient_stops
    st[0].color.rgb = RGBColor.from_string(cols[0])
    st[0].position = 0
    st[1].color.rgb = RGBColor.from_string(cols[1])
    st[1].position = 1.0
    return True


def _text_into(D, s, e, box=None, backdrop=None):
    x, y, w, h = _box(e)
    paras = paragraphs(e)
    if not paras:
        return box
    dflt = _hex(e.get("color")) or _auto_color(backdrop or (255, 255, 255))
    lines = []
    for p in paras:
        o = {"size": p["size"], "bold": p["bold"], "color": _hex(p.get("color"), dflt)}
        if p.get("align") in ("right", "left", "center"):
            o["align"] = p["align"]
        if p.get("font"):
            o["font"] = p["font"]
        lines.append((p["text"], o))
    va = str(e.get("valign") or "top").lower()
    box = D.text(s, x, y, w, h, lines, size=paras[0]["size"],
                 anchor=va if va in ("top", "middle", "bottom") else "top",
                 spacing=_f(e.get("spacing"), 0) or None, box=box)
    ls = _f(e.get("line_spacing"), 0)
    if ls:
        for p in box.text_frame.paragraphs:
            p.line_spacing = ls
    return box


def _image(D, s, e):
    x, y, w, h = _box(e)
    path = e.get("path") or e.get("image")
    shape = str(e.get("shape") or "rect").lower()
    if not (path and os.path.isfile(str(path))):
        return D.frame(s, x, y, w, h, "circle" if shape == "circle" else "square")
    if shape in ("circle", "oval", "rounded"):
        kind = MSO_SHAPE.OVAL if shape != "rounded" else MSO_SHAPE.ROUNDED_RECTANGLE
        sh = D.shape(s, kind, x, y, w, h, fill="FFFFFF", radius=_f(e.get("radius"), 0.1))
        return PD.fill_frame(s, sh, str(path))
    from PIL import Image
    iw, ih = Image.open(str(path)).size
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    fa, ia = w / max(h, 1e-6), iw / max(ih, 1)
    if str(e.get("fit") or "cover") == "cover":
        if ia > fa:
            cut = (1 - fa / ia) / 2
            pic.crop_left = pic.crop_right = cut
        else:
            cut = (1 - ia / fa) / 2
            pic.crop_top = pic.crop_bottom = cut
    return pic


def _chart(D, s, e, chart_png, backdrop):
    x, y, w, h = _box(e)
    spec = dict(e.get("chart") or {k: v for k, v in e.items()
                                   if k not in ("type", "x", "y", "w", "h")})
    if e.get("chart") is None and e.get("kind"):
        spec["type"] = e["kind"]
    spec.setdefault("theme", D.t["id"])
    spec.setdefault("lang", D.lang)
    spec.setdefault("font", D.fa)
    gf = None if spec.get("as_image") else PD.native_chart(D, s, spec, x, y, w, h)
    if gf is None:
        png = chart_png(spec)
        return s.shapes.add_picture(png, Inches(x), Inches(y), Inches(w), Inches(h))
    ch = gf.chart
    tc = _hex(e.get("text_color")) or _auto_color(backdrop)
    ch.font.color.rgb = RGBColor.from_string(tc)
    cols = [c for c in (_hex(v) for v in (e.get("colors") or [])) if c]
    if cols:
        plot = ch.plots[0]
        targets = (list(plot.series[0].points) if len(plot.series) == 1
                   and plot.vary_by_categories else list(plot.series))
        for i, t in enumerate(targets):
            try:
                t.format.fill.solid()
                t.format.fill.fore_color.rgb = RGBColor.from_string(cols[i % len(cols)])
            except Exception:
                pass
    return gf


def _table(D, s, e):
    x, y, w, h = _box(e)
    heads = [str(v) for v in e.get("headers") or []]
    rows = [[str(v) for v in r] for r in e.get("rows") or []]
    ncol = max([len(heads)] + [len(r) for r in rows] + [1])
    heads += [""] * (ncol - len(heads))
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    rtl = PD.is_rtl(" ".join(heads + (rows[0] if rows else [])), D.ar)
    if rtl:                                   # العمودُ الأوّلُ يميناً (كـpptx_design)
        heads = heads[::-1]
        rows = [r[::-1] for r in rows]
    nrow = 1 + len(rows)
    gf = s.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table
    tp = tbl._tbl.find(qn("a:tblPr"))
    tp.set("rtl", "0")
    for flag in ("bandRow", "firstRow"):
        tp.set(flag, "0")
    lens = [max([len(heads[c])] + [len(r[c]) for r in rows]) + 4 for c in range(ncol)]
    for c in range(ncol):
        tbl.columns[c].width = Inches(w * lens[c] / sum(lens))
    for r in range(nrow):
        tbl.rows[r].height = Inches(h / nrow)
    size = _f(e.get("size"), 16)
    hf = _hex(e.get("header_fill"), D.t["primary"])
    hc = _hex(e.get("header_color"), "FFFFFF")
    f1 = _hex(e.get("fill"), "FFFFFF")
    f2 = _hex(e.get("alt_fill"), PD._mix(f1, "000000", 0.05))
    col = _hex(e.get("color"), _auto_color(_rgb(f1)))
    for r in range(nrow):
        for c in range(ncol):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                hf if r == 0 else (f1 if r % 2 else f2))
            txt = heads[c] if r == 0 else rows[r - 1][c]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            rr = PD.is_rtl(txt, rtl)
            p.alignment = PP_ALIGN.RIGHT if rr else PP_ALIGN.LEFT
            p._p.get_or_add_pPr().set("rtl", "1" if rr else "0")
            run = p.add_run()
            run.text = txt
            run.font.size = Pt(size)
            run.font.bold = r == 0
            run.font.color.rgb = RGBColor.from_string(hc if r == 0 else col)
            rPr = run._r.get_or_add_rPr()
            for tag, fam in (("a:latin", D.fa if rr else D.fe), ("a:cs", D.fa)):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {})
                    rPr.append(el)
                el.set("typeface", fam)
    return gf


def _shape(D, s, e, backdrop):
    x, y, w, h = _box(e)
    name = str(e.get("shape") or "rect").lower()
    if name == "line":
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                    Inches(x + w), Inches(y + h))
        ln.line.color.rgb = RGBColor.from_string(_hex(e.get("line") or e.get("fill"),
                                                      "999999"))
        ln.line.width = Pt(_f(e.get("line_w"), 2))
        return ln
    kind = getattr(MSO_SHAPE, SHAPES.get(name, "RECTANGLE"))
    if name == "circle":
        d = min(w, h)
        x, y, w, h = x + (w - d) / 2, y + (h - d) / 2, d, d
    fill = e.get("fill")
    grad = isinstance(fill, dict)
    op = e.get("opacity")
    radius = 0.5 if name == "pill" else _f(e.get("radius"), 0.12)
    sh = D.shape(s, kind, x, y, w, h, fill=None if grad else _hex(fill),
                 line=_hex(e.get("line")), lw=_f(e.get("line_w"), 1.0),
                 dash=bool(e.get("dash")),
                 alpha=None if op is None else max(0.0, min(1.0, _f(op, 1))),
                 radius=radius)
    if grad:
        _fill_gradient(sh, fill)
    if e.get("rotation"):
        sh.rotation = _f(e.get("rotation"))
    if e.get("text") is not None:
        under = _rgb(_hex(fill) or "FFFFFF") if not grad else _rgb(
            (_bg_colors(fill)[0] or ["FFFFFF"])[0])
        _text_into(D, s, dict(e, x=x, y=y, w=w, h=h), box=sh,
                   backdrop=under if (op is None or _f(op, 1) >= 0.5) else backdrop)
    return sh


def build(spec, out, lang, chart_png, font_ar=None, font_en=None):
    """يبني العرض. يعيد (عددَ الشرائح، معرّفَ القالب)."""
    del PD.CHART_LOG[:]
    D = PD.Deck(spec.get("theme"), lang, font_ar, font_en, spec.get("title", ""))
    slides = [sd for sd in spec.get("slides") or [] if isinstance(sd, dict)]
    for sd in slides:
        s = D.prs.slides.add_slide(D.blank)
        bg = sd.get("bg") if sd.get("bg") is not None else spec.get("bg", "FFFFFF")
        _set_bg(s, bg)
        cols, _a = _bg_colors(bg)
        backdrop = _rgb(cols[0]) if len(cols) == 1 else tuple(
            (a + b) // 2 for a, b in zip(_rgb(cols[0]), _rgb(cols[1])))
        for e in sd.get("elements") or []:
            if not isinstance(e, dict):
                continue
            t = str(e.get("type") or ("text" if "text" in e else "shape")).lower()
            if t == "text":
                _text_into(D, s, e, backdrop=_backdrop_at(sd, e, backdrop))
            elif t == "shape":
                _shape(D, s, e, backdrop)
            elif t == "image":
                _image(D, s, e)
            elif t == "chart":
                _chart(D, s, e, chart_png, _backdrop_at(sd, e, backdrop))
            elif t == "table":
                _table(D, s, e)
        D.notes(s, sd.get("notes"))
    D.prs.core_properties.category = "weaver-design:%s" % D.t["id"]
    D.prs.core_properties.title = str(spec.get("title", ""))[:250]
    D.prs.core_properties.keywords = "weaver-free"
    D.prs.save(out)
    return len(slides), D.t["id"]


def _backdrop_at(sd, el, slide_bg):
    """لونُ ما تحت مركز العنصر: آخرُ شكلٍ معتمٍ قبله يحويه، وإلّا الخلفيّة."""
    x, y, w, h = _box(el)
    cx, cy = x + w / 2, y + h / 2
    under = slide_bg
    for e in sd.get("elements") or []:
        if e is el:
            break
        if not isinstance(e, dict) or str(e.get("type") or "") != "shape":
            continue
        ex, ey, ew, eh = _box(e)
        f = e.get("fill")
        if ex <= cx <= ex + ew and ey <= cy <= ey + eh and f and (
                e.get("opacity") is None or _f(e.get("opacity"), 1) >= 0.5):
            under = _rgb(_bg_colors(f)[0][0]) if isinstance(f, dict) else _rgb(f)
    return under


# ─────────────────────── المعاينةُ والفحص (Pillow) ───────────────────────
_FCACHE = {}


def _font_file(family, bold=False):
    try:
        if _FONTS_DIR not in sys.path:
            sys.path.insert(0, _FONTS_DIR)
        import font_catalog
        e = font_catalog.find_font(family) if family else None
        if e and not e.get("bundled") and e.get("stand_in"):
            e = font_catalog.find_font(e["stand_in"])
        if not e:
            e = font_catalog.find_font("Kufyan Arabic Regular")
        return font_catalog.file_for(e, "bold" if bold else None)
    except Exception:
        return None


def _font(family, size_pt, bold=False):
    from PIL import ImageFont
    path = _font_file(family, bold)
    px = max(6, int(round(size_pt * PX / 72)))
    key = (path, px)
    if key not in _FCACHE:
        try:
            lay = ImageFont.Layout.RAQM if hasattr(ImageFont, "Layout") else None
            _FCACHE[key] = ImageFont.truetype(path, px, layout_engine=lay) if lay \
                is not None else ImageFont.truetype(path, px)
        except Exception:
            try:
                _FCACHE[key] = ImageFont.truetype(path, px)
            except Exception:
                _FCACHE[key] = ImageFont.load_default()
    return _FCACHE[key]


class _Bold:
    """خطٌّ بلا ملفٍّ عريض: يُرسم بحدٍّ رفيع، وأعرضُ قليلاً (كتعريض PowerPoint)."""
    def __init__(self, f):
        self.f = f
        self.size = f.size
        self.stroke = max(1, int(round(f.size / 45)))

    def getlength(self, text, **kw):
        return self.f.getlength(text, **kw) * 1.03 + self.stroke

    def getmetrics(self):
        return self.f.getmetrics()


def _ink(d, xy, text, f, color, rtl):
    """ارسم سطراً — والعريضُ المصطنع بحدٍّ من اللون نفسِه."""
    kw = {"fill": color}
    if isinstance(f, _Bold):
        kw.update(stroke_width=f.stroke, stroke_fill=color)
        f = f.f
    try:
        d.text(xy, text, font=f, direction="rtl" if rtl else "ltr", **kw)
    except Exception:
        d.text(xy, text, font=f, **kw)


def _glyph_top(f, lh):
    """أعلى الرسم داخل سطرٍ ارتفاعُه lh: الخطُّ في منتصف السطر."""
    ff = f.f if isinstance(f, _Bold) else f
    try:
        asc, desc = ff.getmetrics()
    except Exception:
        asc, desc = ff.size, ff.size * 0.3
    return (lh - (asc + desc)) / 2


def _length(font, text, rtl):
    if isinstance(font, _Bold):
        try:
            return font.getlength(text, direction="rtl" if rtl else "ltr")
        except Exception:
            return font.getlength(text)
    try:
        return font.getlength(text, direction="rtl" if rtl else "ltr")
    except Exception:
        try:
            return font.getlength(text)
        except Exception:
            return len(text) * font.size * 0.5


def layout(e, w_in, fa, fe):
    """أسطرُ النصّ كما ستُرسم: [(نصّ، خطّ، rtl، محاذاة، لون، ارتفاعُ السطر px)]،
    والارتفاعُ اللازم (px)، وأعرضُ سطر (px)، وكلماتٌ لا تتّسع."""
    paras = paragraphs(e)
    wpx = max(1.0, (w_in - 2 * MARGIN_X) * PX)
    ls = _f(e.get("line_spacing"), 1.0) or 1.0
    gap = _f(e.get("spacing"), 0) * PX / 72
    out, total, widest, broken = [], 0.0, 0.0, []
    for i, p in enumerate(paras):
        txt = p["text"]
        rtl = PD.is_rtl(txt, True) if txt.strip() else True
        fam = p.get("font") or (fa if rtl else (fe or fa))
        if PD._NUMERIC.match(txt) and any(k in (fam or "").lower()
                                          for k in PD._ODD_DIGIT_FONTS):
            fam = "Cairo"                    # كما في Deck.text: أرقامٌ مألوفة
        f = _font(fam, p["size"], p["bold"])
        fake = p["bold"] and _font_file(fam, True) == _font_file(fam, False)
        if fake:                       # لا ملفَّ عريض ⟵ PowerPoint يعرّضه بنفسه
            f = _Bold(f)
        # PowerPoint: السطرُ ١٫٢ × الحجم × التباعد لكلِّ الخطوط — قِيس بـLibreOffice
        # على Kufyan وAmiri وCairo وNoto Naskh وTajawal (لا مقاييسَ الخطّ كما في Word)
        lh = p["size"] * PX / 72 * 1.2 * ls
        al = p.get("align") or ("right" if rtl else "left")
        words = txt.split(" ")
        cur = ""
        lines = []
        for wd in words:
            cand = (cur + " " + wd) if cur else wd
            if not cur or _length(f, cand, rtl) <= wpx:
                cur = cand
            else:
                lines.append(cur)
                cur = wd
            if _length(f, wd, rtl) > wpx and wd.strip():
                broken.append(wd)
        lines.append(cur)
        for k, ln in enumerate(lines):
            widest = max(widest, _length(f, ln, rtl))
            g = gap if (k == len(lines) - 1 and i < len(paras) - 1) else 0.0
            out.append((ln, f, rtl, al, _hex(p.get("color")), lh, g))
            total += lh + g
    return out, total, widest, broken


def _rect_px(x, y, w, h):
    return [int(round(x * PX)), int(round(y * PX)), int(round((x + w) * PX)),
            int(round((y + h) * PX))]


def _draw_bg(img, bg):
    from PIL import Image
    cols, ang = _bg_colors(bg)
    if len(cols) == 1:
        img.paste(_rgb(cols[0]) + (255,), [0, 0, img.width, img.height])
        return
    a, b = _rgb(cols[0]), _rgb(cols[1])
    strip = Image.new("RGBA", (256, 1))
    for i in range(256):
        t = i / 255
        strip.putpixel((i, 0), tuple(int(a[k] + (b[k] - a[k]) * t) for k in range(3))
                       + (255,))
    import math
    rad = math.radians(ang)
    L = int(abs(img.width * math.cos(rad)) + abs(img.height * math.sin(rad))) + 2
    Dg = int(math.hypot(img.width, img.height)) + 2
    # PowerPoint: ٠° من اليسار إلى اليمين، و٩٠° من الأعلى إلى الأسفل
    g = strip.resize((L, Dg)).rotate(-ang, expand=True, resample=Image.BILINEAR)
    l, t = (g.width - img.width) // 2, (g.height - img.height) // 2
    img.paste(g.crop((l, t, l + img.width, t + img.height)), (0, 0))


def _poly(name, r):
    x0, y0, x1, y1 = r
    w, h = x1 - x0, y1 - y0
    mx, my = x0 + w / 2, y0 + h / 2
    return {
        "triangle": [(mx, y0), (x1, y1), (x0, y1)],
        "right_triangle": [(x0, y0), (x1, y1), (x0, y1)],
        "diamond": [(mx, y0), (x1, my), (mx, y1), (x0, my)],
        "hexagon": [(x0 + w * .25, y0), (x1 - w * .25, y0), (x1, my),
                    (x1 - w * .25, y1), (x0 + w * .25, y1), (x0, my)],
        "chevron": [(x0, y0), (x1 - h / 2, y0), (x1, my), (x1 - h / 2, y1), (x0, y1),
                    (x0 + h / 2, my)],
        "arrow": [(x0, y0 + h * .3), (x1 - h / 2, y0 + h * .3), (x1 - h / 2, y0),
                  (x1, my), (x1 - h / 2, y1), (x1 - h / 2, y1 - h * .3),
                  (x0, y1 - h * .3)],
        "pentagon": [(x0, y0), (x1 - h / 2, y0), (x1, my), (x1 - h / 2, y1), (x0, y1)],
        "parallelogram": [(x0 + w * .2, y0), (x1, y0), (x1 - w * .2, y1), (x0, y1)],
    }.get(name)


def _shape_png(draw, e, r, fill, line, lw):
    name = str(e.get("shape") or "rect").lower()
    name = {"right_arrow": "arrow", "rectangle": "rect", "square": "rect",
            "ellipse": "oval"}.get(name, name)
    if name == "left_arrow":
        pts = _poly("arrow", r)
        cx = (r[0] + r[2]) / 2
        draw.polygon([(2 * cx - px, py) for px, py in pts], fill=fill, outline=line,
                     width=lw)
    elif name in ("oval", "circle", "donut"):
        if name == "circle":
            d = min(r[2] - r[0], r[3] - r[1])
            cx, cy = (r[0] + r[2]) / 2, (r[1] + r[3]) / 2
            r = [cx - d / 2, cy - d / 2, cx + d / 2, cy + d / 2]
        draw.ellipse(r, fill=fill, outline=line, width=lw)
    elif name in ("rounded", "pill"):
        rad = (min(r[2] - r[0], r[3] - r[1]) / 2 if name == "pill"
               else min(r[2] - r[0], r[3] - r[1]) * _f(e.get("radius"), 0.12) / 2 * 2)
        draw.rounded_rectangle(r, radius=max(1, int(rad)), fill=fill, outline=line,
                               width=lw)
    elif _poly(name, r):
        draw.polygon(_poly(name, r), fill=fill, outline=line, width=lw)
    else:
        draw.rectangle(r, fill=fill, outline=line, width=lw)


def _draw_text(img, e, fa, fe, issues, sn, label, backdrop_default=None):
    from PIL import ImageDraw
    x, y, w, h = _box(e)
    lines, need, widest, broken = layout(e, w, fa, fe)
    if not lines:
        return None
    r = _rect_px(x, y, w, h)
    # التباين: لونُ ما تحت النصّ فعلاً (من الصورة المرسومة حتى الآن)
    crop = img.crop((max(0, r[0]), max(0, r[1]), max(1, min(img.width, r[2])),
                     max(1, min(img.height, r[3])))).convert("RGB")
    try:
        under = crop.resize((1, 1)).getpixel((0, 0))
    except Exception:
        under = (255, 255, 255)
    dflt = _hex(e.get("color")) or _auto_color(under)
    va = str(e.get("valign") or "top").lower()
    inner = (h - 2 * MARGIN_Y) * PX
    top = r[1] + MARGIN_Y * PX + {"middle": (inner - need) / 2,
                                   "bottom": inner - need}.get(va, 0)
    d = ImageDraw.Draw(img)
    yy = top
    min_size = 99
    worst = 99.0
    for (ln, f, rtl, al, col, lh, g) in lines:
        c = col or dflt
        tw = _length(f, ln, rtl)
        left = r[0] + MARGIN_X * PX
        right = r[2] - MARGIN_X * PX
        xx = {"right": right - tw, "center": (left + right - tw) / 2}.get(al, left)
        _ink(d, (xx, yy + _glyph_top(f, lh)), ln, f, _rgb(c), rtl)
        min_size = min(min_size, f.size * 72 / PX)
        big = f.size * 72 / PX >= 24
        worst = min(worst, contrast(_rgb(c), under) / (3.0 if big else 4.5))
        yy += lh + g
    avail = (h - 2 * MARGIN_Y) * PX
    if need > avail * 1.03 + 2:
        issues.append("slide %d · %s: text needs %.2fin height, box is %.2fin — "
                      "make h ≥ %.2f or reduce size" % (
                          sn, label, need / PX + 2 * MARGIN_Y, h,
                          need / PX + 2 * MARGIN_Y + 0.05))
    for wd in broken[:2]:
        issues.append("slide %d · %s: word «%s» is wider than the box (w %.2fin) — "
                      "widen it or reduce size" % (sn, label, wd[:30], w))
    if min_size < 12 - 0.01:
        issues.append("slide %d · %s: font %.0fpt is too small to read — ≥ 12" % (
            sn, label, min_size))
    if worst < 1:
        issues.append("slide %d · %s: low contrast between text and what is under it "
                      "— change the text color or the background" % (sn, label))
    # مستطيلُ الحبر الفعليّ (للتداخل)
    ink_h = min(need, avail) / PX
    ink_w = min(widest / PX + 2 * MARGIN_X, w)
    al0 = lines[0][3]
    ix = {"right": x + w - ink_w, "center": x + (w - ink_w) / 2}.get(al0, x)
    iy = (top - MARGIN_Y * PX) / PX + MARGIN_Y
    return (ix, iy, ink_w, max(ink_h, 0.01))


def _label(e, i):
    t = str(e.get("type") or "")
    txt = " ".join(p["text"] for p in paragraphs(e))[:28]
    return "%s #%d%s" % (t or "element", i, (" «%s»" % txt) if txt else "")


def _paste_image(img, path, r, shape):
    from PIL import Image, ImageDraw
    w, h = max(1, r[2] - r[0]), max(1, r[3] - r[1])
    im = Image.open(path).convert("RGBA")
    ia, fa_ = im.width / im.height, w / h
    if ia > fa_:
        nw = int(im.height * fa_)
        im = im.crop(((im.width - nw) // 2, 0, (im.width - nw) // 2 + nw, im.height))
    else:
        nh = int(im.width / fa_)
        im = im.crop((0, (im.height - nh) // 2, im.width, (im.height - nh) // 2 + nh))
    im = im.resize((w, h))
    mask = Image.new("L", (w, h), 0)
    md = ImageDraw.Draw(mask)
    if shape in ("circle", "oval"):
        md.ellipse([0, 0, w - 1, h - 1], fill=255)
    elif shape == "rounded":
        md.rounded_rectangle([0, 0, w - 1, h - 1], radius=int(min(w, h) * .1), fill=255)
    else:
        md.rectangle([0, 0, w, h], fill=255)
    img.paste(im, (r[0], r[1]), mask)


def _is_container(e):
    """شكلٌ يحمل محتوى (بطاقة، دائرة، لوح): معتمٌ، داخلَ الشريحة، ليس خلفيّةً كاملة
    ولا شارةً صغيرة ولا خطّاً."""
    x, y, w, h = _box(e)
    if str(e.get("shape") or "") == "line" or not e.get("fill"):
        return False
    if e.get("opacity") is not None and _f(e.get("opacity"), 1) < 0.5:
        return False
    if w * h < 1.5 or w * h > 0.6 * W * H:
        return False
    return x >= -0.01 and y >= -0.01 and x + w <= W + 0.01 and y + h <= H + 0.01


def _balance(boxes, inks, issues, sn):
    """محتوى البطاقة/الدائرة محشورٌ في طرفها وأغلبُها فارغ — قِيس على عرض الهاتف:
    دوائرُ فيها عنوانٌ في أعلاها فقط، وبطاقاتٌ نصُّها الصغيرُ في ثلثها الأعلى."""
    for (bx, by, bw, bh), lab, own in boxes:
        rects = [own] if own else []
        for rect, l2, _k in inks:
            if l2 == lab or rect is own:
                continue
            cx, cy = rect[0] + rect[2] / 2, rect[1] + rect[3] / 2
            if bx <= cx <= bx + bw and by <= cy <= by + bh and rect[3] < bh:
                rects.append(rect)
        if not rects:
            continue
        top = max(by, min(r[1] for r in rects))
        bot = min(by + bh, max(r[1] + r[3] for r in rects))
        fill = (bot - top) / bh if bh else 1
        off = ((top + bot) / 2 - (by + bh / 2)) / bh if bh else 0
        if fill < 0.45 and abs(off) > 0.18:
            issues.append(
                "slide %d · %s: its content fills only %d%% of its height and sits at "
                "the %s — center it (valign middle / move it), enlarge the text, add "
                "the missing explanation, or make the shape smaller" % (
                    sn, lab, round(fill * 100), "top" if off < 0 else "bottom"))


def _coverage(boxes, inks, issues, sn):
    """شريحةٌ محتواها في ركنٍ صغيرٍ منها (٣ عناصرَ فأكثر). الغلافُ والختامُ — عنصران
    — خارجَ الحساب: فراغُهما مقصود."""
    rects = [r for r, _l, _k in inks] + [b for b, _l, _o in boxes]
    if len(inks) < 3 or not rects:
        return
    x0 = max(0.0, min(r[0] for r in rects))
    y0 = max(0.0, min(r[1] for r in rects))
    x1 = min(W, max(r[0] + r[2] for r in rects))
    y1 = min(H, max(r[1] + r[3] for r in rects))
    cov = max(0.0, x1 - x0) * max(0.0, y1 - y0) / (W * H)
    if cov < 0.35:
        issues.append("slide %d: the content occupies only %d%% of the slide — enlarge "
                      "the elements or spread them over the slide" % (sn, round(cov * 100)))


def _overlap(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ix = max(0.0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0.0, min(ay + ah, by + bh) - max(ay, by))
    small = max(1e-6, min(aw * ah, bw * bh))
    return ix * iy / small


def render_check(spec, out_dir, stem, lang, font_ar=None, font_en=None,
                 chart_png=None):
    """صورةٌ لكلِّ شريحة + صورةٌ جامعة، وقائمةُ مشكلاتٍ مقيسة.
    يعيد (المسارات، صورة_جامعة، المشكلات)."""
    from PIL import Image, ImageDraw
    fa = font_ar or "Kufyan Arabic Regular"
    fe = font_en or fa
    os.makedirs(out_dir, exist_ok=True)
    slides = [sd for sd in spec.get("slides") or [] if isinstance(sd, dict)]
    issues, paths = [], []
    for sn, sd in enumerate(slides, 1):
        img = Image.new("RGBA", (int(W * PX), int(H * PX)), (255, 255, 255, 255))
        bg = sd.get("bg") if sd.get("bg") is not None else spec.get("bg", "FFFFFF")
        _draw_bg(img, bg)
        inks = []                         # (مستطيل، تسمية، نوع)
        boxes = []                        # أشكالٌ حاويةٌ (بطاقات، دوائر) للتوازن
        els = [e for e in sd.get("elements") or [] if isinstance(e, dict)]
        if not els:
            issues.append("slide %d: empty — no elements" % sn)
        for i, e in enumerate(els, 1):
            t = str(e.get("type") or ("text" if "text" in e else "shape")).lower()
            x, y, w, h = _box(e)
            lab = _label(e, i)
            if t != "shape" or str(e.get("shape")) != "line":
                if w <= 0 or h <= 0:
                    issues.append("slide %d · %s: needs w and h" % (sn, lab))
                    continue
                if x < -0.02 or y < -0.02 or x + w > W + 0.02 or y + h > H + 0.02:
                    if t != "shape":       # الأشكالُ الزخرفيّةُ قد تخرج عمداً
                        issues.append("slide %d · %s: outside the slide (x %.2f y %.2f "
                                      "w %.2f h %.2f; slide is %.2f×%.2f)" % (
                                          sn, lab, x, y, w, h, W, H))
            r = _rect_px(x, y, w, h)
            if t == "shape":
                layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
                dd = ImageDraw.Draw(layer)
                fill = e.get("fill")
                op = max(0.0, min(1.0, _f(e.get("opacity"), 1)))
                a = int(255 * op)
                if str(e.get("shape")) == "line":
                    dd.line([r[0], r[1], r[2], r[3]], fill=_rgb(e.get("line") or
                                                                e.get("fill") or "999999")
                            + (a,), width=max(1, int(_f(e.get("line_w"), 2) * PX / 72)))
                else:
                    if isinstance(fill, dict):
                        g = Image.new("RGBA", (max(1, r[2] - r[0]), max(1, r[3] - r[1])))
                        _draw_bg(g, fill)
                        m = Image.new("L", g.size, 0)
                        _shape_png(ImageDraw.Draw(m), e, [0, 0, g.width - 1,
                                                          g.height - 1], 255, None, 0)
                        g.putalpha(m.point(lambda v: v * op))
                        layer.paste(g, (r[0], r[1]), g)
                    lc = _hex(e.get("line"))
                    _shape_png(dd, e, r, (_rgb(fill) + (a,)) if fill and not
                               isinstance(fill, dict) else None,
                               (_rgb(lc) + (255,)) if lc else None,
                               max(1, int(_f(e.get("line_w"), 1) * PX / 72)) if lc else 0)
                img.alpha_composite(layer)
                own = None
                if e.get("text") is not None:
                    own = _draw_text(img, e, fa, fe, issues, sn, lab)
                    if own:
                        inks.append((own, lab, "text"))
                if _is_container(e):
                    boxes.append(((x, y, w, h), lab, own))
            elif t == "text":
                ink = _draw_text(img, e, fa, fe, issues, sn, lab)
                if ink:
                    inks.append((ink, lab, "text"))
            elif t == "image":
                p = e.get("path") or e.get("image")
                shp = str(e.get("shape") or "rect").lower()
                if p and os.path.isfile(str(p)):
                    try:
                        _paste_image(img, str(p), r, shp)
                    except Exception:
                        pass
                else:
                    # إطارٌ فارغ كما يُبنى: تعبئةٌ فاتحة وحدٌّ وعلامة «+»
                    d = ImageDraw.Draw(img)
                    if shp == "circle":
                        d.ellipse(r, fill=(232, 234, 240), outline=(200, 170, 100),
                                  width=3)
                    else:
                        d.rounded_rectangle(r, radius=int(min(r[2] - r[0], r[3] - r[1])
                                                          * 0.08), fill=(232, 234, 240),
                                            outline=(200, 170, 100), width=3)
                    fpl = _font("Cairo", max(18, min(w, h) * 22))
                    cx, cy = (r[0] + r[2]) / 2, (r[1] + r[3]) / 2
                    d.text((cx, cy), "+", font=fpl, fill=(140, 145, 160), anchor="mm")
                inks.append(((x, y, w, h), lab, "image"))
            elif t == "chart":
                ok = False
                if chart_png:
                    try:
                        sp = dict(e.get("chart") or {})
                        if e.get("colors"):
                            sp["colors"] = e["colors"]
                        sp.setdefault("lang", lang)
                        sp.setdefault("font", fa)
                        png = chart_png(sp)
                        ci = Image.open(png).convert("RGBA")
                        ci.thumbnail((max(1, r[2] - r[0]), max(1, r[3] - r[1])))
                        img.alpha_composite(ci, (r[0] + (r[2] - r[0] - ci.width) // 2,
                                                 r[1] + (r[3] - r[1] - ci.height) // 2))
                        ok = True
                    except Exception:
                        ok = False
                if not ok:
                    d = ImageDraw.Draw(img)
                    d.rectangle(r, outline=(150, 150, 160), width=2)
                    d.text((r[0] + 10, r[1] + 10), "chart", font=_font(fe, 18),
                           fill=(120, 120, 130))
                inks.append(((x, y, w, h), lab, "chart"))
            elif t == "table":
                _draw_table(img, e, r, fa, fe, issues, sn, lab)
                inks.append(((x, y, w, h), lab, "table"))
        _balance(boxes, inks, issues, sn)
        if len(slides) == 1 or sn not in (1, len(slides)):   # الغلافُ والختامُ: فراغٌ مقصود
            _coverage(boxes, inks, issues, sn)
        for i in range(len(inks)):
            for j in range(i + 1, len(inks)):
                (a, la, ka), (b, lb, kb) = inks[i], inks[j]
                if "text" not in (ka, kb):
                    continue
                if _overlap(a, b) > 0.15:
                    issues.append("slide %d: %s overlaps %s — move one of them" % (
                        sn, la, lb))
        p = os.path.join(out_dir, "%s-s%02d.png" % (stem, sn))
        img.convert("RGB").save(p, optimize=True)
        paths.append(p)
    ov = _overview(paths, os.path.join(out_dir, "%s-overview.png" % stem))
    return paths, ov, issues


def _draw_table(img, e, r, fa, fe, issues, sn, lab):
    from PIL import ImageDraw
    heads = [str(v) for v in e.get("headers") or []]
    rows = [[str(v) for v in rw] for rw in e.get("rows") or []]
    ncol = max([len(heads)] + [len(rw) for rw in rows] + [1])
    heads += [""] * (ncol - len(heads))
    rows = [rw + [""] * (ncol - len(rw)) for rw in rows]
    rtl = PD.is_rtl(" ".join(heads + (rows[0] if rows else [])), True)
    if rtl:
        heads, rows = heads[::-1], [rw[::-1] for rw in rows]
    nrow = 1 + len(rows)
    lens = [max([len(heads[c])] + [len(rw[c]) for rw in rows]) + 4 for c in range(ncol)]
    W_ = r[2] - r[0]
    rh = (r[3] - r[1]) / nrow
    size = _f(e.get("size"), 16)
    hf, hc = _hex(e.get("header_fill"), "1B2A4A"), _hex(e.get("header_color"), "FFFFFF")
    f1 = _hex(e.get("fill"), "FFFFFF")
    f2 = _hex(e.get("alt_fill"), PD._mix(f1, "000000", 0.05))
    col = _hex(e.get("color"), _auto_color(_rgb(f1)))
    d = ImageDraw.Draw(img)
    x = r[0]
    over = False
    for c in range(ncol):
        cw = W_ * lens[c] / sum(lens)
        for rr in range(nrow):
            y = r[1] + rr * rh
            d.rectangle([x, y, x + cw, y + rh], fill=_rgb(hf if rr == 0 else
                                                         (f1 if rr % 2 else f2)),
                        outline=(255, 255, 255), width=1)
            txt = heads[c] if rr == 0 else rows[rr - 1][c]
            cell = {"text": txt, "size": size, "bold": rr == 0,
                    "color": hc if rr == 0 else col}
            lines, need, _w, _b = layout(cell, cw / PX, fa, fe)
            if need > rh * 1.05:
                over = True
            yy = y + max(0, (rh - need) / 2)
            for (ln, f, rt, al, cc, lh, _g) in lines:
                tw = _length(f, ln, rt)
                xx = x + cw - MARGIN_X * PX - tw if rt else x + MARGIN_X * PX
                _ink(d, (xx, yy + _glyph_top(f, lh)), ln, f, _rgb(cc), rt)
                yy += lh
        x += cw
    if over:
        issues.append("slide %d · %s: cell text does not fit its row — make the table "
                      "taller, use fewer rows, or reduce size" % (sn, lab))


def _overview(paths, out):
    from PIL import Image, ImageDraw
    if not paths:
        return None
    cols = 3
    tw, th = 480, 270
    rows = (len(paths) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * (tw + 12) + 12, rows * (th + 34) + 12),
                      (40, 40, 48))
    d = ImageDraw.Draw(sheet)
    for i, p in enumerate(paths):
        im = Image.open(p)
        im.thumbnail((tw, th))
        x = 12 + (i % cols) * (tw + 12)
        y = 12 + (i // cols) * (th + 34)
        sheet.paste(im, (x, y + 22))
        d.text((x, y + 2), "%d" % (i + 1), fill=(230, 230, 235),
               font=_font("Cairo", 13))
    sheet.save(out, optimize=True)
    return out
